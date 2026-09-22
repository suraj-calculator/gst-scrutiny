#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GSTR-1 CANONICAL LAYER  (spec: docs/GSTR1_CANONICAL_SPEC.md)
=============================================================
Only this module knows what a raw (merged) GSTR-1 workbook looks like.

    raw / merged GSTR-1    --build_canonical()-->  canonical GSTR-1 data
    canonical data         --write_canonical()-->  Canonical_GSTR1_<GSTIN>_<FY>.xlsx
    canonical .xlsx        --read_canonical()  -->  canonical data
    canonical data         --parse_gstr1() / parse_b2ba() / read_gstr1_lines() / ...
                            --> exactly what the 13 legacy GSTR-1 readers return today

Every GSTR-1 tab has the same shape (title, ONE heading row, a period-banner row per month with that month's
data under it), so one generic routine normalises them all. Each canonical table keeps the source rows with
canonical column names and the period on every row; numbers are converted (blank = 0), identifiers, text and
dates are kept exactly as in the source. Each reader keeps its own business logic (forward-filling the
continuation rows of a multi-rate invoice, counting invoices, ...) and simply runs it over these rows.

Problems are reported in plain words (ISSUES sheet, [warn]/[error] log lines). A table whose required columns
cannot be found raises CanonicalLayoutError for that month (loud), never a silent zero.
"""

import datetime as _dt
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

import canonical_common as cc
import gst_core as mpu

SCHEMA_VERSION = "1-1.0"
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "gstr1_mapping.json")

LEAD_COLS = ["period", "months_covered", "src_sheet", "src_row", "_any"]
RETURN_COLS = ["month", "fy", "tax_period", "arn", "arn_date", "src_sheet", "marker_text"]
COVERAGE_COLS = ["table", "month", "rows", "note"]
MAP_COLS = ["canonical_sheet", "canonical_field", "raw_sheet", "raw_header_found", "raw_column", "matched_by",
            "confidence", "status", "rows_filled", "rows_blank", "note"]
ISSUE_COLS = cc.ISSUE_COLS

# The ARN / ARN-date printed on a period banner (copied from gst_checks_forensic so both give the same result)
_ARN_IN_MARKER_RE = re.compile(
    r"ARN:\s*([A-Z0-9]{15})\s*(?:[|,])?\s*"
    r"(?:(?:ARN\s*Date|Date\s*of\s*Filing|Filing\s*Date|Date)\s*[:\s]\s*)?"
    r"(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})?",
    re.IGNORECASE)
_DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y", "%d/%m/%y")

_CONTEXT = {"gstin": None}


def set_context(gstin=None):
    if gstin:
        _CONTEXT["gstin"] = gstin


def _mapping():
    return cc.load_json(_MAPPING_FILE)


def _tables():
    return _mapping()["tables"]


def _parse_any_date(s):
    if not s:
        return None
    s = str(s).strip()
    for fmt in _DATE_FORMATS:
        try:
            return _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _canon_cols(tname):
    return LEAD_COLS + list(_tables()[tname]["fields"]) + ["extra"]


def _label(labels):
    return labels[0] if len(labels) == 1 else f"{labels[0]}..{labels[-1]}"


# ----------------------------------------------------------------------------------------
# reading the raw workbook
# ----------------------------------------------------------------------------------------
def _scan_readme(rows):
    """'Read me' key/value rows (last occurrence wins, as the legacy readers did)."""
    labels = _mapping()["readme_labels"]
    out = {}
    for row in rows:
        cells = [str(c).strip() for c in row if c not in (None, "")]
        if len(cells) < 2:
            continue
        key = cells[0].lower()
        for field, names in labels.items():
            if key in names:
                out[field] = cells[1] if field in ("gstin", "legal_name") else cells[-1]
    return out


def _scan_markers(rows, sheet, returns):
    """Every period banner on the sheet -> ARN / ARN date per month (first non-empty value wins)."""
    found_date = False
    for row in rows:
        if not row or not row[0]:
            continue
        cell0 = str(row[0])
        if "Financial Year:" not in cell0 or "Tax Period:" not in cell0:
            continue
        try:
            fy, tp, labels = mpu.parse_marker_text(cell0)
        except mpu.PeriodParseError:
            continue
        m = _ARN_IN_MARKER_RE.search(cell0)
        arn = m.group(1) if m else None
        dt = _parse_any_date(m.group(2)) if (m and m.group(2)) else None
        if dt:
            found_date = True
        for lbl in labels:
            slot = returns.setdefault(lbl, dict(month=lbl, fy=fy, tax_period=tp, arn=None, arn_date=None,
                                                src_sheet=sheet, marker_text=cell0))
            if arn and not slot["arn"]:
                slot["arn"] = arn
            if dt and not slot["arn_date"]:
                slot["arn_date"] = dt
    return found_date


def _resolve_field(hmap, spec):
    """-> (col or None, matched_by, confidence). `hmap` = {clean heading: column} (last duplicate wins)."""
    for pos, name in enumerate(spec.get("exact", [])):
        col = hmap.get(cc.clean(name))
        if col is not None:
            return col, ("exact" if pos == 0 else "alias"), "HIGH"
    for frag in spec.get("contains", []):
        f = cc.clean(frag)
        for h, col in hmap.items():
            if f in h:
                return col, "contains", "MEDIUM"
    return None, "", ""


def _build_table(ctx, tname, tdef, sheet, rows, table_rows, coverage, mapping):
    fields = tdef["fields"]
    first = next((i for i, r in enumerate(rows) if r and mpu.is_marker_row(r)), None)
    if first is None:
        ctx.issue("W101", "WARNING", sheet, "", "all",
                  f"Tab '{sheet}' has no 'Financial Year | Tax Period' banner rows, so no month can be read "
                  f"from it.", f"every check that reads {sheet}", "Give the GSTR-1 as merged from the portal downloads.")
        return
    # The heading row is the row above the first banner that matches the MOST known columns (ties: the
    # nearest). "Nearest non-empty row" alone would take a stray row wedged between the heading and the banner
    # for the heading and silently map nothing.
    hdr_idx, best = None, 0
    for i in range(first - 1, max(-1, first - 13), -1):
        if not any(c not in (None, "") for c in rows[i]):
            continue
        hm = {cc.clean(c): j for j, c in enumerate(rows[i]) if c not in (None, "")}
        score = sum(1 for spec in fields.values() if _resolve_field(hm, spec)[0] is not None)
        if hdr_idx is None or score > best:
            hdr_idx, best = i, score
    if hdr_idx is not None and best == 0:
        ctx.issue("W105", "WARNING", sheet, "", "all",
                  f"Tab '{sheet}': none of the expected column headings were found above the first period banner; "
                  f"its figures will read as blank/0.", f"checks that read {sheet}",
                  "Add the new headings to gstr1_mapping.json.")
    if hdr_idx is None:
        ctx.issue("E102", "ERROR", sheet, "", "all", f"Tab '{sheet}': no heading row found above the first period banner.",
                  f"every check that reads {sheet}", "")
        coverage.append(dict(table=sheet, month="*", rows=-1,
                             note=f"[E102] Tab '{sheet}': no heading row found above the first period banner."))
        return
    header = rows[hdr_idx]
    hmap = {cc.clean(c): i for i, c in enumerate(header) if c not in (None, "")}
    texts = {i: str(c).strip() for i, c in enumerate(header) if c not in (None, "")}
    cols, missing_required = {}, []
    for f, spec in fields.items():
        col, how, conf = _resolve_field(hmap, spec)
        cols[f] = col
        status = "OK" if col is not None else ("NOT_IN_THIS_TABLE" if spec.get("absent_ok") else "MISSING")
        if col is None and spec.get("required"):
            missing_required.append(f)
        mapping.append(dict(canonical_sheet=tdef["canon_sheet"], canonical_field=f, raw_sheet=sheet,
                            raw_header_found=texts.get(col, "") if col is not None else "",
                            raw_column=get_column_letter(col + 1) if col is not None else "", matched_by=how,
                            confidence=conf, status=status, rows_filled=0, rows_blank=0, note=""))
    optional_missing = [f for f, c in cols.items() if c is None and not fields[f].get("required")
                        and not fields[f].get("absent_ok")]
    if optional_missing:
        ctx.issue("I103", "INFO", sheet, ", ".join(optional_missing), "all",
                  f"Tab '{sheet}': column(s) {optional_missing} not found; read as blank/0 where the checks use them.")
    if missing_required:
        names = [fields[f]["exact"][0] for f in missing_required]
        msg = (f"[E101] GSTR-1 tab '{sheet}': required column(s) {names} not found. Headings found: "
               f"{sorted(texts.values())}. Every check that reads {sheet} is skipped for the run; add the new "
               f"heading to gstr1_mapping.json or supply the file as downloaded from the portal.")
        ctx.issue("E101", "ERROR", sheet, ", ".join(missing_required), "all", msg[7:],
                  f"every check that reads {sheet}", "Add the new spelling to gstr1_mapping.json.")
        coverage.append(dict(table=sheet, month="*", rows=-1, note=msg))
        return
    mapped = {c for c in cols.values() if c is not None}
    unmapped = {i: t for i, t in texts.items() if i not in mapped}

    n_by_month, months_order, table_error = {}, [], None
    current = None
    canon = tdef["canon_sheet"]
    fill = {f: [0, 0] for f in fields}
    for ridx in range(hdr_idx + 1, len(rows)):
        r = rows[ridx]
        if r and mpu.is_marker_row(r):
            try:
                _, _, labels = mpu.parse_marker_text(str(r[0]))
            except mpu.PeriodParseError as e:
                table_error = f"[E103] GSTR-1 tab '{sheet}': period banner not understood ({e})."
                ctx.issue("E103", "ERROR", sheet, "", "all", table_error[7:], f"every check that reads {sheet}", "")
                break
            current = labels
            for lbl in labels:
                if lbl not in n_by_month:
                    n_by_month[lbl] = 0
                    months_order.append(lbl)
            continue
        if not r or not any(c not in (None, "") for c in r):
            continue
        if current is None:
            table_error = (f"[E104] GSTR-1 tab '{sheet}': data row {ridx + 1} sits before any period banner, so its "
                           f"month cannot be determined: {r!r}")
            ctx.issue("E104", "ERROR", sheet, "", "all", table_error[7:], f"every check that reads {sheet}", "")
            break
        rec = dict(period=_label(current), months_covered=";".join(current), src_sheet=sheet, src_row=ridx + 1,
                   _any=int(bool(any(r))))
        for f, spec in fields.items():
            col = cols[f]
            if col is None:
                rec[f] = None
                continue
            v = r[col] if col < len(r) else None
            fill[f][0 if v not in (None, "") else 1] += 1
            rec[f] = cc.num(v) if spec["type"] == "num" else v
        extras = [f"{t}={r[i]}" for i, t in unmapped.items() if i < len(r) and r[i] not in (None, "")]
        rec["extra"] = "; ".join(extras)
        table_rows[canon].append(rec)
        for lbl in current:
            n_by_month[lbl] += 1
    if table_error:
        coverage.append(dict(table=sheet, month="*", rows=-1, note=table_error))
    for lbl in months_order:
        coverage.append(dict(table=sheet, month=lbl, rows=n_by_month[lbl], note=""))
    for m in mapping:
        if m["raw_sheet"] == sheet and m["canonical_field"] in fill:
            m["rows_filled"], m["rows_blank"] = fill[m["canonical_field"]]


def build_canonical(raw_paths):
    """Convert one raw/merged GSTR-1 workbook into canonical data (a dict of tables). A file that cannot be
    opened, or that has no GSTR-1 tab at all, raises PeriodParseError ('[E10x] ...')."""
    if isinstance(raw_paths, str):
        raw_paths = [raw_paths]
    path = raw_paths[0]
    ctx = cc.IssueLog("GSTR-1", [os.path.basename(p) for p in raw_paths])
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        ctx.fail("E105", f"GSTR-1 \"{os.path.basename(path)}\": cannot open the file ({type(e).__name__}: {e}). "
                 f"It may be corrupt, password-protected or not an Excel file.", skipped="every check that reads GSTR-1",
                 action="Re-download the file.")
    tables = _tables()
    raw_to_table = {rn.lower(): tname for tname, t in tables.items() for rn in t["raw_names"]}
    returns, coverage, mapping = {}, [], []
    table_rows = {t["canon_sheet"]: [] for t in tables.values()}
    tabs_present, unread, readme = [], [], {}
    found_marker_date = False
    try:
        names = list(wb.sheetnames)
        readme_sheet = next((n for n in names if n.lower() == "read me"), names[0] if names else None)
        for sn in names:
            rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
            if sn == readme_sheet:
                readme = _scan_readme(rows)
            if sn.lower() == "read me":
                continue
            found_marker_date |= _scan_markers(rows, sn, returns)
            tname = raw_to_table.get(sn.lower())
            if tname is None:
                unread.append(sn)
                continue
            tabs_present.append(sn)
            _build_table(ctx, tname, tables[tname], sn, rows, table_rows, coverage, mapping)
    finally:
        wb.close()
    if not tabs_present or "b2b, sez, de_inv" not in tabs_present:
        ctx.fail("E106", f"GSTR-1 \"{os.path.basename(path)}\": no 'b2b, sez, de_inv' tab found (tabs: {names}). "
                 f"This does not look like a merged GSTR-1.", skipped="every check that reads GSTR-1",
                 action="Check that this is the merged GSTR-1 workbook.")
    if not any(m["status"] == "OK" for m in mapping if m["raw_sheet"] == "b2b, sez, de_inv"):
        ctx.fail("E107", "GSTR-1 tab 'b2b, sez, de_inv' has no recognisable column headings.",
                 skipped="every check that reads GSTR-1", action="Check the file.")
    ret_rows = list(returns.values())
    b2b_months = [c["month"] for c in coverage if c["table"] == "b2b, sez, de_inv" and c["month"] != "*"]
    fys = sorted({r["fy"] for r in ret_rows})
    gstin = _CONTEXT["gstin"] or readme.get("gstin") or ""
    if unread:
        ctx.issue("I401", "INFO", "", "", "all", f"Tabs not used by any check: {unread}.")
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    meta = [("schema_version", SCHEMA_VERSION), ("recipient_gstin", gstin), ("legal_name", readme.get("legal_name", "")),
            ("fy", ", ".join(fys)), ("source_files", "; ".join(ctx.source_files)), ("source_layout", "merged with period banners"),
            ("adapter_version", ADAPTER_VERSION), ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("months_present", ";".join(b2b_months)), ("tabs_present", ";".join(tabs_present)),
            ("unread_tabs", "; ".join(unread)), ("marker_dates_found", "Y" if found_marker_date else "N"),
            ("readme_gstin", readme.get("gstin", "")), ("readme_legal_name", readme.get("legal_name", "")),
            ("readme_arn", readme.get("arn", "")), ("readme_arn_date", readme.get("arn_date", "")),
            ("readme_tax_period", readme.get("tax_period", "")),
            ("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK"))]
    return dict(meta=meta, returns=ret_rows, coverage=coverage, tables=table_rows, mapping=mapping,
                issues=ctx.issues)


# ----------------------------------------------------------------------------------------
# write / read the canonical workbook
# ----------------------------------------------------------------------------------------
def canonical_filename(data):
    return cc.canonical_filename("GSTR1", data["meta"])


def write_canonical(data, out_dir):
    sheets = [("RETURNS", RETURN_COLS, data["returns"]), ("COVERAGE", COVERAGE_COLS, data["coverage"])]
    for tname, tdef in _tables().items():
        sheets.append((tdef["canon_sheet"], _canon_cols(tname), data["tables"][tdef["canon_sheet"]]))
    sheets += [("MAPPING_REPORT", MAP_COLS, data["mapping"]), ("ISSUES", ISSUE_COLS, data["issues"])]
    return cc.write_workbook(out_dir, canonical_filename(data), data["meta"], sheets, max_col_width=30)


def is_canonical_file(path):
    return cc.is_canonical(path, "1-", ("META", "RETURNS", "COVERAGE"))


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        tables = {t["canon_sheet"]: cc.read_table(wb, t["canon_sheet"], _canon_cols(n)) for n, t in _tables().items()}
        return dict(meta=cc.read_meta(wb), returns=cc.read_table(wb, "RETURNS", RETURN_COLS),
                    coverage=cc.read_table(wb, "COVERAGE", COVERAGE_COLS), tables=tables,
                    mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS),
                    issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------
# engine-facing access
# ----------------------------------------------------------------------------------------
_SRC = cc.CanonicalCache("GSTR-1", is_canonical_file, read_canonical, build_canonical, write_canonical,
                         lambda data: _index(data))


def get_data(path, out_dir=None):
    return _SRC.get(path, out_dir)


def clear_cache():
    _SRC.clear()


def _index(data):
    meta = dict(data["meta"])
    tabs = [t for t in str(meta.get("tabs_present") or "").split(";") if t]
    cov = {t: {"months": [], "error": None} for t in tabs}
    for c in data["coverage"]:
        info = cov.setdefault(c["table"], {"months": [], "error": None})
        if c["month"] == "*":
            info["error"] = c["note"]
        elif c["month"] not in info["months"]:
            info["months"].append(c["month"])
    rows = {t: [] for t in tabs}
    for tname, tdef in _tables().items():
        for r in data["tables"].get(tdef["canon_sheet"], []):
            r["_ml"] = set(str(r["months_covered"]).split(";"))
            rows.setdefault(r["src_sheet"], []).append(r)
    ok_fields = {(m["raw_sheet"], m["canonical_field"]): m["status"] == "OK" for m in data["mapping"]}
    return dict(meta=meta, tabs=tabs, cov=cov, rows=rows, ok=ok_fields)


def _rows_for_month(ix, tab, month):
    """Rows of `tab` under a banner covering `month`; raises what mpu.rows_for_month raised for the raw sheet."""
    info = ix["cov"].get(tab, {"months": [], "error": None})
    if info["error"]:
        if info["error"].startswith(("[E101]", "[E102]")):
            raise cc.CanonicalLayoutError(info["error"])
        raise mpu.PeriodParseError(info["error"])
    if month not in info["months"]:
        raise mpu.PeriodParseError(f"Month {month!r} not found as a period marker in this sheet. "
                                   f"Months present: {sorted(info['months'])}")
    return [r for r in ix["rows"].get(tab, []) if month in r["_ml"]]


def _f(v):
    return 0.0 if v is None else float(v)


def _s(v):
    return str(v or "").strip()


B2B, CDNR, CDNUR, B2CL, B2CS, EXP, EXEMP, B2BA, CDNRA, DOCS = (
    "b2b, sez, de_inv", "cdnr", "cdnur", "b2cl", "b2cs", "exp", "exemp", "b2ba", "cdnra", "docs")


# ---- gst_core -------------------------------------------------------------------------------------
def months(path):
    """Months that have a period banner on the b2b tab (legacy gst_core._gstr1_months)."""
    return set(get_data(path)["_index"]["cov"].get(B2B, {"months": []})["months"])


def gstin_and_name(path):
    """(GSTIN, Legal Name) from the 'Read me' sheet (legacy gst_core._read_me_gstin_and_name)."""
    m = get_data(path)["_index"]["meta"]
    return (m.get("readme_gstin") or None, m.get("readme_legal_name") or None)


# ---- parse_gstr1 ----------------------------------------------------------------------------------
def parse_gstr1(path, month):
    """Legacy gst_parsers_returns.parse_gstr1(path, month)."""
    ix = get_data(path)["_index"]
    out = {"taxable": 0.0, "IGST": 0.0, "CGST": 0.0, "SGST": 0.0, "CESS": 0.0,
           "b2b_count": 0, "b2b_no_irn": 0, "lines": {}, "blank_invno_lines": 0, "blank_invno_taxable": 0.0,
           "named_taxable": 0.0, "named_IGST": 0.0, "named_CGST": 0.0, "named_SGST": 0.0,
           "cn_taxable": 0.0, "cn_IGST": 0.0, "cn_CGST": 0.0, "cn_SGST": 0.0, "cn_CESS": 0.0,
           "hsn_IGST": 0.0, "hsn_CGST": 0.0, "hsn_SGST": 0.0, "hsn_CESS": 0.0, "hsn_taxable": 0.0,
           "b2cs_taxable": 0.0, "b2cs_IGST": 0.0, "b2cs_CGST": 0.0, "b2cs_SGST": 0.0, "b2cs_CESS": 0.0,
           "nil_taxable": 0.0, "exempt_taxable": 0.0, "nongst_taxable": 0.0, "nil_exempt_taxable": None,
           "month_marker_gaps": []}
    tabs = ix["tabs"]
    if B2B in tabs:
        try:
            rows = _rows_for_month(ix, B2B, month)
            has_irn = ix["ok"].get((B2B, "irn"), False)
            last_gstin, last_invno = None, None
            for r in rows:
                if not r["_any"]:
                    continue
                t, ig, cg, sg, ce = _f(r["taxable"]), _f(r["igst"]), _f(r["cgst"]), _f(r["sgst"]), _f(r["cess"])
                out["taxable"] += t
                out["IGST"] += ig
                out["CGST"] += cg
                out["SGST"] += sg
                out["CESS"] += ce
                if not has_irn or not str(r["irn"]).strip():
                    out["b2b_no_irn"] += 1
                raw_gstin, raw_invno = r["gstin"], r["invno"]
                if raw_gstin not in (None, "") and raw_invno not in (None, ""):
                    last_gstin = str(raw_gstin).strip()
                    last_invno = str(raw_invno).strip()
                    gstin_key, invno = last_gstin, last_invno
                    out["b2b_count"] += 1
                elif last_gstin is not None:
                    gstin_key, invno = last_gstin, last_invno
                else:
                    gstin_key, invno = "", "None"
                k = (invno, gstin_key, _f(r["rate"]))
                L = out["lines"].setdefault(k, [0.0, 0.0])
                L[0] += t
                L[1] += ig
                if not invno or invno.lower() == "none":
                    out["blank_invno_lines"] += 1
                    out["blank_invno_taxable"] += t
                else:
                    out["named_taxable"] += t
                    out["named_IGST"] += ig
                    out["named_CGST"] += cg
                    out["named_SGST"] += sg
        except mpu.PeriodParseError:
            out["month_marker_gaps"].append(B2B)
    for sn in (B2CL, EXP):
        if sn in tabs:
            try:
                for r in _rows_for_month(ix, sn, month):
                    if not r["_any"]:
                        continue
                    out["taxable"] += _f(r["taxable"])
                    out["IGST"] += _f(r["igst"])
                    out["CESS"] += _f(r["cess"])
            except mpu.PeriodParseError:
                out["month_marker_gaps"].append(sn)
    if B2CS in tabs:
        try:
            for r in _rows_for_month(ix, B2CS, month):
                if not r["_any"]:
                    continue
                t, i_, c_, s_, ce = _f(r["taxable"]), _f(r["igst"]), _f(r["cgst"]), _f(r["sgst"]), _f(r["cess"])
                out["taxable"] += t
                out["IGST"] += i_
                out["CGST"] += c_
                out["SGST"] += s_
                out["CESS"] += ce
                out["b2cs_taxable"] += t
                out["b2cs_IGST"] += i_
                out["b2cs_CGST"] += c_
                out["b2cs_SGST"] += s_
                out["b2cs_CESS"] += ce
        except mpu.PeriodParseError:
            out["month_marker_gaps"].append(B2CS)
    for sn in (CDNR, CDNUR):
        if sn in tabs:
            try:
                for r in _rows_for_month(ix, sn, month):
                    if not r["_any"]:
                        continue
                    out["cn_taxable"] += _f(r["taxable"])
                    out["cn_IGST"] += _f(r["igst"])
                    out["cn_CGST"] += _f(r["cgst"])
                    out["cn_SGST"] += _f(r["sgst"])
                    out["cn_CESS"] += _f(r["cess"])
            except mpu.PeriodParseError:
                out["month_marker_gaps"].append(sn)
    hsn_all = read_gstr1_hsn_all_months(path)
    if month in hsn_all:
        for r in hsn_all[month]:
            out["hsn_taxable"] += r["taxable"]
            out["hsn_IGST"] += r["igst"]
            out["hsn_CGST"] += r["cgst"]
            out["hsn_SGST"] += r["sgst"]
            out["hsn_CESS"] += r["cess"]
    else:
        out["month_marker_gaps"].append("hsn")
    if EXEMP in tabs and ix["ok"].get((EXEMP, "nil")) and ix["ok"].get((EXEMP, "nongst")):
        try:
            for r in _rows_for_month(ix, EXEMP, month):
                out["nil_taxable"] += _f(r["nil"])
                out["exempt_taxable"] += _f(r["exempted"])
                out["nongst_taxable"] += _f(r["nongst"])
            out["nil_exempt_taxable"] = out["nil_taxable"] + out["exempt_taxable"]
        except mpu.PeriodParseError:
            out["month_marker_gaps"].append("exemp")
    return out


# ---- HSN ------------------------------------------------------------------------------------------
def read_gstr1_hsn_all_months(path):
    """Legacy gst_parsers_returns.read_gstr1_hsn_all_months(path): Table 12 for every month, preferring the
    single 'hsn' tab per month and falling back to 'hsn(b2b)' + 'hsn(b2c)'."""
    ix = get_data(path)["_index"]
    cache = ix.setdefault("_hsn", None)
    if cache is not None:
        return cache

    def read_tab(tab, month):
        recs = []
        for r in _rows_for_month(ix, tab, month):
            if not r["_any"]:
                continue
            recs.append(dict(hsn=_s(r["hsn"]), desc=_s(r["desc"]), uqc=_s(r["uqc"]), qty=_f(r["qty"]),
                             rate=_f(r["rate"]), taxable=_f(r["taxable"]), igst=_f(r["igst"]), cgst=_f(r["cgst"]),
                             sgst=_f(r["sgst"]), cess=_f(r["cess"]), source_tab=tab))
        return recs

    all_months = set()
    for tab in ("hsn", "hsn(b2b)", "hsn(b2c)"):
        if tab in ix["tabs"]:
            all_months.update(ix["cov"][tab]["months"])
    out = {}
    for month in all_months:
        recs, used_old = [], False
        if "hsn" in ix["tabs"]:
            try:
                recs = read_tab("hsn", month)
                used_old = True
            except mpu.PeriodParseError:
                used_old = False
        if not used_old:
            for tab in ("hsn(b2b)", "hsn(b2c)"):
                if tab not in ix["tabs"]:
                    continue
                try:
                    recs.extend(read_tab(tab, month))
                except mpu.PeriodParseError:
                    continue
        out[month] = recs
    ix["_hsn"] = out
    return out


# ---- amendments and documents ---------------------------------------------------------------------
def parse_b2ba(path, month):
    ix = get_data(path)["_index"]
    if B2BA not in ix["tabs"] or not (ix["ok"].get((B2BA, "orig_invno")) and ix["ok"].get((B2BA, "rev_invno"))):
        return []
    out, last = [], None
    for r in _rows_for_month(ix, B2BA, month):
        if not r["_any"]:
            continue
        raw_orig, raw_gstin = r["orig_invno"], r["gstin"]
        if raw_orig not in (None, "") and raw_gstin not in (None, ""):
            last = (str(raw_gstin).strip(), _s(r["receiver"]), str(raw_orig).strip(), r["orig_invdate"],
                    _s(r["rev_invno"]), r["rev_invdate"], _f(r["invval"]), _s(r["pos"]))
        elif last is None:
            continue
        gstin_v, recipient_v, orig_invno_v, orig_date_v, revised_invno_v, revised_date_v, invval_v, pos_v = last
        out.append(dict(gstin=gstin_v, recipient=recipient_v, orig_invno=orig_invno_v, orig_date=orig_date_v,
                        revised_invno=revised_invno_v, revised_date=revised_date_v, invval=invval_v, pos=pos_v,
                        rate=_f(r["rate"]), taxable=_f(r["taxable"]), igst=_f(r["igst"]), cgst=_f(r["cgst"]),
                        sgst=_f(r["sgst"]), cess=_f(r["cess"])))
    return out


def parse_cdnra(path, month):
    ix = get_data(path)["_index"]
    if CDNRA not in ix["tabs"] or not (ix["ok"].get((CDNRA, "orig_noteno")) and ix["ok"].get((CDNRA, "rev_noteno"))):
        return []
    out, last = [], None
    for r in _rows_for_month(ix, CDNRA, month):
        if not r["_any"]:
            continue
        raw_orig, raw_gstin = r["orig_noteno"], r["gstin"]
        if raw_orig not in (None, "") and raw_gstin not in (None, ""):
            last = (str(raw_gstin).strip(), str(raw_orig).strip(), r["orig_notedate"], _s(r["rev_noteno"]),
                    r["rev_notedate"], _s(r["notetype"]))
        elif last is None:
            continue
        gstin_v, orig_noteno_v, orig_date_v, revised_noteno_v, revised_date_v, note_type_v = last
        out.append(dict(gstin=gstin_v, orig_noteno=orig_noteno_v, orig_date=orig_date_v, revised_noteno=revised_noteno_v,
                        revised_date=revised_date_v, note_type=note_type_v, taxable=_f(r["taxable"]),
                        igst=_f(r["igst"]), cgst=_f(r["cgst"]), sgst=_f(r["sgst"])))
    return out


def parse_docs(path, month):
    ix = get_data(path)["_index"]
    if DOCS not in ix["tabs"] or not (ix["ok"].get((DOCS, "sr_from")) and ix["ok"].get((DOCS, "sr_to"))):
        return []
    out = []
    for r in _rows_for_month(ix, DOCS, month):
        if not r["nature"]:
            continue
        out.append(dict(nature=_s(r["nature"]), sr_from=_s(r["sr_from"]), sr_to=_s(r["sr_to"]),
                        total=_f(r["total"]), cancelled=_f(r["cancelled"])))
    return out


# ---- line-level readers (gst_checks_monthly) ------------------------------------------------------
def read_gstr1_lines(path, month):
    from gst_checks_monthly import _parse_date
    ix = get_data(path)["_index"]
    out = []
    for sn, kind, id_f, date_f in ((B2B, "INV", "invno", "invdate"), (CDNR, "CN", "noteno", "notedate"),
                                   (CDNUR, "CN", "noteno", "notedate")):
        if sn not in ix["tabs"]:
            continue
        last = None
        for r in _rows_for_month(ix, sn, month):
            if not r["_any"]:
                continue
            raw_gstin, raw_invno = r.get("gstin"), r[id_f]
            if raw_gstin not in (None, "") and raw_invno not in (None, ""):
                last = (raw_gstin, raw_invno, r[date_f], r["pos"], r.get("rcm"))
                gstin_v, invno_v, invdate_v, pos_v, rcm_v = last
            elif last is not None:
                gstin_v, invno_v, invdate_v, pos_v, rcm_v = last
            else:
                gstin_v, invno_v, invdate_v, pos_v, rcm_v = "", "", None, "", ""
            out.append(dict(sheet=sn, kind=kind, gstin=str(gstin_v or "").strip(), invno=str(invno_v or "").strip(),
                            invdate=_parse_date(invdate_v), pos=str(pos_v or "").strip(), rate=_f(r["rate"]),
                            taxable=_f(r["taxable"]), igst=_f(r["igst"]), cgst=_f(r["cgst"]), sgst=_f(r["sgst"]),
                            cess=_f(r["cess"]), rcm=str(rcm_v or "").strip(), irn=str(r["irn"] or "").strip(),
                            irndate=_parse_date(r["irndate"])))
    return out


def read_gstr1_invoices(path, month):
    ix = get_data(path)["_index"]
    if B2B not in ix["tabs"]:
        raise KeyError(B2B)
    inv = {}
    last_no, last_pos, last_gstin = None, "", ""
    for r in _rows_for_month(ix, B2B, month):
        if not r["_any"]:
            continue
        raw_no, raw_gstin = r["invno"], r["gstin"]
        if raw_no not in (None, "") and raw_gstin not in (None, ""):
            last_no = str(raw_no).strip()
            last_pos = _s(r["pos"])
            last_gstin = str(raw_gstin).strip()
        no = last_no if last_no is not None else str(raw_no or "").strip()
        pos_v = last_pos if last_no is not None else _s(r["pos"])
        gstin_v = last_gstin if last_no is not None else str(raw_gstin or "").strip()
        d = inv.setdefault(no, dict(taxable=0.0, igst=0.0, cgst=0.0, sgst=0.0, invval=0.0, pos=pos_v, gstin=gstin_v,
                                    rates=set()))
        d["taxable"] += _f(r["taxable"])
        d["igst"] += _f(r["igst"])
        d["cgst"] += _f(r["cgst"])
        d["sgst"] += _f(r["sgst"])
        d["invval"] += _f(r["invval"])
        d["rates"].add(_f(r["rate"]))
    return inv


# ---- credit-note / B2CL readers (gst_checks_hsn_fraud) --------------------------------------------
def cdnr_rows_by_month(path):
    ix = get_data(path)["_index"]
    if CDNR not in ix["tabs"]:
        return {}
    info = ix["cov"][CDNR]
    if info["error"]:
        raise mpu.PeriodParseError(info["error"])
    out = {}
    for m in info["months"]:
        lst = []
        last_gstin = last_noteno = last_notedate = last_notetype = None
        for r in ix["rows"][CDNR]:
            if m not in r["_ml"] or not r["_any"]:
                continue
            raw_gstin, raw_noteno = r["gstin"], r["noteno"]
            if raw_gstin not in (None, "") and raw_noteno not in (None, ""):
                last_gstin = str(raw_gstin).strip()
                last_noteno = str(raw_noteno).strip()
                last_notedate = str(r["notedate"] or "").strip()
                last_notetype = str(r["notetype"] or "").strip()
                gstin_v, noteno_v, notedate_v, notetype_v = last_gstin, last_noteno, last_notedate, last_notetype
            elif last_gstin is not None:
                gstin_v, noteno_v, notedate_v, notetype_v = last_gstin, last_noteno, last_notedate, last_notetype
            else:
                gstin_v, noteno_v, notedate_v, notetype_v = "", "", "", ""
            lst.append(dict(gstin=gstin_v, noteno=noteno_v, notedate=notedate_v, notetype=notetype_v,
                            taxable=_f(r["taxable"]), igst=_f(r["igst"]), cgst=_f(r["cgst"]), sgst=_f(r["sgst"])))
        out[m] = lst
    return out


def b2cl_rows_by_month(path):
    ix = get_data(path)["_index"]
    if B2CL not in ix["tabs"]:
        return {}
    info = ix["cov"][B2CL]
    if info["error"]:
        raise mpu.PeriodParseError(info["error"])
    out = {}
    for m in info["months"]:
        lst = []
        last_invno = last_invdate = last_invval = last_pos = None
        for r in ix["rows"][B2CL]:
            if m not in r["_ml"] or not r["_any"]:
                continue
            raw_invno = r["invno"]
            if raw_invno not in (None, ""):
                last_invno = str(raw_invno).strip()
                last_invdate = str(r["invdate"] or "").strip()
                last_invval = _f(r["invval"])
                last_pos = str(r["pos"] or "").strip()
                invno_v, invdate_v, invval_v, pos_v = last_invno, last_invdate, last_invval, last_pos
            elif last_invno is not None:
                invno_v, invdate_v, invval_v, pos_v = last_invno, last_invdate, last_invval, last_pos
            else:
                invno_v, invdate_v, invval_v, pos_v = "", "", 0.0, ""
            lst.append(dict(invno=invno_v, invdate=invdate_v, invval=invval_v, pos=pos_v, rate=_f(r["rate"]),
                            taxable=_f(r["taxable"]), igst=_f(r["igst"])))
        out[m] = lst
    return out


# ---- ARN dates (gst_checks_forensic) --------------------------------------------------------------
def arn_dates_by_month(path):
    """Legacy gst_checks_forensic.gstr1_arn_dates_by_month(path) -> (out, warnings)."""
    d = get_data(path)
    ix = d["_index"]
    meta = ix["meta"]
    out, warnings = {}, []
    for r in d["returns"]:
        dt = r["arn_date"]
        if isinstance(dt, _dt.datetime):
            dt = dt.date()
        out[r["month"]] = {"arn": r["arn"] or None, "date": dt or None}
    if meta.get("marker_dates_found") != "Y":
        date_val = _parse_any_date(meta.get("readme_arn_date")) if meta.get("readme_arn_date") else None
        if date_val:
            warnings.append(
                "GSTR-1 per-month ARN date not found on any period-marker row for this file -- "
                "falling back to the single 'ARN date' value on the 'Read me' sheet "
                f"({date_val}). This is a WHOLE-FILE value; it is only applied to a specific "
                "month if 'Read me' also states a Tax Period, and is otherwise NOT applied to "
                "any month automatically (see '_readme_fallback' in the result) -- confirm "
                "against the portal before using it for a late-fee calculation on a specific month.")
            arn_val, tp = (meta.get("readme_arn") or None), meta.get("readme_tax_period")
            if tp:
                try:
                    _, _, labels = mpu.parse_marker_text(f"Financial Year: 0000-00 | Tax Period: {tp}")
                except Exception:  # noqa: BLE001
                    labels = []
                for lbl in labels:
                    out.setdefault(lbl, {"arn": arn_val, "date": date_val})
            else:
                out["_readme_fallback"] = {"arn": arn_val, "date": date_val}
    return out, warnings


# ---- flow-module readers --------------------------------------------------------------------------
def read_gstr1_b2b_ff(path, month):
    ix = get_data(path)["_index"]
    if B2B not in ix["tabs"]:
        return []
    out, carry = [], None
    for r in _rows_for_month(ix, B2B, month):
        if not r["_any"]:
            continue
        gstin = str(r["gstin"] or "").strip()
        if gstin:
            carry = dict(gstin=gstin, name=_s(r["receiver"]), invno=_s(r["invno"]), invdate=r["invdate"],
                         invval=_f(r["invval"]), pos=_s(r["pos"]), invtype=_s(r["invtype"]))
        if carry is None:
            continue
        out.append(dict(carry, month=month, rate=_f(r["rate"]), taxable=_f(r["taxable"]), igst=_f(r["igst"]),
                        cgst=_f(r["cgst"]), sgst=_f(r["sgst"]), cess=_f(r["cess"]), header_row=bool(gstin)))
    return out


def read_gstr1_b2c(path, month):
    res = dict(b2cs_taxable=0.0, b2cs_tax=0.0, b2cl_taxable=0.0, b2cl_tax=0.0, b2cl_invoices=0, available=True)
    ix = get_data(path)["_index"]
    for sheet, pfx in ((B2CS, "b2cs"), (B2CL, "b2cl")):
        if sheet not in ix["tabs"] or not ix["ok"].get((sheet, "taxable")):
            continue
        try:
            month_rows = _rows_for_month(ix, sheet, month)
        except mpu.PeriodParseError:
            continue
        for r in month_rows:
            if not r["_any"]:
                continue
            res[pfx + "_taxable"] += _f(r["taxable"])
            res[pfx + "_tax"] += _f(r["igst"]) + _f(r["cgst"]) + _f(r["sgst"])
            if pfx == "b2cl" and r["invno"]:
                res["b2cl_invoices"] += 1
    return res
