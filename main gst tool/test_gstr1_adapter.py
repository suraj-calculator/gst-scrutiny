#!/usr/bin/env python3
"""
Regression tests for the GSTR-1 canonical layer (gstr1_adapter.py, docs/GSTR1_CANONICAL_SPEC.md).

Self-contained: builds small synthetic merged GSTR-1 workbooks (multi-rate invoices with blank continuation
rows, credit notes, amendments, documents, both HSN tab styles, a quarterly banner) and checks that the
ORIGINAL readers and the canonical route return identical results for every reader and every month, that
broken inputs give the documented errors, and the one deliberate fix (the 'Non-GST Supplies' column).

Run:  python3 test_gstr1_adapter.py
"""
import os
import shutil
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.environ["GST_1_CANONICAL"] = "0"          # original readers for the legacy side

import gst_core as mpu                        # noqa: E402
import gst_parsers_returns as pr              # noqa: E402
import gst_checks_monthly as mo               # noqa: E402
import gst_checks_hsn_fraud as hf             # noqa: E402
import gst_checks_forensic as fo              # noqa: E402
import gst_checks_flow as fl                  # noqa: E402
import gstr1_adapter as ad                    # noqa: E402
import canonical_common as cc                 # noqa: E402

GSTIN = "05AAAAA0001A1Z1"
G1, G2 = "07BBBBB0002B1Z2", "09CCCCC0003C1Z3"
B2B_HDR = ["GSTIN/UIN of Recipient", "Receiver Name", "Invoice Number", "Invoice date", "Invoice Value",
           "Place Of Supply", "Reverse Charge", "Invoice Type", "Rate", "Taxable Value", "Integrated Tax",
           "Central Tax", "State/UT Tax", "Cess Amount", "Source", "IRN", "IRN date"]
CDNR_HDR = ["GSTIN/UIN of Recipient", "Receiver Name", "Note Number", "Note Date", "Note Type", "Place Of Supply",
            "Reverse Charge", "Note Supply Type", "Note Value", "Rate", "Taxable Value", "Integrated Tax",
            "Central Tax", "State/UT Tax", "Cess Amount", "Source", "IRN", "IRN date"]


def banner(fy, tp, arn="AA0504240000001", date="11-05-2024"):
    return f"Financial Year: {fy}  |  Tax Period: {tp}  |  ARN: {arn}  |  ARN Date: {date}  |  Date and Time of Generation: x"


def add_table(wb, name, title, header, blocks):
    """blocks = [(banner text, [data rows])]"""
    ws = wb.create_sheet(name)
    ws.append(["Goods and Services Tax - Form GSTR-1"])
    ws.append([])
    ws.append([title])
    ws.append(header)
    for b, rows in blocks:
        ws.append([b])
        for r in rows:
            ws.append(list(r))
    return ws


def make_fixture(path, hsn_style="hsn", quarterly=False, corrupt_b2b_header=False, tp_apr="April"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rm = wb.create_sheet("Read me")
    for row in (["Goods and Services Tax - Form GSTR-1"], [], [], ["Financial Year", None, "2024-25"],
                ["Tax Period", None, "April"], ["GSTIN", None, GSTIN], ["Legal Name", None, "TEST TRADERS"],
                ["ARN", None, "AA0504240000001"], ["ARN date", None, "11-05-2024"]):
        rm.append(row)
    A, M = banner("2024-25", tp_apr), banner("2024-25", "May", "AA0505240000002", "12-06-2024")
    hdr = list(B2B_HDR)
    if corrupt_b2b_header:
        hdr[hdr.index("Taxable Value")] = "Taxable Amt"
    add_table(wb, "b2b, sez, de_inv", "4A, 4B - Taxable supplies", hdr, [
        (A, [(G1, "BUYER ONE", "INV1", "05-04-2024", 1180.0, "Delhi", "N", "R", 18.0, 1000.0, 0, 90.0, 90.0, 0, None, "IRN1", "05-04-2024"),
             (G2, "BUYER TWO", "INV2", "06-04-2024", 1400.0, "Punjab", "N", "R", 12.0, 500.0, 0, 30.0, 30.0, 0, None, None, None),
             (None, None, None, None, None, None, None, None, 18.0, 300.0, 0, 27.0, 27.0, 0, None, None, None)]),
        (M, [(G1, "BUYER ONE", "INV3", "02-05-2024", 590.0, "Delhi", "N", "R", 18.0, 500.0, 90.0, 0, 0, 5.0, None, "IRN3", "02-05-2024")])])
    add_table(wb, "b2cl", "5A - Taxable outward inter-State supplies to unregistered persons",
              ["Invoice Number", "Invoice date", "Invoice Value", "Place Of Supply", "Rate", "Taxable Value",
               "Integrated Tax", "Cess Amount"],
              [(A, [("BIG1", "07-04-2024", 200000.0, "Delhi", 28.0, 150000.0, 42000.0, 1000.0),
                    (None, None, None, None, 18.0, 20000.0, 3600.0, 0)]), (M, [])])
    add_table(wb, "exp", "6A - Exports", ["Export Type", "Invoice Number", "Invoice date", "Invoice Value", "Port Code",
                                          "Shipping Bill Number", "Shipping Bill Date", "Taxable Value", "Rate",
                                          "Integrated Tax", "Cess Amount", "Source", "IRN", "IRN date"],
              [(A, [("WOPAY", "EX1", "09-04-2024", 5000.0, "INDEL4", "SB1", "10-04-2024", 5000.0, 0.0, 0, 0, None, None, None)]),
               (M, [])])
    add_table(wb, "b2cs", "7 - Taxable supplies (Net of debit notes and credit notes) to unregistered persons",
              ["Place Of Supply", "Rate", "Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax", "Cess Amount"],
              [(A, [("Uttarakhand", 18.0, 800.0, 0, 72.0, 72.0, 0)]), (M, [("Uttarakhand", 12.0, 100.0, 0, 6.0, 6.0, 0)])])
    add_table(wb, "exemp", "8 - Nil rated, exempted and non-GST outward supplies",
              ["Description", "Nil Rated Supplies", "Exempted (Other than Nil rated/non-GST supply)", "Non-GST Supplies"],
              [(A, [("Intra-State supplies to registered persons", 10.0, 20.0, 30.0)]),
               (M, [("Intra-State supplies to registered persons", 1.0, 2.0, 3.0)])])
    add_table(wb, "b2ba", "9A - Amendment to taxable outward supplies",
              ["GSTIN/UIN of Recipient", "Receiver Name", "Original Invoice Number", "Original Invoice date",
               "Revised Invoice Number", "Revised Invoice date", "Invoice Value", "Place Of Supply", "Reverse Charge",
               "Invoice Type", "Rate", "Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax", "Cess Amount"],
              [(A, []), (M, [(G1, "BUYER ONE", "OLD9", "01-01-2024", "NEW9", "01-01-2024", 300.0, "Delhi", "N", "R", 18.0, 250.0, 0, 22.5, 22.5, 0),
                             (None, None, None, None, None, None, None, None, None, None, 12.0, 50.0, 0, 3.0, 3.0, 0)])])
    add_table(wb, "cdnr", "9B - Credit/Debit notes", CDNR_HDR,
              [(A, [(G1, "BUYER ONE", "CN1", "20-04-2024", "C", "Delhi", "N", "R", 118.0, 18.0, 100.0, 0, 9.0, 9.0, 0, None, None, None),
                    (None, None, None, None, None, None, None, None, None, 12.0, 40.0, 0, 2.4, 2.4, 0, None, None, None)]), (M, [])])
    add_table(wb, "cdnur", "9B - Credit/Debit notes (unregistered)",
              ["UR Type", "Note Number", "Note Date", "Note Type", "Place Of Supply", "Note Value", "Rate", "Taxable Value",
               "Integrated Tax", "Cess Amount", "Source", "IRN", "IRN date"],
              [(A, [("B2CL", "UN1", "21-04-2024", "C", "Delhi", 500.0, 18.0, 400.0, 72.0, 0, None, None, None)]), (M, [])])
    add_table(wb, "cdnra", "9C - Amendment to credit/debit notes",
              ["GSTIN/UIN of Recipient", "Receiver Name", "Original Note Number", "Original Note Date", "Revised Note Number",
               "Revised Note Date", "Note Type", "Place Of Supply", "Reverse Charge", "Note Supply Type", "Note Value", "Rate",
               "Taxable Value", "Integrated Tax", "Central Tax", "State/UT Tax", "Cess Amount"],
              [(A, []), (M, [(G1, "BUYER ONE", "OLDCN", "01-02-2024", "NEWCN", "01-02-2024", "C", "Delhi", "N", "R", 100.0, 18.0, 80.0, 0, 7.2, 7.2, 0)])])
    hsn_hdr = ["HSN", "Description", "UQC", "Total Quantity", "Rate", "Taxable Value", "Integrated Tax", "Central Tax",
               "State/UT Tax", "Cess Amount"]
    if hsn_style == "hsn":
        add_table(wb, "hsn", "12 - HSN-wise summary", hsn_hdr,
                  [(A, [("8708", "Parts", "NOS", 10.0, 28.0, 1000.0, 0, 140.0, 140.0, 0)]),
                   (M, [("8708", "Parts", "NOS", 5.0, 28.0, 500.0, 140.0, 0, 0, 0)])])
    elif hsn_style == "split":
        add_table(wb, "hsn(b2b)", "12 - HSN (b2b)", hsn_hdr, [(A, [("8708", "Parts", "NOS", 10.0, 28.0, 900.0, 0, 126.0, 126.0, 0)]),
                                                              (M, [("8708", "Parts", "NOS", 5.0, 28.0, 400.0, 112.0, 0, 0, 0)])])
        add_table(wb, "hsn(b2c)", "12 - HSN (b2c)", hsn_hdr, [(A, [("8708", "Parts", "NOS", 1.0, 28.0, 100.0, 0, 14.0, 14.0, 0)]),
                                                              (M, [("8708", "Parts", "NOS", 1.0, 28.0, 100.0, 28.0, 0, 0, 0)])])
    else:   # mixed: 'hsn' carries a banner for April only, the split tabs for May
        add_table(wb, "hsn", "12 - HSN-wise summary", hsn_hdr, [(A, [("8708", "Parts", "NOS", 10.0, 28.0, 1000.0, 0, 140.0, 140.0, 0)])])
        add_table(wb, "hsn(b2b)", "12 - HSN (b2b)", hsn_hdr, [(M, [("8708", "Parts", "NOS", 5.0, 28.0, 400.0, 112.0, 0, 0, 0)])])
        add_table(wb, "hsn(b2c)", "12 - HSN (b2c)", hsn_hdr, [(M, [("8708", "Parts", "NOS", 1.0, 28.0, 100.0, 28.0, 0, 0, 0)])])
    add_table(wb, "docs", "13 - Documents issued", ["Nature of Document", "Sr. No. From", "Sr. No. To", "Total Number", "Cancelled"],
              [(A, [("Invoices for outward supply", "INV1", "INV2", 2.0, 0.0)]), (M, [("Invoices for outward supply", "INV3", "INV3", 1.0, 0.0)])])
    add_table(wb, "at", "11A - Advances", ["Place Of Supply", "Rate", "Gross Advance Received"], [(A, [])])   # a tab nobody reads
    wb.save(path)
    return path


def check(name, cond, detail=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail and not cond else ""))
    return cond


def run(fn, *a):
    try:
        return ("ok", fn(*a))
    except Exception as e:  # noqa: BLE001
        return ("EXC", type(e).__name__)


def compare_all(raw, months):
    """Every reader, every month: (legacy, canonical) pairs that differ."""
    diffs = []
    ad.clear_cache()
    pr._HSN_ALL_MONTHS_CACHE.clear()
    for m in months:
        for name, old, new in (
                ("parse_gstr1", pr.parse_gstr1, ad.parse_gstr1), ("parse_b2ba", pr.parse_b2ba, ad.parse_b2ba),
                ("parse_cdnra", pr.parse_cdnra, ad.parse_cdnra), ("parse_docs", pr.parse_docs, ad.parse_docs),
                ("read_gstr1_lines", mo.read_gstr1_lines, ad.read_gstr1_lines),
                ("read_gstr1_invoices", mo.read_gstr1_invoices, ad.read_gstr1_invoices),
                ("b2b_ff", fl.read_gstr1_b2b_ff, ad.read_gstr1_b2b_ff), ("b2c", fl.read_gstr1_b2c, ad.read_gstr1_b2c)):
            a, b = run(old, raw, m), run(new, raw, m)
            if name == "parse_gstr1" and a[0] == "ok" and b[0] == "ok":
                a[1].pop("nongst_taxable", None)      # the one deliberate difference, checked separately
                b[1].pop("nongst_taxable", None)
            if a != b:
                diffs.append((name, m, str(a)[:120], str(b)[:120]))
    for name, old, new in (("hsn_all", pr.read_gstr1_hsn_all_months, ad.read_gstr1_hsn_all_months),
                           ("cdnr_by_month", hf._cdnr_rows_by_month, ad.cdnr_rows_by_month),
                           ("b2cl_by_month", hf._b2cl_rows_by_month, ad.b2cl_rows_by_month),
                           ("arn_dates", fo.gstr1_arn_dates_by_month, ad.arn_dates_by_month)):
        a, b = run(old, raw), run(new, raw)
        if a != b:
            diffs.append((name, "-", str(a)[:120], str(b)[:120]))
    if mpu._gstr1_months_raw(raw) != ad.months(raw):
        diffs.append(("months", "-", "", ""))
    if mpu._read_me_gstin_and_name_raw(raw) != ad.gstin_and_name(raw):
        diffs.append(("gstin_and_name", "-", "", ""))
    return diffs


def main():
    ok = True
    tmp = tempfile.mkdtemp()
    cwd = os.getcwd()
    os.chdir(tmp)
    try:
        for style in ("hsn", "split", "mixed"):
            raw = make_fixture(os.path.join(tmp, f"g1_{style}.xlsx"), hsn_style=style)
            d = compare_all(raw, ["Apr-24", "May-24", "Jun-24"])
            ok &= check(f"[HSN tabs: {style}] every reader, every month identical (Jun-24 has no banner: same error)", not d, str(d[:2]))
        raw = os.path.join(tmp, "g1_hsn.xlsx")
        ad.clear_cache()

        # ---- the awkward cases actually exercised
        g = ad.parse_gstr1(raw, "Apr-24")
        ok &= check("multi-rate invoice INV2: continuation row counted, 2 invoices + 1 line of INV2 12%/18%",
                    g["b2b_count"] == 2 and g["taxable"] == 1000 + 500 + 300 + 150000 + 20000 + 5000 + 800 and
                    ("INV2", G2, 18.0) in g["lines"], str({k: g[k] for k in ("b2b_count", "taxable")}))
        ok &= check("credit notes CN1 (+ continuation) and cdnur summed", abs(g["cn_taxable"] - 540.0) < 1e-9, str(g["cn_taxable"]))
        lines = ad.read_gstr1_lines(raw, "Apr-24")
        ok &= check("continuation rows inherit their invoice identity in read_gstr1_lines",
                    [x["invno"] for x in lines if x["sheet"] == "b2b, sez, de_inv"] == ["INV1", "INV2", "INV2"])
        ok &= check("quarterly-style ARN dates come from the period banners",
                    ad.arn_dates_by_month(raw)[0]["Apr-24"]["arn"] == "AA0504240000001" and
                    str(ad.arn_dates_by_month(raw)[0]["May-24"]["date"]) == "2024-06-12")

        # ---- the deliberate fix: 'Non-GST Supplies' is its own column
        old_n, new_n = pr.parse_gstr1(raw, "Apr-24")["nongst_taxable"], ad.parse_gstr1(raw, "Apr-24")["nongst_taxable"]
        ok &= check("Non-GST column: canonical reads 'Non-GST Supplies' (30); the original read the Exempted column (20)",
                    new_n == 30.0 and old_n == 20.0, f"old={old_n} new={new_n}")

        # ---- quarterly banner (Apr-Jun) fans out to all three months, same as the original
        rawq = make_fixture(os.path.join(tmp, "g1_q.xlsx"), tp_apr="Apr-Jun")
        d = compare_all(rawq, ["Apr-24", "May-24", "Jun-24"])
        ok &= check("quarterly banner: every reader identical for all three months", not d, str(d[:2]))

        # ---- canonical workbook
        out = ad.write_canonical(ad.build_canonical(raw), os.path.join(tmp, "_canonical"))
        ok &= check("canonical file recognised", ad.is_canonical_file(out))
        dd = ad.read_canonical(out)
        ok &= check("RETURNS holds ARN + date per month",
                    {r["month"]: r["arn"] for r in dd["returns"]} == {"Apr-24": "AA0504240000001", "May-24": "AA0505240000002"})
        ok &= check("COVERAGE lists the months every tab has a banner for",
                    {(c["table"], c["month"]) for c in dd["coverage"] if c["table"] == "b2cl"} == {("b2cl", "Apr-24"), ("b2cl", "May-24")})
        ok &= check("unread tab recorded, not an error", "at" in dict(dd["meta"])["unread_tabs"] and dict(dd["meta"])["status"] == "OK")
        c = mpu.classify_folder(os.path.dirname(out))
        ok &= check("canonical GSTR-1 is classified as GSTR-1 with its months (needs a 3B too to be a run, so checked on the list)",
                    len(c["gstr1_files"]) == 1)
        # a file supplied as input is read through the adapter even with the switch off
        ad.clear_cache()
        ok &= check("canonical file as input gives the same numbers",
                    ad.parse_gstr1(out, "May-24")["taxable"] == pr.parse_gstr1(raw, "May-24")["taxable"])
        wb = openpyxl.load_workbook(out)
        ws = wb["B2B_LINES"]
        hdr = [x.value for x in ws[1]]
        ws.cell(row=2, column=hdr.index("taxable") + 1).value = 12345.0
        wb.save("mutant.xlsx")
        ad.clear_cache()
        ok &= check("mutation: editing the canonical file changes what the engine reads",
                    ad.parse_gstr1("mutant.xlsx", "Apr-24")["taxable"] != g["taxable"])

        # ---- classification: split HSN tabs are a GSTR-1 too
        os.makedirs("cls", exist_ok=True)
        shutil.copy(os.path.join(tmp, "g1_split.xlsx"), "cls/a.xlsx")
        ok &= check("classify_folder accepts a GSTR-1 whose HSN is only hsn(b2b)/hsn(b2c)", len(mpu.classify_folder("cls")["gstr1_files"]) == 1)

        # ---- documented errors
        rawb = make_fixture(os.path.join(tmp, "g1_bad.xlsx"), corrupt_b2b_header=True)
        ad.clear_cache()
        e = run(ad.parse_gstr1, rawb, "Apr-24")
        ok &= check("required b2b column renamed -> CanonicalLayoutError [E101] (loud, not a silent zero)",
                    e[0] == "EXC" and e[1] == "CanonicalLayoutError", str(e))
        ad.clear_cache()
        try:
            ad.parse_gstr1(rawb, "Apr-24")
        except cc.CanonicalLayoutError as ex:
            ok &= check("... and the message names the tab and the heading it looked for",
                        "[E101]" in str(ex) and "b2b, sez, de_inv" in str(ex) and "Taxable Value" in str(ex), str(ex)[:160])
        open("corrupt.xlsx", "wb").write(open(raw, "rb").read()[:2000])
        ad.clear_cache()
        ok &= check("E105 corrupt file", "[E105]" in (str(run(ad.build_canonical, "corrupt.xlsx")) + str(_exc_text(ad.build_canonical, "corrupt.xlsx"))))
        wb = openpyxl.Workbook()
        wb.active.title = "Summary"
        wb.save("wrong.xlsx")
        ad.clear_cache()
        ok &= check("E106 not a GSTR-1", "[E106]" in _exc_text(ad.build_canonical, "wrong.xlsx"))
        wb = openpyxl.load_workbook(raw)
        ws = wb["b2cl"]
        ws.insert_rows(5)                                     # a data row above the first banner
        ws["A5"] = "STRAY"
        wb.save("stray.xlsx")
        ad.clear_cache()
        e2 = run(ad.parse_gstr1, "stray.xlsx", "Apr-24")
        ok &= check("data before any banner on a tab: that tab is unusable for every month, like the original (gap noted)",
                    e2[0] == "ok" and "b2cl" in e2[1]["month_marker_gaps"], str(e2)[:120])
    finally:
        os.chdir(cwd)
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nALL PASSED" if ok else "\nSOME TESTS FAILED")
    return 0 if ok else 1


def _exc_text(fn, *a):
    try:
        fn(*a)
        return ""
    except Exception as e:  # noqa: BLE001
        return str(e)


if __name__ == "__main__":
    sys.exit(main())
