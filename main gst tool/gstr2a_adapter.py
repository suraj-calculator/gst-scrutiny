#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GSTR-2A CANONICAL LAYER  (spec: docs/GSTR2A_CANONICAL_SPEC.md)
================================================================
Only this module knows what a raw (merged) GSTR-2A workbook looks like.

    raw / merged GSTR-2A   --build_canonical()-->  canonical GSTR-2A data
    canonical data         --write_canonical()-->  Canonical_GSTR2A_<GSTIN>_<FY>.xlsx
    canonical .xlsx        --read_canonical()  -->  canonical data
    canonical data         --parse_r2a_excel()  -->  exactly what gst_parsers_dept.parse_r2a_excel() returns

Every 2A sheet (B2B, CDNR, B2BA, CDNRA, ISD) has a title, a wrapped 1-3 row heading block, a period banner per
month and that month's rows under it. Each canonical table keeps the source rows with canonical column names and
the month(s) their banner covers; numbers are converted (blank = 0), identifiers, text and dates are kept exactly
as in the source. The reader then applies the same business logic as the original (drop repeated sub-heading and
stray rows by GSTIN shape, keep only the '-Total' rollup row of every document, ...) over those rows.

The original took every column by FIXED POSITION. The converter finds each column by its HEADING first (and reports
what it did in the MAPPING_REPORT), falling back to the original position when the heading does not match - so an
unchanged layout gives the identical result, a moved column is followed, and a heading that no longer matches is
reported in plain words instead of silently reading whatever now sits in that column.
"""

import datetime as _dt
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

import canonical_common as cc
import gst_core as mpu

SCHEMA_VERSION = "2A-1.0"
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "gstr2a_mapping.json")

LEAD_COLS = ["months_covered", "src_row"]
RETURN_COLS = ["month", "fy", "tax_period", "marker_text"]
COVERAGE_COLS = ["table", "month", "rows", "note"]
MAP_COLS = ["canonical_sheet", "canonical_field", "raw_sheet", "raw_header_found", "raw_column", "matched_by",
            "confidence", "status", "rows_filled", "rows_blank", "note"]
ISSUE_COLS = cc.ISSUE_COLS
ORDER = ["b2b", "cdnr", "b2ba", "cdnra", "isd"]          # the order the original reader walks the sheets
_GSTIN_RE = re.compile(r"^\d{2}[A-Z]{5}\d{4}[A-Z]\dZ[A-Z\d]$")

_CONTEXT = {"gstin": None}


def set_context(gstin=None):
    if gstin:
        _CONTEXT["gstin"] = gstin


def _mapping():
    return cc.load_json(_MAPPING_FILE)


def _tables():
    return _mapping()["tables"]


def _canon_cols(tname):
    return LEAD_COLS + list(_tables()[tname]["fields"]) + ["extra"]


def _gstin_ok(v):
    return bool(v) and bool(_GSTIN_RE.match(str(v).strip().upper()))


def _blank(c):
    return c is None or (isinstance(c, str) and c.strip() == "")


# ----------------------------------------------------------------------------------------
# heading detection
# ----------------------------------------------------------------------------------------
def _find_header_row(rows, must_contain, max_scan=15):
    """0-based index of the first of the top rows with a cell containing `must_contain` (same test as the original
    gst_parsers_dept._find_header_row, which is 1-based)."""
    want = must_contain.lower()
    for i, row in enumerate(rows[:max_scan]):
        if any(want in (str(c).strip() if c else "").lower() for c in row):
            return i
    return None


def _column_labels(rows, hdr_idx, start, gstin_pos):
    """Combined heading text per column: the group row directly above the heading row (its group heading applied to
    the columns it spans), the heading row(s) down to the first banner, and - for the amendment sheets, which repeat a
    sub-heading row right after the first banner - that row. Returns (cleaned labels, display labels)."""
    width = max((len(r) for r in rows[:start + 6]), default=0)
    parts = [[] for _ in range(width)]
    if hdr_idx >= 1:
        above = rows[hdr_idx - 1]
        if sum(1 for c in above if not _blank(c)) >= 2:          # a real group-heading row (e.g. Original / Revised)
            cur = ""
            for c in range(width):
                v = above[c] if c < len(above) else None
                if not _blank(v):
                    cur = str(v).strip()
                if cur:
                    parts[c].append(cur)
    for i in range(hdr_idx, start):
        for c, v in enumerate(rows[i]):
            if not _blank(v):
                parts[c].append(str(v).strip())
    for i in range(start + 1, min(start + 4, len(rows))):
        r = rows[i]
        if mpu.is_marker_row(r) or not any(not _blank(c) for c in r):
            continue
        texts = [c for c in r if not _blank(c)]
        if len(texts) >= 3 and all(isinstance(c, str) for c in texts) and not any(_gstin_ok(c) for c in texts):
            for c, v in enumerate(r):
                if not _blank(v):
                    parts[c].append(str(v).strip())
        break
    return [cc.clean(" ".join(p)) for p in parts], [" / ".join(p) for p in parts]


def _matches(label, spec):
    if any(tok in label for tok in spec.get("none", [])):
        return False
    return any(all(tok in label for tok in alt) for alt in spec["match"])


def _resolve_columns(tdef, labels, display, ctx, sheet):
    """field -> (column or None, matched_by, confidence, note). See gstr2a_mapping.json's _about."""
    fields = tdef["fields"]
    width = len(labels)
    cols, how, used = {}, {}, set()
    for f, spec in fields.items():                                          # 1: the original position, if its heading fits
        p = spec["pos"]
        if p < width and _matches(labels[p], spec):
            cols[f], how[f], _ = p, ("position", "HIGH", ""), None
            used.add(p)
    for f, spec in fields.items():
        if f in cols:
            continue
        p = spec["pos"]
        cands = [c for c in range(width) if c not in used and _matches(labels[c], spec)]
        if len(cands) == 1:                                                  # 2: the heading moved
            c = cands[0]
            cols[f], how[f] = c, ("heading", "MEDIUM",
                                  f"expected at column {get_column_letter(p + 1)}, found by heading at "
                                  f"{get_column_letter(c + 1)}")
            used.add(c)
            ctx.issue("W703", "WARNING", sheet, f, "all",
                      f"GSTR-2A sheet '{sheet}': '{f}' is not at its usual column {get_column_letter(p + 1)}; found by "
                      f"heading ('{display[c]}') at column {get_column_letter(c + 1)} and read from there.")
        elif p < width:                                                       # 3: read by position as the original did
            cols[f], how[f] = p, ("position", "LOW",
                                  f"heading at this column reads '{display[p]}', expected {spec['match'][0]}")
            used.add(p)
            ctx.issue("W704", "WARNING", sheet, f, "all",
                      f"GSTR-2A sheet '{sheet}': cannot confirm the column for '{f}' from its heading (column "
                      f"{get_column_letter(p + 1)} reads '{display[p]}'); read by position as before - check that the "
                      f"figures are what they should be.")
        else:
            cols[f], how[f] = None, ("", "", "column does not exist in this sheet")
    return cols, how


# ----------------------------------------------------------------------------------------
# reading the raw workbook
# ----------------------------------------------------------------------------------------
def _scan_readme(wb):
    name = next((n for n in wb.sheetnames if n.lower() == "read me"), None)
    out = {}
    if not name:
        return out
    labels = _mapping()["readme_labels"]
    for row in wb[name].iter_rows(min_row=1, max_row=12, values_only=True):
        cells = [str(c).strip() for c in row if c not in (None, "")]
        if len(cells) < 2:
            continue
        for i, c in enumerate(cells[:-1]):
            key = c.lower()
            for field, names in labels.items():
                if key in names and field not in out:
                    out[field] = cells[i + 1]
    return out


def _build_table(ctx, tname, tdef, sheet, rows, table_rows, coverage, mapping, returns):
    hdr_idx = _find_header_row(rows, tdef["header_text"])
    if hdr_idx is None:
        ctx.issue("W702", "WARNING", sheet, "", "all",
                  f"GSTR-2A sheet '{sheet}' is present but the heading '{tdef['header_text']}' was not found in its "
                  f"top rows, so the sheet is skipped (as the original did, silently).",
                  f"the {sheet} figures in every GSTR-2A check", "Check the sheet is the standard GSTR-2A layout.")
        return False
    start = next((i for i in range(hdr_idx, len(rows)) if rows[i] and mpu.is_marker_row(rows[i])), None)
    if start is None:
        ctx.issue("W701", "WARNING", sheet, "", "all",
                  f"GSTR-2A sheet '{sheet}' has no 'Financial Year | Tax Period' banner rows, so no month can be "
                  f"read from it.", f"the {sheet} figures in every GSTR-2A check", "")
        return True
    labels, display = _column_labels(rows, hdr_idx, start, tdef["fields"][tdef["gstin_field"]]["pos"])
    cols, how = _resolve_columns(tdef, labels, display, ctx, sheet)
    missing = [f for f, spec in tdef["fields"].items() if spec.get("required") and cols.get(f) is None]
    canon = tdef["canon_sheet"]
    fill = {f: [0, 0] for f in tdef["fields"]}
    mapped = {c for c in cols.values() if c is not None}
    for f, spec in tdef["fields"].items():
        c = cols[f]
        mb, conf, note = how[f]
        mapping.append(dict(canonical_sheet=canon, canonical_field=f, raw_sheet=sheet,
                            raw_header_found=display[c] if c is not None else "",
                            raw_column=get_column_letter(c + 1) if c is not None else "", matched_by=mb,
                            confidence=conf, status="OK" if c is not None else "MISSING", rows_filled=0,
                            rows_blank=0, note=note))
    if missing:
        msg = (f"GSTR-2A sheet '{sheet}': required column(s) {missing} could not be located (headings found: "
               f"{sorted(set(d for d in display if d))[:40]}). This sheet is unusable, so it reads as empty in every "
               f"GSTR-2A check.")
        ctx.issue("E701", "ERROR", sheet, ", ".join(missing), "all", msg, f"the {sheet} figures in every GSTR-2A check",
                  "Add the new heading to gstr2a_mapping.json or supply the file as downloaded from the portal.")
        coverage.append(dict(table=sheet, month="*", rows=-1, note="[E701] " + msg))
        return True
    n_by_month, order, current = {}, [], None
    for ridx in range(start, len(rows)):
        r = rows[ridx]
        if r and mpu.is_marker_row(r):
            try:
                fy, tp, labels_m = mpu.parse_marker_text(str(r[0]))
            except mpu.PeriodParseError as e:
                ctx.fail("E703", f"GSTR-2A sheet '{sheet}': the period banner in row {ridx + 1} is not understood ({e}).",
                         sheet=sheet, skipped="every GSTR-2A check", action="Check the file.")
            current = labels_m
            for lbl in labels_m:
                returns.setdefault(lbl, dict(month=lbl, fy=fy, tax_period=tp, marker_text=str(r[0])))
                if lbl not in n_by_month:
                    n_by_month[lbl] = 0
                    order.append(lbl)
            continue
        if not any(not _blank(c) for c in r):
            continue
        rec = dict(months_covered=";".join(current), src_row=ridx + 1)
        for f, spec in tdef["fields"].items():
            c = cols[f]
            v = r[c] if (c is not None and c < len(r)) else None
            fill[f][0 if not _blank(v) else 1] += 1
            rec[f] = cc.num(v) if spec["type"] == "num" else v
        rec["extra"] = "; ".join(f"{display[i] or get_column_letter(i + 1)}={r[i]}" for i in range(min(len(r), len(display)))
                                 if i not in mapped and not _blank(r[i]))
        table_rows[canon].append(rec)
        for lbl in current:
            n_by_month[lbl] += 1
    for lbl in order:
        coverage.append(dict(table=sheet, month=lbl, rows=n_by_month[lbl], note=""))
    for m in mapping:
        if m["canonical_sheet"] == canon and m["canonical_field"] in fill:
            m["rows_filled"], m["rows_blank"] = fill[m["canonical_field"]]
    return True


def build_canonical(raw_paths):
    """Convert one raw/merged GSTR-2A workbook into canonical data (a dict of tables). A file that cannot be opened
    raises PeriodParseError ('[E70x] ...')."""
    if isinstance(raw_paths, str):
        raw_paths = [raw_paths]
    path = raw_paths[0]
    ctx = cc.IssueLog("GSTR-2A", [os.path.basename(p) for p in raw_paths])
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        ctx.fail("E702", f"GSTR-2A \"{os.path.basename(path)}\": cannot open the file ({type(e).__name__}: {e}). It may be "
                 f"corrupt, password-protected or not an Excel file.", skipped="every GSTR-2A check",
                 action="Re-download the file.")
    tables = _tables()
    returns, coverage, mapping = {}, [], []
    table_rows = {t["canon_sheet"]: [] for t in tables.values()}
    tabs_present, absent = [], []
    try:
        names = list(wb.sheetnames)
        readme = _scan_readme(wb)
        for tname in ORDER:
            tdef = tables[tname]
            sheet = next((n for n in tdef["raw_names"] if n in names), None)
            if sheet is None:
                absent.append(tdef["raw_names"][0])
                continue
            rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
            if _build_table(ctx, tname, tdef, sheet, rows, table_rows, coverage, mapping, returns):
                tabs_present.append(f"{tname}={sheet}")
    finally:
        wb.close()
    if absent:
        ctx.issue("I702", "INFO", "", "", "all", f"Sheets not in the workbook (read as empty): {absent}.")
    if not tabs_present:
        ctx.issue("E704", "ERROR", "", "", "all",
                  "Workbook found but none of B2B/CDNR/B2BA/CDNRA/ISD sheets were readable.",
                  "every GSTR-2A check", "Check that this is the merged GSTR-2A workbook.")
    fys = sorted({r["fy"] for r in returns.values()}) or ([readme["fy"]] if readme.get("fy") else [])
    gstin = _CONTEXT["gstin"] or readme.get("gstin") or ""
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    meta = [("schema_version", SCHEMA_VERSION), ("recipient_gstin", gstin), ("legal_name", readme.get("legal_name", "")),
            ("fy", ", ".join(fys)), ("source_files", "; ".join(ctx.source_files)),
            ("source_layout", "merged with period banners"), ("adapter_version", ADAPTER_VERSION),
            ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("months_present", ";".join(sorted(returns))), ("tabs_present", ";".join(tabs_present)),
            ("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK"))]
    return dict(meta=meta, returns=list(returns.values()), coverage=coverage, tables=table_rows, mapping=mapping,
                issues=ctx.issues)


# ----------------------------------------------------------------------------------------
# write / read the canonical workbook
# ----------------------------------------------------------------------------------------
def canonical_filename(data):
    return cc.canonical_filename("GSTR2A", data["meta"])


def write_canonical(data, out_dir):
    sheets = [("RETURNS", RETURN_COLS, data["returns"]), ("COVERAGE", COVERAGE_COLS, data["coverage"])]
    for tname in ORDER:
        tdef = _tables()[tname]
        sheets.append((tdef["canon_sheet"], _canon_cols(tname), data["tables"][tdef["canon_sheet"]]))
    sheets += [("MAPPING_REPORT", MAP_COLS, data["mapping"]), ("ISSUES", ISSUE_COLS, data["issues"])]
    return cc.write_workbook(out_dir, canonical_filename(data), data["meta"], sheets, max_col_width=30)


def is_canonical_file(path):
    return cc.is_canonical(path, "2A-", ("META", "RETURNS", "COVERAGE"))


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        tables = {t["canon_sheet"]: cc.read_table(wb, t["canon_sheet"], _canon_cols(n)) for n, t in _tables().items()}
        return dict(meta=cc.read_meta(wb), returns=cc.read_table(wb, "RETURNS", RETURN_COLS),
                    coverage=cc.read_table(wb, "COVERAGE", COVERAGE_COLS), tables=tables,
                    mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS),
                    issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------
# engine-facing access
# ----------------------------------------------------------------------------------------
def _index(data):
    """Per table: found?, its sheet name, error note, and {month: [rows]} in banner order."""
    meta = dict(data["meta"])
    found = dict(p.split("=", 1) for p in str(meta.get("tabs_present") or "").split(";") if p)
    ix = {}
    for tname in ORDER:
        tdef = _tables()[tname]
        sheet = found.get(tname)
        info = dict(found=sheet is not None, sheet=sheet, error=None, months=[], by_month={})
        if sheet is not None:
            for c in data["coverage"]:
                if c["table"] != sheet:
                    continue
                if c["month"] == "*":
                    info["error"] = c["note"]
                elif c["month"] not in info["by_month"]:
                    info["months"].append(c["month"])
                    info["by_month"][c["month"]] = []
            for r in data["tables"].get(tdef["canon_sheet"], []):
                for m in str(r["months_covered"]).split(";"):
                    if m in info["by_month"]:
                        info["by_month"][m].append(r)
        ix[tname] = info
    return dict(meta=meta, tables=ix)


_SRC = cc.CanonicalCache("GSTR-2A", is_canonical_file, read_canonical, build_canonical, write_canonical, _index)


def get_data(path, out_dir=None):
    return _SRC.get(path, out_dir)


def clear_cache():
    _SRC.clear()


def _s(v):
    return str(v or "").strip()


def _clean_gstin(v, malformed):
    """Legacy gst_parsers_dept._r2a_clean_gstin."""
    s = str(v or "").strip().upper()
    if s and not _GSTIN_RE.match(s):
        malformed.append(s)
    return s


def _valid_and_totals(rows, doc_field):
    """The original's two row filters: genuine data rows only (GSTIN-shaped), then - where there is a document column -
    only each document's '-Total' rollup row (suffix stripped), reporting any document with no '-Total' row.
    Returns (rows, missing_total_docnos)."""
    valid = [r for r in rows if _gstin_ok(r["gstin"])]
    if doc_field is None:
        return valid, []
    total_rows, rateline, total_docs = [], set(), set()
    for r in valid:
        doc = r[doc_field]
        if not doc or not isinstance(doc, str):
            continue
        doc = doc.strip()
        if not doc:
            continue
        if doc.endswith("-Total"):
            base = doc[:-len("-Total")].strip()
            total_docs.add(base)
            total_rows.append(dict(r, **{doc_field: base}))
        else:
            rateline.add(doc)
    return total_rows, sorted(rateline - total_docs)


def parse_r2a_excel(path):
    """Legacy gst_parsers_dept.parse_r2a_excel(path): the whole merged GSTR-2A workbook in one pass."""
    from gst_parsers_dept import r2a_clean_date, r2a_invtype_canon
    out = dict(available=False, reason=None, months_present=set(), b2b={}, b2ba={}, cdnr={}, cdnra={}, isd={},
               total_row_missing={}, malformed_gstin=[])
    if not path or not os.path.exists(path):
        out["reason"] = "GSTR-2A not supplied for this taxpayer/FY."
        return out
    try:
        ix = get_data(path)["_index"]["tables"]
    except mpu.PeriodParseError as ex:
        out["reason"] = f"Could not read GSTR-2A workbook: {ex}"
        return out
    malformed = out["malformed_gstin"]
    any_found = False
    D, S = r2a_clean_date, _s

    def tail_b2b(r):
        return dict(g1_status=S(r["g1_status"]), g1_filing_date=S(r["g1_filing_date"]),
                    g1_filing_period=S(r["g1_filing_period"]), g3b_status=S(r["g3b_status"]))

    def rec_b2b(m, r):
        return dict(month=m, gstin=_clean_gstin(r["gstin"], malformed), supplier=S(r["supplier"]), invno=S(r["invno"]),
                    invtype=r2a_invtype_canon(r["invtype"]), invdate=D(r["invdate"]), invval=r["invval"] or 0.0,
                    pos=S(r["pos"]), rcm=S(r["rcm"]).upper(), taxable=r["taxable"] or 0.0, igst=r["igst"] or 0.0,
                    cgst=r["cgst"] or 0.0, sgst=r["sgst"] or 0.0, cess=r["cess"] or 0.0, **tail_b2b(r),
                    amend_type=S(r["amend_type"]), amend_period=S(r["amend_period"]), cancel_date=S(r["cancel_date"]),
                    source=S(r["source"]), irn=S(r["irn"]), irn_date=S(r["irn_date"]))

    def rec_cdnr(m, r):
        return dict(month=m, gstin=_clean_gstin(r["gstin"], malformed), supplier=S(r["supplier"]),
                    note_type=S(r["note_type"]), note_no=S(r["note_no"]), supply_type=S(r["supply_type"]),
                    note_date=D(r["note_date"]), note_val=r["note_val"] or 0.0, pos=S(r["pos"]), rcm=S(r["rcm"]).upper(),
                    taxable=r["taxable"] or 0.0, igst=r["igst"] or 0.0, cgst=r["cgst"] or 0.0, sgst=r["sgst"] or 0.0,
                    cess=r["cess"] or 0.0, **tail_b2b(r), amend_type=S(r["amend_type"]),
                    amend_period=S(r["amend_period"]), cancel_date=S(r["cancel_date"]), source=S(r["source"]),
                    irn=S(r["irn"]), irn_date=S(r["irn_date"]))

    def rec_b2ba(m, r):
        return dict(month=m, orig_invno=S(r["orig_invno"]), orig_invdate=D(r["orig_invdate"]),
                    gstin=_clean_gstin(r["gstin"], malformed), supplier=S(r["supplier"]),
                    invtype=r2a_invtype_canon(r["invtype"]), invno=S(r["invno"]), invdate=D(r["invdate"]),
                    invval=r["invval"] or 0.0, pos=S(r["pos"]), rcm=S(r["rcm"]).upper(), taxable=r["taxable"] or 0.0,
                    igst=r["igst"] or 0.0, cgst=r["cgst"] or 0.0, sgst=r["sgst"] or 0.0, cess=r["cess"] or 0.0,
                    **tail_b2b(r), cancel_date=S(r["cancel_date"]), amend_type=S(r["amend_type"]),
                    orig_tax_period=S(r["orig_tax_period"]))

    def rec_cdnra(m, r):
        return dict(month=m, orig_note_type=S(r["orig_note_type"]), orig_note_no=S(r["orig_note_no"]),
                    orig_note_date=D(r["orig_note_date"]), gstin=_clean_gstin(r["gstin"], malformed),
                    supplier=S(r["supplier"]), note_type=S(r["note_type"]), note_no=S(r["note_no"]),
                    supply_type=S(r["supply_type"]), note_date=D(r["note_date"]), note_val=r["note_val"] or 0.0,
                    pos=S(r["pos"]), rcm=S(r["rcm"]).upper(), taxable=r["taxable"] or 0.0, igst=r["igst"] or 0.0,
                    cgst=r["cgst"] or 0.0, sgst=r["sgst"] or 0.0, cess=r["cess"] or 0.0, **tail_b2b(r),
                    amend_type=S(r["amend_type"]), orig_tax_period=S(r["orig_tax_period"]),
                    cancel_date=S(r["cancel_date"]))

    def rec_isd(m, r):
        return dict(month=m, eligibility=S(r["eligibility"]).upper(), gstin=_clean_gstin(r["gstin"], malformed),
                    isd_name=S(r["isd_name"]), doc_type=S(r["doc_type"]), isd_invno=S(r["isd_invno"]),
                    isd_invdate=D(r["isd_invdate"]), isd_cnno=S(r["isd_cnno"]), isd_cndate=D(r["isd_cndate"]),
                    orig_invno=S(r["orig_invno"]), orig_invdate=D(r["orig_invdate"]), igst=r["igst"] or 0.0,
                    cgst=r["cgst"] or 0.0, sgst=r["sgst"] or 0.0, cess=r["cess"] or 0.0,
                    filing_status=S(r["filing_status"]), amend=S(r["amend"]), amend_period=S(r["amend_period"]))

    builders = {"b2b": rec_b2b, "cdnr": rec_cdnr, "b2ba": rec_b2ba, "cdnra": rec_cdnra, "isd": rec_isd}
    for tname in ORDER:
        info = ix[tname]
        if not info["found"]:
            continue
        any_found = True
        if info["error"]:                    # required column missing: the sheet reads as empty (reported as E701)
            continue
        doc_field = _tables()[tname]["doc_field"]
        for month in info["months"]:
            out["months_present"].add(month)
            clean, missing = _valid_and_totals(info["by_month"][month], doc_field)
            if missing:
                out["total_row_missing"].setdefault(tname.upper(), []).extend(f"{month}: {d}" for d in missing)
            out[tname][month] = [builders[tname](month, r) for r in clean]
    if not any_found:
        out["reason"] = "Workbook found but none of B2B/CDNR/B2BA/CDNRA/ISD sheets were readable."
        return out
    out["available"] = True
    return out
