#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
LEDGER CANONICAL LAYER  (spec: docs/LEDGER_CANONICAL_SPEC.md)
==============================================================
Only this module knows what the four portal ledger CSV downloads look like:

    Electronic Cash Ledger              kind "cash"
    Electronic Credit Ledger            kind "credit"
    Electronic Liability Register       kind "liability"         (Part I, return related)
    Electronic Liability Ledger         kind "liability_demand"  (Part II, DRC / demand / voluntary payments)

    raw CSV  --build_canonical()-->  canonical ledger data
    data     --write_canonical()-->  Canonical_<CASHLEDGER|CREDITLEDGER|LIABREGISTER|LIABLEDGER>_<GSTIN>_<FY>.xlsx
    .xlsx    --read_canonical()  -->  data
    data     --parse_cash_or_liability_ledger() / parse_credit_ledger()  -->  exactly what gst_parsers_dept returns today

Each CSV has a title, GSTIN / From / To lines, a two-row heading (group heading + sub-heading), one row per transaction, an
'Opening Balance' row and a 'Closing Balance' row. The canonical file keeps every row under the heading with canonical column
names: text exactly as in the source, every amount converted to a number (blank / '-' = 0; the portal's Indian '3,06,53,037.00'
grouping is fine) - all six parts (tax, interest, penalty, fee, others, total) of every head in both the Debited/Credited block
and the Balance block, although the engine reads only some of them. The reader keeps the original's business logic (which rows
are transactions, credit vs debit sign, the monthly roll-ups) over those rows.

The original took every column by FIXED POSITION. The converter also reads the headings and reports, in plain words, a column
that has moved (W902) or whose heading can no longer be confirmed (W901); a file it cannot use at all raises PeriodParseError, and
the readers turn that into the same 'not supplied' degrade the engine already has for a missing ledger, with the message printed.
"""

import csv
import datetime as _dt
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

import canonical_common as cc
import gst_core as mpu

SCHEMA_VERSION = "LED-1.0"
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "ledger_mapping.json")

ROWS_SHEET = "LEDGER_ROWS"
LEAD_COLS = ["src_row", "src_width"]
MAP_COLS = ["canonical_field", "raw_header_found", "raw_column", "matched_by", "confidence", "status", "rows_filled",
            "rows_blank", "note"]
ISSUE_COLS = cc.ISSUE_COLS
KINDS = ("cash", "credit", "liability", "liability_demand")
_MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
HEADS = ["IGST", "CGST", "SGST", "CESS"]


def _kinds():
    return cc.load_json(_MAPPING_FILE)["kinds"]


def _canon_cols(kind):
    return LEAD_COLS + list(_kinds()[kind]["fields"]) + ["extra"]


# ----------------------------------------------------------------------------------------
# reading the raw CSV
# ----------------------------------------------------------------------------------------
def _read_csv(path, ctx):
    try:
        with open(path, newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f))
    except UnicodeDecodeError:
        with open(path, newline="", encoding="cp1252", errors="replace") as f:
            rows = list(csv.reader(f))
        ctx.issue("I901", "INFO", "", "", "all", "The file is not UTF-8; it was read as Windows-1252 (Excel's 'CSV' export).")
        return rows


def _fail(ctx, path, kind, iid, msg, action="Download the ledger again from the GST portal (CSV) and supply it unedited."):
    ctx.fail(iid, f"{_LABEL[kind]} \"{os.path.basename(path)}\": {msg}", skipped=f"every check that reads the {_LABEL[kind]}",
             action=action)


_LABEL = {"cash": "Electronic Cash Ledger", "credit": "Electronic Credit Ledger",
          "liability": "Electronic Liability Register", "liability_demand": "Electronic Liability Ledger"}


def _fy_from(from_text, taxperiod_text):
    m = re.match(r"\s*(\d{2})[-/](\d{2})[-/](\d{4})", from_text or "")
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y0 = y if mo >= 4 else y - 1
        return f"{y0}-{str(y0 + 1)[2:]}"
    m = re.match(r"\s*([A-Za-z]{3})[A-Za-z]*[-\s]*(\d{2,4})", taxperiod_text or "")
    if m and m.group(1).title() in _MON:
        mo, y = _MON.index(m.group(1).title()) + 1, int(m.group(2))
        y = y + 2000 if y < 100 else y
        y0 = y if mo >= 4 else y - 1
        return f"{y0}-{str(y0 + 1)[2:]}"
    return ""


def _labels(rows, hdr_idx, width):
    """group heading (forward-filled across its span) and sub-heading per column, cleaned; plus the sub-row length."""
    hdr = rows[hdr_idx]
    sub_row = rows[hdr_idx + 1] if hdr_idx + 1 < len(rows) else []
    group, cur = [], ""
    for c in range(width):
        v = hdr[c] if c < len(hdr) else ""
        if v.strip():
            cur = cc.clean(v)
        group.append(cur)
    sub = [cc.clean(sub_row[c]) if c < len(sub_row) else "" for c in range(width)]
    display = [((hdr[c].strip() if c < len(hdr) else "") + " / " + (sub_row[c].strip() if c < len(sub_row) else "")).strip(" /")
               for c in range(width)]
    return group, sub, len(sub_row), display


def _candidates(spec, group, sub, sub_len):
    """Columns whose heading fits `spec` (see ledger_mapping.json _about)."""
    width = len(group)
    if "label" in spec:
        return [c for c in range(width) if all(t in (group[c] + sub[c]) for t in spec["label"])]
    out = []
    for c in range(width):
        if "group" in spec and not all(t in group[c] for t in spec["group"]):
            continue
        if sub[c] == spec["sub"] or (sub[c] == "" and c >= sub_len and "group" in spec):
            out.append(c)
    if "nth" in spec:
        out = [c for c in out if sub[c] == spec["sub"]]
        return [out[spec["nth"] - 1]] if len(out) >= spec["nth"] else []
    return out


def _resolve(kdef, group, sub, sub_len, display, ctx, kind):
    fields, width = kdef["fields"], len(group)
    cols, how, used = {}, {}, set()
    cands = {f: _candidates(spec, group, sub, sub_len) for f, spec in fields.items()}
    for f, spec in fields.items():                                        # 1: the original position, if its heading fits
        p = spec["pos"]
        if p in cands[f] and p not in used:
            weak = "sub" in spec and sub[p] == "" and p >= sub_len
            cols[f] = p
            how[f] = ("position", "MEDIUM" if weak else "HIGH",
                      "sub-heading not present in this file; group heading confirmed" if weak else "")
            used.add(p)
    for f, spec in fields.items():
        if f in cols:
            continue
        p = spec["pos"]
        free = [c for c in cands[f] if c not in used]
        if len(free) == 1:                                                # 2: the heading moved
            c = free[0]
            cols[f] = c
            how[f] = ("heading", "MEDIUM", f"expected at column {get_column_letter(p + 1)}, found by heading at {get_column_letter(c + 1)}")
            used.add(c)
            ctx.issue("W902", "WARNING", "", f, "all",
                      f"{_LABEL[kind]}: '{f}' is not at its usual column {get_column_letter(p + 1)}; found by heading "
                      f"('{display[c]}') at column {get_column_letter(c + 1)} and read from there.")
        elif p < width:                                                    # 3: read by position as the original did
            cols[f] = p
            how[f] = ("position", "LOW", f"heading at this column reads '{display[p]}'")
            used.add(p)
            ctx.issue("W901", "WARNING", "", f, "all",
                      f"{_LABEL[kind]}: cannot confirm the column for '{f}' from its heading (column "
                      f"{get_column_letter(p + 1)} reads '{display[p]}'); read by position as before - check the figures.")
        else:
            cols[f], how[f] = None, ("", "", "column does not exist in this file")
    return cols, how


def _unparsable(v):
    s = str(v).strip()
    if s in ("", "-", "–", "NA", "N/A"):
        return False
    try:
        float(s.replace(",", "").replace("₹", "").strip())
        return False
    except ValueError:
        return True


def build_canonical(path, kind):
    """Convert one ledger CSV (of the given kind) into canonical data. A file that cannot be used raises
    PeriodParseError ('[E90x] ...')."""
    if isinstance(path, (list, tuple)):
        path = path[0]
    ctx = cc.IssueLog(_LABEL[kind], [os.path.basename(path)])
    kdef = _kinds()[kind]
    try:
        rows = _read_csv(path, ctx)
    except Exception as e:  # noqa: BLE001 - OSError, csv.Error, ...
        _fail(ctx, path, kind, "E901", f"cannot open the file ({type(e).__name__}: {e}).")
    hdr_idx = next((i for i, r in enumerate(rows[:30]) if r and cc.clean(r[0]).startswith("srno")), None)
    if hdr_idx is None:
        _fail(ctx, path, kind, "E902", f"no heading row (the row starting 'Sr.No') found in the first 30 lines. This does not look "
              f"like a {_LABEL[kind]} download.")
    title = next((" ".join(c for c in r if c).strip() for r in rows[:hdr_idx] if any(r)), "")
    if kdef["title_keyword"] not in title.lower():
        ctx.issue("W903", "WARNING", "", "", "all",
                  f"The file's title reads '{title}' but it is being read as a {_LABEL[kind]}.")
    pre = {}
    for r in rows[:hdr_idx]:
        cells = [c.strip() for c in r if c.strip()]
        if len(cells) >= 2 and cells[0].lower() in ("gstin", "from", "to", "tax period") and cells[0].lower() not in pre:
            pre[cells[0].lower()] = cells[1]
    data_rows = [(i, r) for i, r in enumerate(rows[hdr_idx + 2:], start=hdr_idx + 3) if any(c.strip() for c in r)]
    width = max([len(rows[hdr_idx]), len(rows[hdr_idx + 1]) if hdr_idx + 1 < len(rows) else 0] + [len(r) for _, r in data_rows[:2000]])
    group, sub, sub_len, display = _labels(rows, hdr_idx, width)
    cols, how = _resolve(kdef, group, sub, sub_len, display, ctx, kind)
    missing = [f for f, spec in kdef["fields"].items() if spec.get("required") and cols.get(f) is None]
    mapping = []
    for f, spec in kdef["fields"].items():
        c = cols[f]
        mb, conf, note = how[f]
        mapping.append(dict(canonical_field=f, raw_header_found=display[c] if c is not None else "",
                            raw_column=get_column_letter(c + 1) if c is not None else "", matched_by=mb, confidence=conf,
                            status="OK" if c is not None else "MISSING", rows_filled=0, rows_blank=0, note=note))
    if missing:
        _fail(ctx, path, kind, "E903", f"required column(s) {missing} could not be located (headings found: "
              f"{sorted(set(d for d in display if d))[:30]}). The ledger is unusable, so the ledger checks show N/A.",
              "Supply the ledger as downloaded from the portal, or add the new heading to ledger_mapping.json.")
    mapped = {c for c in cols.values() if c is not None}
    fill = {f: [0, 0] for f in kdef["fields"]}
    out_rows, bad_amounts, unkeyed = [], 0, 0
    for ridx, r in data_rows:
        rec = dict(src_row=ridx, src_width=len(r))
        for f, spec in kdef["fields"].items():
            c = cols[f]
            v = r[c] if (c is not None and c < len(r)) else None
            fill[f][0 if (v is not None and str(v).strip() != "") else 1] += 1
            if spec["type"] == "num":
                if v is not None and _unparsable(v):
                    bad_amounts += 1
                rec[f] = cc.num(v)
            else:
                rec[f] = v
        rec["extra"] = "; ".join(f"{display[i] or get_column_letter(i + 1)}={r[i]}" for i in range(min(len(r), width))
                                 if i not in mapped and r[i].strip())
        out_rows.append(rec)
    for m in mapping:
        m["rows_filled"], m["rows_blank"] = fill[m["canonical_field"]]
    if bad_amounts:
        ctx.issue("W904", "WARNING", "", "", "all",
                  f"{bad_amounts} amount cell(s) contain text that is not a number and are read as 0 (as the original did).")
    if not any(str(r.get("description") or "").strip() == "Opening Balance" for r in out_rows):
        ctx.issue("I902", "INFO", "", "", "all", "No 'Opening Balance' row found; the opening balance is treated as not stated.")
    fy = _fy_from(pre.get("from", ""), pre.get("tax period", ""))
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    meta = [("schema_version", SCHEMA_VERSION), ("ledger_kind", kind), ("ledger_title", title),
            ("recipient_gstin", pre.get("gstin", "")), ("fy", fy), ("from_date", pre.get("from", "")),
            ("to_date", pre.get("to", "")), ("tax_period_text", pre.get("tax period", "")),
            ("source_files", "; ".join(ctx.source_files)), ("adapter_version", ADAPTER_VERSION),
            ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")), ("row_count", str(len(out_rows))),
            ("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK"))]
    return dict(meta=meta, rows=out_rows, mapping=mapping, issues=ctx.issues)


# ----------------------------------------------------------------------------------------
# write / read the canonical workbook
# ----------------------------------------------------------------------------------------
_WRITTEN = {}


def write_canonical(data, out_dir):
    m = dict(data["meta"])
    prefix = _kinds()[m["ledger_kind"]]["canon_prefix"]
    name = cc.canonical_filename(prefix, data["meta"])
    src = m.get("source_files")
    base, ext = os.path.splitext(name)
    n = 1
    while _WRITTEN.get(os.path.abspath(os.path.join(out_dir, name)), src) != src:
        n += 1
        name = f"{base}_{n}{ext}"
    sheets = [(ROWS_SHEET, _canon_cols(m["ledger_kind"]), data["rows"]), ("MAPPING_REPORT", MAP_COLS, data["mapping"]),
              ("ISSUES", ISSUE_COLS, data["issues"])]
    path = cc.write_workbook(out_dir, name, data["meta"], sheets, max_col_width=28)
    _WRITTEN[os.path.abspath(path)] = src
    return path


def is_canonical_file(path):
    return cc.is_canonical(path, "LED-", ("META", ROWS_SHEET, "MAPPING_REPORT"))


def canonical_kind(path):
    """The ledger kind recorded in a canonical ledger workbook's META (None if it is not one)."""
    if not is_canonical_file(path):
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return dict(cc.read_meta(wb)).get("ledger_kind")
    finally:
        wb.close()


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        meta = cc.read_meta(wb)
        kind = dict(meta).get("ledger_kind")
        return dict(meta=meta, rows=cc.read_table(wb, ROWS_SHEET, _canon_cols(kind)),
                    mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS), issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------
# engine-facing access
# ----------------------------------------------------------------------------------------
_SRC = {k: cc.CanonicalCache(_LABEL[k], is_canonical_file, read_canonical, lambda p, k=k: build_canonical(p, k),
                             write_canonical, lambda d: dict(meta=dict(d["meta"]), rows=d["rows"]))
        for k in KINDS}


def get_data(path, kind, out_dir=None):
    return _SRC[kind].get(path, out_dir)


def clear_cache():
    for c in _SRC.values():
        c.clear()
    _WRITTEN.clear()


def _s(v):
    return str(v or "").strip()


def _f(v):
    return 0.0 if v is None else float(v)


def _empty(kind):
    if kind == "credit":
        return dict(opening=None, transactions=[], monthly_by_tax_period={})
    return dict(opening=None, transactions=[], monthly_by_tax_period={}, monthly_by_txn_date={})


def _rows_or_empty(path, kind):
    """The canonical rows, or None after printing the plain reason (the caller returns the empty shape)."""
    try:
        return get_data(path, kind)["_index"]["rows"]
    except mpu.PeriodParseError as e:
        print(f"[error] {_LABEL[kind]} could not be read -> its checks show N/A: {e}")
        return None


def parse_cash_or_liability_ledger(path, kind):
    """Legacy gst_parsers_dept.parse_cash_or_liability_ledger(path, kind)."""
    from gst_parsers_dept import _period_key, _tax_period_key
    rows = _rows_or_empty(path, kind)
    if rows is None:
        return _empty(kind)
    kdef = _kinds()[kind]
    fields = kdef["fields"]
    stay_pos = fields["stay_status"]["pos"] if "stay_status" in fields else None
    opening, transactions = None, []
    for r in rows:
        if r["src_width"] < kdef["min_width"]:
            continue
        first, desc = _s(r["sr_no"]), _s(r["description"])
        if not first and desc != "Opening Balance":
            continue
        if not (first.isdigit() or first in ("-",) or desc == "Opening Balance"):
            continue
        heads = {}
        for h in HEADS:
            k = h.lower()
            heads[h] = dict(tax=_f(r[f"amt_{k}_tax"]), total=_f(r[f"amt_{k}_total"]), balance=_f(r[f"bal_{k}_total"]))
        rec = dict(date=_s(r["date"]), ref=_s(r["ref"]), tax_period=_s(r.get("tax_period")) if "tax_period" in fields else "-",
                   description=desc, ttype=_s(r["ttype"]), heads=heads, total=sum(h["total"] for h in heads.values()),
                   balance_total=sum(h["balance"] for h in heads.values()),
                   ledger_used=_s(r["ledger_used"]) if "ledger_used" in fields else None,
                   demand_id=_s(r["demand_id"]) if "demand_id" in fields else None,
                   stay_status=(_s(r["stay_status"]) if r["src_width"] > stay_pos else None) if stay_pos is not None else None)
        if desc == "Opening Balance":
            opening = rec
            continue
        if desc == "Closing Balance":
            continue
        transactions.append(rec)
    by_tp, by_date = {}, {}
    for t in transactions:
        sign = 1 if t["ttype"].lower() == "credit" else -1
        amt = t["total"]
        tpk = _tax_period_key(t["tax_period"]) if kind == "cash" else None
        dpk = _period_key(t["date"])
        for key, bucket in ((tpk, by_tp), (dpk, by_date)):
            if key:
                bucket.setdefault(key, dict(credited=0.0, debited=0.0, net=0.0))
                bucket[key]["credited" if t["ttype"].lower() == "credit" else "debited"] += amt
                bucket[key]["net"] += sign * amt
    return dict(opening=opening, transactions=transactions, monthly_by_tax_period=by_tp, monthly_by_txn_date=by_date)


def parse_credit_ledger(path):
    """Legacy gst_parsers_dept.parse_credit_ledger(path)."""
    from gst_parsers_dept import _tax_period_key
    rows = _rows_or_empty(path, "credit")
    if rows is None:
        return _empty("credit")
    kdef = _kinds()["credit"]
    opening, transactions = None, []
    for r in rows:
        if r["src_width"] < kdef["min_width"]:
            continue
        desc, first = _s(r["description"]), _s(r["sr_no"])
        if not (first.isdigit() or desc == "Opening Balance"):
            continue
        rec = dict(date=_s(r["date"]), ref=_s(r["ref"]), tax_period=_s(r["tax_period"]), description=desc,
                   ttype=_s(r["ttype"]), igst=_f(r["amt_igst"]), cgst=_f(r["amt_cgst"]), sgst=_f(r["amt_sgst"]),
                   cess=_f(r["amt_cess"]), total=_f(r["amt_total"]), bal_igst=_f(r["bal_igst"]), bal_cgst=_f(r["bal_cgst"]),
                   bal_sgst=_f(r["bal_sgst"]), bal_cess=_f(r["bal_cess"]), bal_total=_f(r["bal_total"]))
        if desc == "Opening Balance":
            opening = rec
            continue
        transactions.append(rec)
    by_tp = {}
    for t in transactions:
        tpk = _tax_period_key(t["tax_period"])
        if not tpk:
            continue
        m = by_tp.setdefault(tpk, dict(credited=0.0, debited=0.0, net=0.0, accrued_desc=set()))
        sign = 1 if t["ttype"].lower() == "credit" else -1
        if t["ttype"].lower() == "credit":
            m["credited"] += t["total"]
            m["accrued_desc"].add(t["description"])
        else:
            m["debited"] += t["total"]
        m["net"] += sign * t["total"]
    for m in by_tp.values():
        m["accrued_desc"] = sorted(m["accrued_desc"])
    return dict(opening=opening, transactions=transactions, monthly_by_tax_period=by_tp)
