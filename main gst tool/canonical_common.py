#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CANONICAL LAYER - SHARED BUILDING BLOCKS
========================================
Everything the per-source converters (gstr2b_adapter.py, gstr3b_adapter.py, and the ones still to come)
have in common, so a new converter only has to say what its OWN form/sheet layout looks like:

  * IssueLog            - collects warnings/errors (ISSUES sheet), the field-mapping report and fill counts;
                          prints each one as "[warn] GSTR-3B W303: ..."; .fail() raises the plain-message error
  * load_json           - the mapping table (column spellings / line rules), loaded once
  * clean / text / num / is_blank - the cell helpers every source uses
  * write_workbook      - META + data sheets + MAPPING_REPORT + ISSUES, in one fixed look
  * read_meta / read_table / is_canonical / canonical_filename - reading a canonical file back
  * CanonicalCache      - "raw file -> build -> write canonical file -> read it back (the engine reads the
                          workbook, not memory) -> cache"; failures are cached too, so a bad file is
                          reported once, not once per month

ERROR-level problems are raised as gst_core.PeriodParseError with a full plain message, which the engine
already turns into "<source> not supplied (<reason>)" / "ERROR processing <month> ... SKIPPED".
"""

import json
import os
import re

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import gst_core as mpu

HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")

# The two sheets every canonical workbook carries besides its data sheets.
ISSUE_COLS = ["id", "severity", "sheet", "field", "periods_affected", "message", "what_is_skipped",
              "suggested_action"]


# ----------------------------------------------------------------------------------------
# cell helpers
# ----------------------------------------------------------------------------------------
def clean(s):
    """Lower-case and strip every non-alphanumeric character (header / label comparison)."""
    return re.sub(r"[^a-z0-9]+", "", str(s or "").lower())


def text(v):
    return str(v or "").strip()


def num(v):
    """Cell -> float. '-', '', None and anything unreadable -> 0.0; commas and the rupee sign are ignored."""
    if v is None:
        return 0.0
    s = str(v).strip()
    if s in ("", "-", "–"):
        return 0.0
    s = s.replace(",", "").replace("₹", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


# ----------------------------------------------------------------------------------------
# mapping table
# ----------------------------------------------------------------------------------------
_JSON_CACHE = {}


def load_json(path):
    """Read a mapping-table JSON file once per process."""
    if path not in _JSON_CACHE:
        with open(path, encoding="utf-8") as f:
            _JSON_CACHE[path] = json.load(f)
    return _JSON_CACHE[path]


# ----------------------------------------------------------------------------------------
# issues
# ----------------------------------------------------------------------------------------
class IssueLog:
    """Collects the issues, the field-mapping report and per-field fill counts of ONE build.
    `label` (e.g. "GSTR-3B") prefixes every printed line."""

    def __init__(self, label, source_files=None):
        self.label = label
        self.source_files = source_files or []
        self.issues = []
        self.mapping = []
        self.fill = {}

    def issue(self, iid, sev, sheet, field, periods, message, skipped="", action=""):
        self.issues.append(dict(id=iid, severity=sev, sheet=sheet, field=field, periods_affected=periods,
                                message=message, what_is_skipped=skipped, suggested_action=action))
        tag = {"ERROR": "[error]", "WARNING": "[warn]", "INFO": "[info]"}[sev]
        print(f"{tag} {self.label} {iid}: {message}")

    def count(self, sheet, field, filled):
        """Count one source cell as filled / blank for the mapping report."""
        c = self.fill.setdefault((sheet, field), [0, 0])
        c[0 if filled else 1] += 1

    def fail(self, iid, message, sheet="", field="", skipped="", action=""):
        """Record an ERROR and raise it as PeriodParseError("[<id>] <message>")."""
        self.issue(iid, "ERROR", sheet, field, "all", message, skipped, action)
        raise mpu.PeriodParseError(f"[{iid}] {message}")


# ----------------------------------------------------------------------------------------
# writing
# ----------------------------------------------------------------------------------------
def table_sheet(wb, title, cols, rows, max_width=40):
    """One data sheet: styled header row, one row per dict, columns in `cols` order."""
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([r.get(c) for c in cols])
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(max_width, len(c) + 4))
    return ws


def canonical_filename(prefix, meta):
    """Canonical_<prefix>_<GSTIN>_<FY>.xlsx from the META (key, value) pairs."""
    m = dict(meta)
    return f"Canonical_{prefix}_{m.get('recipient_gstin') or 'UNKNOWN'}_{(m.get('fy') or 'FY').replace(', ', '_')}.xlsx"


def write_workbook(out_dir, filename, meta, tables, meta_key_width=30, max_col_width=40):
    """META (key/value) followed by `tables` = [(sheet title, columns, list of row dicts), ...].
    Returns the path written."""
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "META"
    ws.append(["key", "value"])
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
    for k, v in meta:
        ws.append([k, v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = meta_key_width, 90
    for title, cols, rows in tables:
        table_sheet(wb, title, cols, rows, max_col_width)
    wb.save(path)
    return path


# ----------------------------------------------------------------------------------------
# reading
# ----------------------------------------------------------------------------------------
def is_canonical(path, schema_prefix, required_sheets):
    """True if `path` is an .xlsx with the given sheets and a META schema_version starting with
    `schema_prefix` (e.g. '3B-'). Never raises."""
    try:
        if not str(path).lower().endswith((".xlsx", ".xlsm")):
            return False
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if any(s not in wb.sheetnames for s in required_sheets):
                return False
            for row in wb["META"].iter_rows(min_row=1, max_row=3, values_only=True):
                if row and row[0] == "schema_version":
                    return str(row[1]).startswith(schema_prefix)
        finally:
            wb.close()
    except Exception:  # noqa: BLE001 - any failure just means "not one of ours"
        return False
    return False


def read_meta(wb):
    return [(r[0], r[1]) for r in wb["META"].iter_rows(min_row=2, values_only=True) if r and r[0]]


def read_table(wb, name, cols):
    """A data sheet back into a list of dicts (only the wanted columns; missing ones read as None)."""
    if name not in wb.sheetnames:
        return []
    it = wb[name].iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        return []
    idx = {h: i for i, h in enumerate(header)}
    return [{c: (r[idx[c]] if c in idx and idx[c] < len(r) else None) for c in cols} for r in it if any(r)]


# ----------------------------------------------------------------------------------------
# build once, cache
# ----------------------------------------------------------------------------------------
class CanonicalCache:
    """The engine-facing flow every converter shares.

    get(path): a canonical file is read as-is; a raw file is built, WRITTEN to <cwd>/_canonical/ (a
    sub-folder, so it is never mistaken for an input) and READ BACK, so the workbook you can open is
    exactly what the engine used. Results - and errors - are cached per path.
    """

    def __init__(self, label, is_canonical_fn, read_fn, build_fn, write_fn, finalize_fn):
        self.label = label
        self._is_canonical = is_canonical_fn
        self._read = read_fn
        self._build = build_fn
        self._write = write_fn
        self._finalize = finalize_fn
        self._cache = {}

    def get(self, path, out_dir=None):
        key = os.path.abspath(path)
        if key in self._cache:
            hit = self._cache[key]
            if isinstance(hit, Exception):
                raise hit
            return hit
        try:
            if self._is_canonical(path):
                data = self._read(path)
            else:
                data = self._build(path)
                try:
                    out = self._write(data, out_dir or os.path.join(os.getcwd(), "_canonical"))
                    print(f"[info] {self.label} canonical file written: {out}")
                    data = self._read(out)          # the engine reads the workbook, not memory
                except OSError as e:
                    print(f"[info] {self.label} canonical file could not be written ({e}); "
                          f"using it from memory only.")
        except mpu.PeriodParseError as e:
            self._cache[key] = e
            raise
        data["_index"] = self._finalize(data)
        self._cache[key] = data
        return data

    def clear(self):
        self._cache.clear()
