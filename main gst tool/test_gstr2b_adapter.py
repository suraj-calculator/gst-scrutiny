#!/usr/bin/env python3
"""
Regression tests for the GSTR-2B canonical layer (gstr2b_adapter.py, docs/GSTR2B_CANONICAL_SPEC.md).

Self-contained: builds a small synthetic merged GSTR-2B (no taxpayer data needed) containing the
awkward cases the real files have, then checks that the ORIGINAL reader
(gst_parsers_returns.parse_2b_excel, direct) and the canonical route (gstr2b_adapter.parse_month)
return identical results, and that broken inputs give the documented error IDs.

Run:  python3 test_gstr2b_adapter.py
"""
import os
import shutil
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.environ["GST_2B_CANONICAL"] = "0"          # original reader for the legacy side

import gst_core as mpu                        # noqa: E402
import gst_parsers_returns as pr              # noqa: E402
import gstr2b_adapter as ad                   # noqa: E402

BANNER = "Financial Year: 2024-25  |  Tax Period: {}  |  Date of Generation: 01/09/2026"


def _b2b_header(ws):
    ws.append(["Goods and Services Tax  - GSTR-2B"] + [None] * 17)
    ws.append([None] * 18)
    ws.append([None] * 18)
    ws.append(["Taxable inward supplies received from registered persons"] + [None] * 17)
    ws.append(["GSTIN of supplier", "Trade/Legal name", "Invoice Details", None, None, None, "Place of supply",
               "Supply Attract Reverse Charge", "Rate(%)", "Taxable Value (₹)", "Tax Amount", None, None, None,
               "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability", "Reason"])
    ws.append([None, None, "Invoice number", "Invoice type", "Invoice Date", "Invoice Value(₹)", None, None, None,
               None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)", None, None, None, None])


def _row(gstin, inv, date, val, rate, taxable, igst, cgst, sgst, period, avail="Yes", shifted=False):
    base = [gstin, "SUPPLIER " + gstin[-3:], inv, "Regular", date, val, "Uttarakhand", "No"]
    tail = [taxable, igst, cgst, sgst, 0, period, "11/05/2024", avail, None]
    return base + (tail if shifted else [rate] + tail)


def make_fixture(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "B2B"
    _b2b_header(ws)
    G1, G2, G3 = "05AAAAA0001A1Z1", "05BBBBB0002B1Z2", "05CCCCC0003C1Z3"
    ws.append([BANNER.format("April")])
    ws.append(_row(G1, "INV1", "05/04/2024", 1180, 18.0, 1000, 0, 90, 90, "Apr'24"))
    ws.append(_row(G2, "INV2", "06/04/2024", 2360, None, 2000, 0, 180, 180, "Apr'24", shifted=True))  # blank Rate
    ws.append(_row(G3, "INV3", "07/04/2024", 590, 18.0, 500, 0, 45, 45, "Apr'24"))                     # amended later
    ws.append([BANNER.format("May")])
    ws.append(_row(G1, "INV4", "02/05/2024", 1180, 18.0, 1000, 180, 0, 0, "May'24"))
    ws.append([BANNER.format("May")])                                                                  # 2nd download
    ws.append(_row(G1, "INV4", "02/05/2024", 1180, 18.0, 1000, 180, 0, 0, "May'24"))                   # duplicate
    ws.append(_row(G2, "INV5", "03/05/2024", 1180, 18.0, 1000, 0, 90, 90, "May'24", avail="No"))       # new row

    cd = wb.create_sheet("B2B-CDNR")
    cd.append(["Goods and Services Tax  - GSTR-2B"] + [None] * 18)
    cd.append([None] * 19)
    cd.append([None] * 19)
    cd.append(["Debit/Credit notes (Original)"] + [None] * 18)
    cd.append(["GSTIN of supplier", "Trade/Legal name", "Credit note/Debit note details", None, None, None, None,
               "Place of supply", "Supply Attract Reverse Charge", "Rate(%)", "Taxable Value (₹)", "Tax Amount",
               None, None, None, "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability",
               "Reason"])
    cd.append([None, None, "Note number", "Note type", "Note Supply type", "Note date", "Note Value (₹)", None, None,
               None, None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)", None, None, None,
               None])
    cd.append([BANNER.format("April")])
    cd.append([G1, "SUPPLIER Z11", "CN1", "Credit Note", "Regular", "20/04/2024", 118, "Uttarakhand", "No", 18.0,
               100, 0, 9, 9, 0, "Apr'24", "11/05/2024", "Yes", None])
    cd.append([BANNER.format("May")])

    ba = wb.create_sheet("B2BA")
    for _ in range(4):
        ba.append([None] * 19)
    ba.append(["Original Details", None, None, None, "Revised Details"] + [None] * 14)
    ba.append(["Invoice number", "Invoice Date", "GSTIN of supplier", "Trade/Legal name", "Invoice Details", None,
               None, None, "Place of supply", "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount",
               None, None, None, "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability",
               "Reason"])
    ba.append([None, None, None, None, "Invoice number", "Invoice type", "Invoice Date", "Invoice Value(₹)", None,
               None, None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)", None, None, None,
               None])
    ba.append([BANNER.format("April")])
    ba.append([BANNER.format("May")])
    ba.append(["INV3", "07/04/2024", G3, "SUPPLIER 3C3", "INV3-A", "Regular", "07/04/2024", 708, "Uttarakhand", "No",
               600, 0, 54, 54, 0, "May'24", "11/06/2024", "Yes", None])
    wb.save(path)
    return path


def _strip(d):
    d = dict(d)
    d.pop("canonical_issues", None)
    d.pop("amendment_warnings", None)
    return d


def check(name, cond, detail=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail and not cond else ""))
    return cond


def main():
    ok = True
    tmp = tempfile.mkdtemp()
    cwd = os.getcwd()
    os.chdir(tmp)
    try:
        raw = make_fixture(os.path.join(tmp, "synthetic_2b.xlsx"))
        ad.clear_cache()
        pr._2B_FILE_CACHE.clear()

        # ---- parity: original reader vs canonical route, every month the file has
        for month in ("Apr-24", "May-24"):
            old, new = _strip(pr.parse_2b_excel(raw, month)), _strip(ad.parse_month(raw, month))
            ok &= check(f"parity {month}: summary", old["summary"] == new["summary"])
            ok &= check(f"parity {month}: b2b rows identical ({len(old['b2b'])})", old["b2b"] == new["b2b"],
                        f"{old['b2b']} != {new['b2b']}")
            ok &= check(f"parity {month}: cdnr rows identical ({len(old['cdnr'])})", old["cdnr"] == new["cdnr"])
        # a month with no block: same exception type on both sides
        e_old = e_new = None
        try:
            pr.parse_2b_excel(raw, "Jun-24")
        except mpu.PeriodParseError as e:
            e_old = type(e)
        try:
            ad.parse_month(raw, "Jun-24")
        except mpu.PeriodParseError as e:
            e_new = type(e)
        ok &= check("parity Jun-24: both raise PeriodParseError (month not in file)", e_old is e_new is not None)

        # ---- the awkward cases behave as specified
        apr = ad.parse_month(raw, "Apr-24")["b2b"]
        may = ad.parse_month(raw, "May-24")["b2b"]
        ok &= check("shifted row (blank Rate): rate None, amounts read from shifted columns",
                    [r for r in apr if r["invno"] == "INV2"][0]["rate"] is None and
                    [r for r in apr if r["invno"] == "INV2"][0]["taxable"] == 2000)
        ok &= check("superseded original INV3 dropped from April", all(r["invno"] != "INV3" for r in apr))
        ok &= check("amendment INV3-A spliced into May (its own filing period)",
                    any(r["invno"] == "INV3-A" and r["via_amendment"] for r in may))
        ok &= check("May came in 2 downloads: duplicate INV4 skipped, INV5 kept",
                    sorted(r["invno"] for r in may if not r["via_amendment"]) == ["INV4", "INV5"])

        # ---- the canonical workbook itself
        out = ad.write_canonical(ad.build_canonical(raw), os.path.join(tmp, "_canonical"))
        ok &= check("canonical file is recognised as canonical", ad.is_canonical_file(out))
        ok &= check("canonical file months", ad.months_from_canonical(out) == {"Apr-24", "May-24"})
        d = ad.read_canonical(out)
        ok &= check("canonical B2B rows carry period + block_no",
                    {(r["period"], r["block_no"]) for r in d["b2b"]} == {("Apr-24", 1), ("May-24", 1), ("May-24", 2)})
        ids = {i["id"] for i in d["issues"]}
        ok &= check("issues report the shifted row (W303) and the 2 downloads (W304)", {"W303", "W304"} <= ids, str(ids))
        ad.clear_cache()
        ok &= check("engine reads the canonical file directly (same result)",
                    _strip(ad.parse_month(out, "Apr-24"))["b2b"] == _strip(pr.parse_2b_excel(raw, "Apr-24"))["b2b"])

        # ---- documented errors
        def err(label, path, expect):
            ad.clear_cache()
            try:
                ad.build_canonical(path)
                return check(label, False, "no error raised")
            except mpu.PeriodParseError as e:
                return check(label, str(e).startswith(expect), str(e)[:120])

        wb = openpyxl.load_workbook(raw)
        for row in wb["B2B"].iter_rows(min_row=1, max_row=6):
            for c in row:
                if c.value and "Taxable Value" in str(c.value):
                    c.value = "Taxable Amt"
        wb.save("renamed.xlsx")
        ok &= err("E201 required header renamed", "renamed.xlsx", "[E201]")
        open("corrupt.xlsx", "wb").write(open(raw, "rb").read()[:2000])
        ok &= err("E204 corrupt file", "corrupt.xlsx", "[E204]")
        wb = openpyxl.Workbook()
        wb.active.title = "Summary"
        wb.save("wrong.xlsx")
        ok &= err("E203 no 2B sheets", "wrong.xlsx", "[E203]")
        wb = openpyxl.load_workbook(raw)
        for sn in wb.sheetnames:
            ws = wb[sn]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("Financial Year:"):
                        c.value = None
        wb.save("nobanner.xlsx")
        ok &= err("E202 no period anywhere (no banner, no Read me, junk filename)", "nobanner.xlsx", "[E202]")
    finally:
        os.chdir(cwd)
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nALL PASSED" if ok else "\nSOME TESTS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
