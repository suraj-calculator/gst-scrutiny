#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PORTAL COMPARISON CANONICAL LAYER  (spec: docs/PORTAL_COMPARISON_CANONICAL_SPEC.md)
=====================================================================================
Only this module knows what the portal's "Tax liability and ITC comparison" report looks like.

    raw report  --build_canonical()-->  canonical data
    data        --write_canonical()-->  Canonical_PORTALCOMPARISON_<GSTIN>_<FY>.xlsx
    .xlsx       --read_canonical()  -->  data
    data        --parse_portal_comparison()  -->  exactly what gst_parsers_dept.parse_portal_comparison() returns

The engine reads only the 'Comparison Summary' sheet: one row per tax period ('Apr-25'), with the liability per GSTR-1 and GSTR-3B, the ITC per
GSTR-3B and GSTR-2B (excluding reversal) and, in newer reports, the ITC considering reversal - each with the shortfall and the cumulative shortfall.
The original took every figure from a fixed column. The converter reads each column by its heading first and falls back to the original position,
reporting a moved column (W1402) or an unconfirmable one (W1403) in plain words.
"""

import datetime as _dt
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

import canonical_common as cc
import gst_core as mpu

SCHEMA_VERSION = "PC-1.0"
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "portal_mapping.json")
ROWS_SHEET = "COMPARISON_ROWS"
MAP_COLS = ["canonical_field", "raw_header_found", "raw_column", "matched_by", "confidence", "status", "rows_filled", "rows_blank", "note"]
ISSUE_COLS = cc.ISSUE_COLS
_PERIOD_RE = re.compile(r"([A-Za-z]{3})-(\d{2})$")
_CONTEXT = {"gstin": None}


def set_context(gstin=None):
    if gstin:
        _CONTEXT["gstin"] = gstin


def _map():
    return cc.load_json(_MAPPING_FILE)


def _canon_cols():
    return ["src_row", "period"] + list(_map()["fields"])


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def _labels(rows, hdr, first_data):
    """Heading block = the row containing 'Tax Period' down to the row before the first period row. The LAST row of the block holds the
    sub-headings; every earlier row is a group-heading row whose text spans the columns to its right (forward-filled), and the texts of
    all group rows are joined (older reports have three heading rows). Returns (group, sub, display) per column, cleaned / display."""
    block = rows[hdr:first_data]
    subrow, grows = block[-1], block[:-1]
    width = max(len(r) for r in block)
    group = [""] * width
    for gr in grows:
        cur = ""
        for c in range(width):
            v = gr[c] if c < len(gr) else None
            if not _blank(v):
                cur = cc.clean(v)
            group[c] += cur
    subc = [cc.clean(subrow[c]) if c < len(subrow) and not _blank(subrow[c]) else "" for c in range(width)]
    disp = [(str(subrow[c]).strip() if c < len(subrow) and not _blank(subrow[c]) else "") for c in range(width)]
    return group, subc, disp


def _fits(spec, group, sub, c):
    return (all(t in group[c] for t in spec["group"]) and not any(t in group[c] for t in spec.get("group_none", []))
            and all(t in sub[c] for t in spec["sub_has"]) and not any(t in sub[c] for t in spec.get("sub_none", [])))


def build_canonical(path):
    """Convert one portal comparison report into canonical data. A file that cannot be opened raises PeriodParseError ('[E1401] ...')."""
    if isinstance(path, (list, tuple)):
        path = path[0]
    ctx = cc.IssueLog("Portal comparison", [os.path.basename(path)])
    try:
        wb = openpyxl.load_workbook(path, data_only=True)       # full (not read-only) load: portal files can carry a wrong <dimension>
    except Exception as e:  # noqa: BLE001
        ctx.fail("E1401", f"Portal comparison \"{os.path.basename(path)}\": cannot open the file ({type(e).__name__}: {e}). It may be corrupt, "
                 f"password-protected or not an Excel file.", skipped="the portal comparison checks", action="Re-download the report.")
    m = _map()
    fields = m["fields"]
    rows_out, mapping, meta = [], [], {}
    if m["sheet"] not in wb.sheetnames:
        ctx.issue("W1401", "WARNING", "", "", "all", f"The workbook has no '{m['sheet']}' sheet (sheets: {wb.sheetnames}); the portal comparison reads as empty.",
                  "the portal comparison checks", "Check that this is the 'Tax liability and ITC comparison' report.")
    else:
        rows = [list(r) for r in wb[m["sheet"]].iter_rows(values_only=True)]
        for r in rows[:10]:
            for c in r:
                if isinstance(c, str):
                    for lab, k in (("GSTIN:", "gstin"), ("Legal name:", "legal_name"), ("Financial Year:", "fy")):
                        if c.strip().lower().startswith(lab.lower()) and k not in meta:
                            meta[k] = c.split(":", 1)[1].strip()
        hdr = next((i for i, r in enumerate(rows[:25]) if r and cc.clean(r[0]) == "taxperiod"), None)
        cols, how = {}, {}
        if hdr is None:
            ctx.issue("W1404", "WARNING", m["sheet"], "", "all",
                      "No heading row starting 'Tax Period' was found, so the columns cannot be confirmed from their headings; every figure is read by position as before.")
            for f, spec in fields.items():
                cols[f], how[f] = spec["pos"], ("position", "LOW", "no heading row found")
            disp = []
        else:
            first = next((i for i in range(hdr + 1, len(rows)) if rows[i] and rows[i][0] and _PERIOD_RE.match(str(rows[i][0]).strip())), None)
            group, sub, disp = _labels(rows, hdr, first if first is not None else min(hdr + 3, len(rows)))
            used = set()
            cands = {f: [c for c in range(len(group)) if _fits(spec, group, sub, c)] for f, spec in fields.items()}
            for f, spec in fields.items():
                if spec["pos"] in cands[f] and spec["pos"] not in used:
                    cols[f], how[f] = spec["pos"], ("position", "HIGH", "")
                    used.add(spec["pos"])
            for f, spec in fields.items():
                if f in cols:
                    continue
                p, free = spec["pos"], [c for c in cands[f] if c not in used]
                if len(free) == 1:
                    cols[f], how[f] = free[0], ("heading", "MEDIUM", f"expected at column {get_column_letter(p + 1)}")
                    used.add(free[0])
                    ctx.issue("W1402", "WARNING", m["sheet"], f, "all",
                              f"Portal comparison: '{f}' is not at its usual column {get_column_letter(p + 1)}; found by heading at column "
                              f"{get_column_letter(free[0] + 1)} ('{disp[free[0]]}') and read from there.")
                elif p < len(group):
                    cols[f], how[f] = p, ("position", "LOW", f"heading at this column reads '{disp[p]}'")
                    used.add(p)
                    ctx.issue("W1403", "WARNING", m["sheet"], f, "all",
                              f"Portal comparison: cannot confirm the column for '{f}' from its heading (column {get_column_letter(p + 1)} reads "
                              f"'{disp[p]}'); read by position as before - check the figures.")
                else:
                    cols[f], how[f] = None, ("", "", "column does not exist")
        missing = [f for f, s in fields.items() if s.get("required") and cols[f] is None]
        for f in fields:
            c = cols[f]
            mapping.append(dict(canonical_field=f, raw_header_found=(disp[c] if (c is not None and c < len(disp)) else ""),
                                raw_column=get_column_letter(c + 1) if c is not None else "", matched_by=how[f][0], confidence=how[f][1],
                                status="OK" if c is not None else "NOT_IN_THIS_REPORT" if not fields[f].get("required") else "MISSING",
                                rows_filled=0, rows_blank=0, note=how[f][2]))
        if missing:
            ctx.fail("E1403", f"Portal comparison \"{os.path.basename(path)}\": required column(s) {missing} could not be located "
                     f"(headings found: {sorted(set(d for d in disp if d))[:12]}). The portal comparison reads as empty.",
                     sheet=m["sheet"], skipped="the portal comparison checks", action="Supply the report as downloaded from the portal.")
        for ridx, r in enumerate(rows, start=1):
            if not r or not r[0]:
                continue
            period = str(r[0]).strip()
            if not _PERIOD_RE.match(period):
                continue
            rec = dict(src_row=ridx, period=period)
            for f, spec in fields.items():
                c = cols[f]
                v = r[c] if (c is not None and c < len(r)) else None
                rec[f] = None if (spec["type"] == "num_or_none" and _blank(v)) else cc.num(v)
            rows_out.append(rec)
        for mp in mapping:
            mp["rows_filled"] = sum(1 for r in rows_out if r[mp["canonical_field"]] is not None)
            mp["rows_blank"] = len(rows_out) - mp["rows_filled"]
        if not rows_out:
            ctx.issue("W1405", "WARNING", m["sheet"], "", "all", "No tax-period rows ('Apr-25' style) were found under the heading; the portal comparison reads as empty.",
                      "the portal comparison checks", "")
    unread = [n for n in wb.sheetnames if n != m["sheet"]]
    wb.close()
    if unread:
        ctx.issue("I1401", "INFO", "", "", "all", f"Sheets not used by any check: {unread}.")
    fy = meta.get("fy", "")
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    cmeta = [("schema_version", SCHEMA_VERSION), ("recipient_gstin", _CONTEXT["gstin"] or meta.get("gstin", "")),
             ("legal_name", meta.get("legal_name", "")), ("fy", fy), ("source_files", "; ".join(ctx.source_files)),
             ("adapter_version", ADAPTER_VERSION), ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
             ("row_count", str(len(rows_out))), ("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK"))]
    return dict(meta=cmeta, rows=rows_out, mapping=mapping, issues=ctx.issues)


_WRITTEN = {}


def write_canonical(data, out_dir):
    name = cc.canonical_filename("PORTALCOMPARISON", data["meta"])
    src = dict(data["meta"]).get("source_files")
    base, ext = os.path.splitext(name)
    n = 1
    while _WRITTEN.get(os.path.abspath(os.path.join(out_dir, name)), src) != src:
        n += 1
        name = f"{base}_{n}{ext}"
    sheets = [(ROWS_SHEET, _canon_cols(), data["rows"]), ("MAPPING_REPORT", MAP_COLS, data["mapping"]), ("ISSUES", ISSUE_COLS, data["issues"])]
    path = cc.write_workbook(out_dir, name, data["meta"], sheets, max_col_width=30)
    _WRITTEN[os.path.abspath(path)] = src
    return path


def is_canonical_file(path):
    return cc.is_canonical(path, "PC-", ("META", ROWS_SHEET, "MAPPING_REPORT"))


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return dict(meta=cc.read_meta(wb), rows=cc.read_table(wb, ROWS_SHEET, _canon_cols()),
                    mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS), issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


_SRC = cc.CanonicalCache("Portal comparison", is_canonical_file, read_canonical, build_canonical, write_canonical,
                         lambda d: dict(meta=dict(d["meta"])))


def get_data(path, out_dir=None):
    return _SRC.get(path, out_dir)


def clear_cache():
    _SRC.clear()
    _WRITTEN.clear()


def _f(v):
    return None if v is None else float(v)


def parse_portal_comparison(path):
    """Legacy gst_parsers_dept.parse_portal_comparison(path)."""
    try:
        data = get_data(path)
    except mpu.PeriodParseError as e:
        print(f"[error] Portal comparison could not be read -> that comparison shows N/A: {e}")
        return {}
    out = {}
    for r in data["rows"]:
        mm = _PERIOD_RE.match(str(r["period"]))
        rec = {}
        for f, spec in _map()["fields"].items():
            rec[f] = _f(r[f]) if spec["type"] == "num_or_none" else float(r[f] or 0.0)
        out[f"{mm.group(1)}-{mm.group(2)}"] = rec
    return out
