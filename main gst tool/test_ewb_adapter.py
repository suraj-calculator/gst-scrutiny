#!/usr/bin/env python3
"""
Regression tests for the E-Way Bill canonical layer (ewb_adapter.py, docs/EWB_CANONICAL_SPEC.md).

Self-contained: builds small synthetic merged e-way bill workbooks (the flat, single-header-row, no-period-
marker shape a real inward/outward export has) and checks that the ORIGINAL parse_annual_ewb()/filter_by_month()
and the canonical route return identical results, that broken inputs give the documented errors, and that the
header-repeat-row / blank-row / unparsable-date handling matches the original exactly.

Run:  python3 test_ewb_adapter.py
"""
import os
import shutil
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.environ["GST_EWB_CANONICAL"] = "0"          # original reader for the legacy side

import gst_core as mpu                        # noqa: E402
import gst_parsers_returns as pr              # noqa: E402
import ewb_adapter as ad                      # noqa: E402

HDR = ["EWB No.", "From GSTIN & Name", "To GSTIN & Name", "From Place & Pin", "To Place & Pin",
       "EWB No. & Dt.", "Doc No. & Dt.", "Assess Val.", "Tax Val.", "HSN Code", "HSN Desc.", "Latest Vehicle No."]
SELF = "05AACFT2702L1ZD"
OTHER1 = "07AABCI6357A1ZV"
OTHER2 = "09CCCCC0003C1Z3"


def row(ewbno, from_gstin, to_gstin, from_place, to_place, ewb_dt, doc_dt, assess, taxval, hsn="87033299",
        hsn_desc="PARTS", vehicle="UP02TC0046"):
    return [ewbno, f"{from_gstin} / Some Sender", f"{to_gstin} / Some Receiver", from_place, to_place,
            f"{ewbno} - {ewb_dt}", f"{doc_dt[0]} - {doc_dt[1]}", assess, taxval, hsn, hsn_desc, vehicle]


def make_workbook(path, rows, hdr=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(hdr or HDR)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def fixture(path):
    return make_workbook(path, [
        row(301546430758, SELF, OTHER1, "Haldwani / 263139", "Agra / 282003", "10/04/2025 17:01:00",
            ("INV24A000152", "10/04/2025"), 1047083.79, 502600.22),
        row(301546430759, SELF, OTHER1, "Haldwani / 263139", "Agra / 282003", "20/05/2025 13:32:00",
            ("INV24A000731", "20/05/2025"), 1321576.41, 634356.67),
        row(301546430760, OTHER2, SELF, "Delhi / 110001", "Haldwani / 263139", "03/06/2025 09:10:00",
            ("SUP24A000001", "02/06/2025"), 55000.0, 9900.0),                       # inward-direction row (self is 'To')
        [None] * 12,                                                                # blank spacer row
    ])


def check(name, cond, detail=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail and not cond else ""))
    return cond


def run(fn, *a):
    try:
        return ("ok", fn(*a))
    except Exception as e:  # noqa: BLE001
        return ("EXC", type(e).__name__)


def compare_all(raw):
    ad.clear_cache()
    old, new = run(pr.parse_annual_ewb, raw), run(ad.parse_annual_ewb, raw)
    diffs = [] if old == new else [("parse_annual_ewb", str(old)[:200], str(new)[:200])]
    if old[0] == "ok":
        for m in sorted({r["month"] for r in old[1] if r["month"]}):
            a, b = pr.filter_by_month(old[1], m), ad.filter_by_month(new[1], m)
            if a != b:
                diffs.append((f"filter_by_month {m}", str(a)[:150], str(b)[:150]))
    return diffs


def _exc_text(fn, *a):
    try:
        fn(*a)
        return ""
    except Exception as e:  # noqa: BLE001
        return str(e)


def main():
    ok = True
    tmp = tempfile.mkdtemp()
    cwd = os.getcwd()
    os.chdir(tmp)
    try:
        raw = fixture(os.path.join(tmp, "ewb.xlsx"))
        d = compare_all(raw)
        ok &= check("parse_annual_ewb + filter_by_month identical to the original, every month", not d, str(d[:2]))

        # ---- the awkward cases actually exercised
        rows = ad.parse_annual_ewb(raw)
        ok &= check("3 real rows read (the blank spacer row dropped)", len(rows) == 3, str(len(rows)))
        r0 = rows[0]
        ok &= check("EWB no./date/time split out of the combined text",
                    r0["ewbno"] == "301546430758" and str(r0["ewbdate"]) == "2025-04-10" and str(r0["ewbtime"]) == "17:01:00",
                    str({k: r0[k] for k in ("ewbno", "ewbdate", "ewbtime")}))
        ok &= check("Doc No. & Dt. split, docno upper-cased", r0["docno"] == "INV24A000152" and str(r0["docdate"]) == "2025-04-10")
        ok &= check("month derived from EWB date (not doc date)", r0["month"] == "Apr-25", r0["month"])
        ok &= check("From/To GSTIN extracted from the combined 'GSTIN / Name' text",
                    r0["from_gstin"] == SELF and r0["to_gstin"] == OTHER1 and r0["from_name"] == "Some Sender")
        ok &= check("an inward-direction row (self in To) reads the same way", rows[2]["to_gstin"] == SELF and rows[2]["from_gstin"] == OTHER2)
        may = [r for r in rows if r["month"] == "May-25"][0]
        ok &= check("Assess Val. / Tax Val. converted to numbers", may["assess"] == 1321576.41 and may["taxval"] == 634356.67)

        # ---- canonical workbook on disk
        path = ad.write_canonical(ad.build_canonical(raw), tmp, source_path=raw)
        ok &= check("canonical file is recognised", ad.is_canonical_file(path) and os.path.basename(path).startswith("Canonical_EWB_"))
        ok &= check("canonical file supplied as input reads identically", ad.parse_annual_ewb(path) == ad.parse_annual_ewb(raw))
        wb = openpyxl.load_workbook(path, read_only=True)
        ok &= check("sheets: META, EWB_ROWS, MAPPING_REPORT, ISSUES", wb.sheetnames == ["META", "EWB_ROWS", "MAPPING_REPORT", "ISSUES"], str(wb.sheetnames))
        wb.close()
        ok &= check("classification: canonical workbook recognised as an e-way bill file",
                    mpu.use_canonical_ewb(path) and ad.is_canonical_file(path))
        other = fixture(os.path.join(tmp, "ewb_other.xlsx"))
        p2 = ad.write_canonical(ad.build_canonical(other), tmp, source_path=other)
        ok &= check("a second e-way bill file (different stem) does not overwrite the first", p2 != path and os.path.exists(path))
        p3 = ad.write_canonical(ad.build_canonical(raw), tmp, source_path=raw)
        ok &= check("the SAME source written twice -> same file (not '_2')", p3 == path)

        # ---- broken inputs
        ad.clear_cache()
        hdr_bad = [h for h in HDR if h != "Tax Val."]
        bad = make_workbook(os.path.join(tmp, "bad_cols.xlsx"), [row(1, SELF, OTHER1, "A", "B", "01/04/2025 00:00:00", ("X", "01/04/2025"), 1, 1)[:8] + ["87033299", "PARTS", "UP02TC0046"]], hdr=hdr_bad)
        t = _exc_text(ad.build_canonical, bad)
        ok &= check("E603 required column missing -> plain message naming the column and the headings found",
                    "[E603]" in t and "Tax Val." in t and "Headings found" in t, t[:200])
        wb = openpyxl.Workbook()
        wb.active.append(["Not", "An", "EWB", "File"])
        wb.save("wrong.xlsx")
        ok &= check("E602 not an e-way bill file", "[E602]" in _exc_text(ad.build_canonical, "wrong.xlsx"))
        with open("junk.xlsx", "w") as f:
            f.write("not an excel file")
        ok &= check("E601 unreadable file", "[E601]" in _exc_text(ad.build_canonical, "junk.xlsx"))

        # ---- header-repeat row + unparsable date (mirrors the mutation test run against real data)
        mut = make_workbook(os.path.join(tmp, "mut.xlsx"), [
            list(HDR),                                                              # a literal header-repeat data row
            row(2, SELF, OTHER1, "A", "B", "01/04/2025 00:00:00", ("Y", "01/04/2025"), 10, 1),
        ])
        ad.clear_cache()
        rows = ad.parse_annual_ewb(mut)
        ok &= check("a literal header-repeat row is dropped, like the original", len(rows) == 1, str(len(rows)))
        legacy_rows = pr.parse_annual_ewb(mut)
        ok &= check("...and the original drops it too (same row count)", len(legacy_rows) == 1, str(len(legacy_rows)))

        bad_date = make_workbook(os.path.join(tmp, "bad_date.xlsx"), [
            [3, f"{SELF} / X", f"{OTHER1} / Y", "A", "B", "not a real value", "Z - 01/04/2025", 5, 1, "", "", ""],
        ])
        ad.clear_cache()
        rows2 = ad.parse_annual_ewb(bad_date)
        ok &= check("unparsable EWB date -> month is None (row kept, just unusable for month-based checks)",
                    len(rows2) == 1 and rows2[0]["month"] is None and rows2[0]["ewbdate"] is None)
        ok &= check("W601 reported for the unparsable date", any(i["id"] == "W601" for i in ad.get_data(bad_date)["issues"]))
        ok &= check("the original reader gives the exact same result for this row", pr.parse_annual_ewb(bad_date) == rows2)

        # ---- optional column missing (TAX RATE never present in real exports)
        ad.clear_cache()
        no_rate = fixture(os.path.join(tmp, "no_rate.xlsx"))
        e = ad.get_data(no_rate)
        ok &= check("TAX RATE missing is absent_ok -- no ERROR/WARNING for it", not any("TAX RATE" in i["message"] for i in e["issues"]))
        ok &= check("rate reads as 0.0 when the column doesn't exist", ad.parse_annual_ewb(no_rate)[0]["rate"] == 0.0)
    finally:
        os.chdir(cwd)
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nALL PASSED" if ok else "\nSOME TESTS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
