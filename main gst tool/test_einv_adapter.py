#!/usr/bin/env python3
"""
Regression tests for the E-Invoice canonical layer (einv_adapter.py, docs/EINV_CANONICAL_SPEC.md).

Self-contained: builds small synthetic merged E-Invoice workbooks (several months, multi-rate invoices,
cancelled invoices, error rows, a numeric MMYYYY tax period, a quarterly banner) and checks that the ORIGINAL
readers and the canonical route return identical results for every reader and every month, that broken inputs
give the documented errors, and the deliberate fixes (case-insensitive headings such as 'Invoice value';
quarterly banner rows placed by their own invoice date).

Run:  python3 test_einv_adapter.py
"""
import datetime
import os
import shutil
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.environ["GST_EINV_CANONICAL"] = "0"          # original readers for the legacy side

import gst_core as mpu                        # noqa: E402
import gst_parsers_returns as pr              # noqa: E402
import gst_checks_monthly as mo               # noqa: E402
import einv_adapter as ad                     # noqa: E402

G1, G2, G3 = "07BBBBB0002B1Z2", "09CCCCC0003C1Z3", "29DDDDD0004D1Z4"
HDR = ["GSTIN/UIN of Recipient", "Reciever Name", "Invoice number", "Invoice date", "Invoice value", "Place of Supply",
       "Reverse Charge", "Applicable % of Tax Rate", "Invoice Type", "E-Commerce GSTIN", "Rate", "Taxable Value",
       "Integrated Tax", "Central Tax", "State/UT Tax", "Cess Amount", "IRN", "IRN date", "E-invoice status",
       "GSTR-1 auto-population/ deletion upon cancellation date", "GSTR-1 auto-population/ deletion status",
       "Error in auto-population/ deletion"]


def banner(fy, tp):
    return f"Financial Year: {fy}  |  Tax Period: {tp}  |  Date Updated till: 03-Aug-2026 11:57:25"


def inv(gstin, no, date, rate, taxable, igst=0.0, cgst=0.0, sgst=0.0, status="Valid", err=None, value=None):
    return [gstin, "Some Customer", no, date, value if value is not None else taxable + igst + cgst + sgst,
            "07 - Delhi", "N", None, "Regular B2B", None, rate, taxable, igst, cgst, sgst, 0.0, "irn" + str(no),
            date, status, "10-Apr-2025", "Auto-populated", err]


def make_workbook(path, blocks, hdr=None, extra_tabs=True, title_rows=3):
    """blocks = [(banner text, [rows])] written under the heading row of tab 'b2b, sez, de'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "b2b, sez, de"
    ws.append(["E-invoice details auto-populated for Form GSTR-1"])
    ws.append([])
    ws.append(["Taxable supplies made to registered taxpayers (invoices only) - B2B, SEZ, DE Invoices"])
    ws.append(hdr or HDR)
    for b, rows in blocks:
        ws.append([b])
        for r in rows:
            ws.append(list(r))
    if extra_tabs:
        for n in ("cdnr", "cdnur", "exp"):
            w = wb.create_sheet(n)
            w.append(["x"])
    wb.save(path)
    return path


def fixture(path, hdr=None):
    return make_workbook(path, [
        (banner("2025-26", "042025"), [
            inv(G1, "A/1", "09-Apr-2025", 18.0, 1000.0, igst=180.0),
            inv(G1, "A/2", "10-Apr-2025", 18.0, 500.0, cgst=45.0, sgst=45.0),
            inv(G2, "A/3", "11-Apr-2025", 12.0, 300.0, igst=36.0, status="Cancelled"),
            inv(G2, "A/3", "11-Apr-2025", 18.0, 200.0, igst=36.0, status="Cancelled"),   # 2nd rate line, same invoice
            inv(G3, "A/4", "12-Apr-2025", 5.0, 100.0, igst=5.0, err="Invalid GSTIN"),
            [None] * 22,                                                                    # blank spacer row
        ]),
        (banner("2025-26", "052025"), [
            inv(G1, "B/1", "03-May-2025", 18.0, 2500.0, igst=450.0),
            inv(G3, "B/2", datetime.datetime(2025, 5, 4), 28.0, 700.0, igst=196.0),          # real date cell
        ]),
        (banner("2025-26", "062025"), []),                                                   # month with no invoices
    ], hdr=hdr)


def check(name, cond, detail=""):
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f"  -- {detail}" if detail and not cond else ""))
    return cond


def run(fn, *a):
    try:
        return ("ok", fn(*a))
    except Exception as e:  # noqa: BLE001
        return ("EXC", type(e).__name__)


def _drop_invval(x):
    if isinstance(x, tuple) and x and x[0] == "ok" and isinstance(x[1], dict):
        return ("ok", {k: {a: b for a, b in v.items() if a != "invval"} for k, v in x[1].items()})
    return x


def compare_all(raw, months, ignore_invval=False):
    """Every reader, every month: original vs canonical. Returns the list of differences."""
    ad.clear_cache()
    diffs = []
    for m in months:
        for label, old, new in (("parse_einv", pr.parse_einv, ad.parse_einv),
                                ("read_einv_lines", mo.read_einv_lines, ad.read_einv_lines),
                                ("read_einv_invoices", mo.read_einv_invoices, ad.read_einv_invoices)):
            a, b = run(old, raw, m), run(new, raw, m)
            if ignore_invval and label == "read_einv_invoices":
                a, b = _drop_invval(a), _drop_invval(b)
            if a != b:
                diffs.append((label, m, str(a)[:160], str(b)[:160]))
    if mpu._einv_months_raw(raw) != ad.months(raw):
        diffs.append(("months", "", "", ""))
    return diffs


def _exc_text(fn, *a):
    try:
        fn(*a)
        return ""
    except Exception as e:  # noqa: BLE001
        return str(e)


def main():
    ok = True
    tmp = tempfile.mkdtemp()
    cwd = os.getcwd()
    os.chdir(tmp)
    try:
        raw = fixture(os.path.join(tmp, "einv.xlsx"))
        d = compare_all(raw, ["Apr-25", "May-25", "Jun-25", "Jul-25"], ignore_invval=True)
        ok &= check("every reader, every month identical to the original (Jun-25: banner without rows; Jul-25: not "
                    "covered)", not d, str(d[:2]))

        # ---- the awkward cases actually exercised
        e = ad.parse_einv(raw, "Apr-25")
        ok &= check("cancelled e-invoice lines are listed and left out of the totals",
                    len(e["cancelled"]) == 2 and e["count"] == 3 and abs(e["taxable"] - 1600.0) < 1e-9
                    and e["cancel_col_found"] and not e["cancel_date_col_found"], str({k: e[k] for k in ("count", "taxable")}))
        ok &= check("rows with an auto-population error are counted", e["errors"] == 1, str(e["errors"]))
        ok &= check("the portal's 'auto-population / deletion' date is NOT reported as a cancellation date",
                    all(c["cancel_date"] is None for c in e["cancelled"]))
        ok &= check("a month whose banner has no rows is 'available' with nothing in it",
                    ad.parse_einv(raw, "Jun-25")["available"] and ad.parse_einv(raw, "Jun-25")["count"] == 0)
        ok &= check("a month not in the file is reported as not covered",
                    not ad.parse_einv(raw, "Jul-25")["available"] and ad.read_einv_lines(raw, "Jul-25") == [])
        ln = ad.read_einv_lines(raw, "May-25")
        ok &= check("real date cells are read as dates", ln[1]["invdate"] == datetime.date(2025, 5, 4), str(ln[1]["invdate"]))

        # ---- deliberate fix: heading spelt 'Invoice value' (lower-case v) is now found
        inv_new = ad.read_einv_invoices(raw, "Apr-25")
        ok &= check("'Invoice value' (lower-case v) heading is read (the original looked for 'Invoice Value' only and read 0)",
                    abs(inv_new["A/1"]["invval"] - 1180.0) < 1e-9 and mo.read_einv_invoices(raw, "Apr-25")["A/1"]["invval"] == 0.0,
                    str(inv_new["A/1"]))

        # ---- canonical workbook on disk
        path = ad.write_canonical(ad.build_canonical(raw), tmp)
        ok &= check("canonical file is recognised", ad.is_canonical_file(path) and os.path.basename(path).startswith("Canonical_EINV_"))
        ok &= check("canonical file supplied as input reads identically",
                    ad.parse_einv(path, "Apr-25") == ad.parse_einv(raw, "Apr-25"))
        wb = openpyxl.load_workbook(path, read_only=True)
        ok &= check("sheets: META, RETURNS, COVERAGE, EINV_B2B, MAPPING_REPORT, ISSUES",
                    wb.sheetnames == ["META", "RETURNS", "COVERAGE", "EINV_B2B", "MAPPING_REPORT", "ISSUES"], str(wb.sheetnames))
        wb.close()
        p2 = ad.write_canonical(ad.build_canonical(raw), tmp)
        ok &= check("same source written twice -> same file (not '_2')", p2 == path)
        other = fixture(os.path.join(tmp, "einv_other.xlsx"))
        p3 = ad.write_canonical(ad.build_canonical(other), tmp)
        ok &= check("a different E-Invoice file of the same GSTIN/FY does not overwrite the first", p3 != path and os.path.exists(path))
        ok &= check("classification: canonical workbook recognised as E-Invoice",
                    mpu._einv_months(path) == {"Apr-25", "May-25", "Jun-25"}, str(mpu._einv_months(path)))

        # ---- numeric tax period, other heading spellings
        raw2 = make_workbook(os.path.join(tmp, "einv2.xlsx"), [
            (banner("2022-23", "042022"), [inv(G1, "X/1", "05-Apr-2022", 18.0, 400.0, igst=72.0)])],
            hdr=[h.replace("Invoice number", "Invoice Number").replace("E-invoice status", "IRN status") for h in HDR])
        d = compare_all(raw2, ["Apr-22"], ignore_invval=True)
        ok &= check("numeric MMYYYY tax period + 'Invoice Number' / 'IRN status' spellings identical", not d, str(d[:2]))

        # ---- broken inputs
        ad.clear_cache()
        hdr_bad = [h for h in HDR if h != "Central Tax"]
        bad = make_workbook(os.path.join(tmp, "bad_cols.xlsx"), [(banner("2025-26", "042025"), [])], hdr=hdr_bad)
        t = _exc_text(ad.build_canonical, bad)
        ok &= check("E504 required column missing -> plain message naming the column and the headings found",
                    "[E504]" in t and "Central tax" in t and "Headings found" in t, t[:200])
        nob = make_workbook(os.path.join(tmp, "no_banner.xlsx"), [])
        ok &= check("E503 no period banner", "[E503]" in _exc_text(ad.build_canonical, nob))
        wb = openpyxl.Workbook()
        wb.active.title = "Summary"
        wb.save("wrong.xlsx")
        ok &= check("E502 not an E-Invoice file", "[E502]" in _exc_text(ad.build_canonical, "wrong.xlsx"))
        with open("junk.xlsx", "w") as f:
            f.write("not an excel file")
        ok &= check("E501 unreadable file", "[E501]" in _exc_text(ad.build_canonical, "junk.xlsx"))
        stray = make_workbook(os.path.join(tmp, "stray.xlsx"), [])
        wb = openpyxl.load_workbook(stray)
        wb["b2b, sez, de"].append(inv(G1, "S/1", "01-Apr-2025", 18.0, 10.0, igst=1.8))
        wb["b2b, sez, de"].append([banner("2025-26", "042025")])
        wb.save(stray)
        ok &= check("E505 data row before any period banner", "[E505]" in _exc_text(ad.build_canonical, stray) or
                    "[E503]" in _exc_text(ad.build_canonical, stray))
        no_status = make_workbook(os.path.join(tmp, "no_status.xlsx"),
                                  [(banner("2025-26", "042025"), [inv(G1, "N/1", "01-Apr-2025", 18.0, 10.0, igst=1.8)])],
                                  hdr=[h for h in HDR if "status" not in h.lower()])
        ad.clear_cache()
        e = ad.parse_einv(no_status, "Apr-25")
        ok &= check("no status column: everything counted, cancel_col_found False (the check then says 'cannot test')",
                    e["count"] == 1 and not e["cancel_col_found"])
        ok &= check("I501 lists the missing optional columns and explains the cancelled-invoice consequence",
                    any(i["id"] == "I501" and "cancelled" in i["message"] for i in ad.get_data(no_status)["issues"]))

        # ---- quarterly banner: rows go to the month of their own invoice date
        q = make_workbook(os.path.join(tmp, "quarter.xlsx"), [
            (banner("2025-26", "Apr-Jun"), [
                inv(G1, "Q/1", "09-Apr-2025", 18.0, 100.0, igst=18.0),
                inv(G1, "Q/2", "10-May-2025", 18.0, 200.0, igst=36.0),
                inv(G1, "Q/3", datetime.datetime(2025, 6, 2), 18.0, 300.0, igst=54.0),
                inv(G1, "Q/4", "not a date", 18.0, 400.0, igst=72.0),
                inv(G1, "Q/5", "05-Jul-2025", 18.0, 500.0, igst=90.0)])])
        ad.clear_cache()
        got = {m: ad.parse_einv(q, m)["taxable"] for m in ("Apr-25", "May-25", "Jun-25")}
        ok &= check("quarterly banner: Apr=100, May=200, Jun=300 (the original copied the whole quarter into each month)",
                    got == {"Apr-25": 100.0, "May-25": 200.0, "Jun-25": 300.0}, str(got))
        issues = {i["id"] for i in ad.get_data(q)["issues"]}
        ok &= check("W501 (quarterly) and W502 (rows with no usable date in the quarter) reported",
                    {"W501", "W502"} <= issues, str(issues))
        ok &= check("legacy months() still lists the three months of the quarter", ad.months(q) == {"Apr-25", "May-25", "Jun-25"})
        old = run(pr.parse_einv, q, "May-25")
        ok &= check("(documenting the fixed bug) the original put the whole quarter in May", old[0] == "ok" and old[1]["taxable"] == 1500.0,
                    str(old)[:100])
    finally:
        os.chdir(cwd)
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nALL PASSED" if ok else "\nSOME TESTS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
