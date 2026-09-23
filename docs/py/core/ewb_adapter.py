#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
E-WAY BILL CANONICAL LAYER  (spec: docs/EWB_CANONICAL_SPEC.md)
================================================================
Only this module knows what a merged e-way bill workbook looks like.

    raw / merged EWB export  --build_canonical()-->  canonical EWB data
    canonical data           --write_canonical()-->  Canonical_EWB_<GSTIN>_<file-stem>.xlsx
    canonical .xlsx          --read_canonical()  -->  canonical data
    canonical data           --parse_annual_ewb() --> exactly what gst_parsers_returns.parse_annual_ewb() returns

Unlike every other source converted so far, an e-way bill export is ONE FLAT SHEET with a single header
row and no period markers at all -- 'month' is not a column, it is derived entirely from each row's OWN
'EWB No. & Dt.' text (e.g. '301546430758 - 10/01/2023 17:01:00'). So the canonical file keeps the raw
combined text columns exactly as the source has them (never pre-split), and the number/date/time split
happens in the reader, at read time, the same way the legacy reader always did it -- this is what makes the
result identical to the original, and it also means a portal wording change only ever needs a new heading
alias here, never a change to the parsing regex.

Problems are reported in plain words (ISSUES sheet, [warn]/[error] log lines). A file that cannot be used at
all (cannot be opened, doesn't have the e-way bill header signature, a required column missing) raises
PeriodParseError with the full message.
"""

import datetime as _dt
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

import canonical_common as cc
import gst_core as mpu

SCHEMA_VERSION = "EWB-1.0"
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "ewb_mapping.json")

ROWS_SHEET = "EWB_ROWS"
LEAD_COLS = ["src_row"]
MAP_COLS = ["canonical_field", "raw_header_found", "raw_column", "matched_by", "confidence", "status",
            "rows_filled", "rows_blank", "note"]
ISSUE_COLS = cc.ISSUE_COLS

_GSTIN_RE = re.compile(r"(\d{2}[A-Z]{5}\d{4}[A-Z]\dZ[A-Z\d])")
_MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

_CONTEXT = {"gstin": None}


def set_context(gstin=None):
    if gstin:
        _CONTEXT["gstin"] = gstin


def _mapping():
    return cc.load_json(_MAPPING_FILE)


def _fields():
    return _mapping()["fields"]


def _canon_cols():
    return LEAD_COLS + list(_fields()) + ["extra"]


# ----------------------------------------------------------------------------------------
# the two combined-text columns -> number / date / time (ported from
# gst_parsers_returns._split_ewb_no_dt, _gstin_of, _num_ewb so this module never has to import the
# module that dispatches TO it; kept byte-for-byte identical on purpose -- see that function's own
# docstring for why the 'time' group exists)
# ----------------------------------------------------------------------------------------
def _split_no_dt(v):
    """'351569103816 - 04/03/2023 11:43:00' -> (number, date, time)."""
    if not v:
        return "", None, None
    s = str(v)
    parts = s.split(" - ", 1)
    no = parts[0].strip()
    date = None
    time = None
    if len(parts) > 1:
        m = re.search(r"(\d{2})/(\d{2})/(\d{4})", parts[1])
        if m:
            dd, mm, yyyy = m.groups()
            date = _dt.date(int(yyyy), int(mm), int(dd))
        mt = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", parts[1])
        if mt:
            hh, mi, ss = mt.groups()
            try:
                time = _dt.time(int(hh), int(mi), int(ss or 0))
            except ValueError:
                time = None
    return no, date, time


def _gstin_of(cell):
    if not cell:
        return ""
    m = _GSTIN_RE.search(str(cell))
    return m.group(1) if m else str(cell).split("/")[0].strip()


def _num(v):
    if v is None:
        return 0.0
    s = str(v).replace(",", "").strip()
    if s in ("", "-", "#N/A", "NA"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _s(v):
    return str(v or "").strip()


# ----------------------------------------------------------------------------------------
# reading the raw workbook
# ----------------------------------------------------------------------------------------
def _resolve_field(hmap, spec):
    for pos, name in enumerate(spec.get("exact", [])):
        col = hmap.get(cc.clean(name))
        if col is not None:
            return col, ("exact" if pos == 0 else "alias"), "HIGH"
    return None, "", ""


def _find_data_sheet(wb, sig):
    """The sheet + header row whose first 3 rows carry the e-way bill signature headings
    ('EWB No.', 'From GSTIN & Name', 'To GSTIN & Name') -- same test gst_core.classify_folder uses to
    recognise an e-way bill file in the first place, so a file that gets this far is never rejected here
    for that reason."""
    sig_clean = {cc.clean(h) for h in sig}
    for sn in wb.sheetnames:
        ws = wb[sn]
        for ridx, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True)):
            hdr = [str(c).strip() if c else "" for c in row]
            if sig_clean <= {cc.clean(h) for h in hdr if h}:
                return sn, ridx, list(row)
    return None, None, None


def _fail(ctx, path, iid, msg, action="Check that this is the merged e-way bill workbook (inward_eway_bill_merged.xlsx / outward_eway_bill_merged.xlsx)."):
    ctx.fail(iid, f"E-Way Bill \"{os.path.basename(path)}\": {msg}", skipped="every check that reads this file",
             action=action)


def build_canonical(raw_paths):
    """Convert one merged e-way bill workbook (inward OR outward -- this module does not decide which;
    gst_core.classify_folder still does that from the parsed rows, exactly as before) into canonical data.
    A file that cannot be used raises PeriodParseError ('[E60x] ...')."""
    if isinstance(raw_paths, str):
        raw_paths = [raw_paths]
    path = raw_paths[0]
    ctx = cc.IssueLog("E-Way Bill", [os.path.basename(p) for p in raw_paths])
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        _fail(ctx, path, "E601", f"cannot open the file ({type(e).__name__}: {e}). It may be corrupt, "
              f"password-protected or not an Excel file.", "Re-download / re-merge the file.")
    fields = _fields()
    sig = _mapping()["signature_headers"]
    try:
        sheet, hdr_idx, header = _find_data_sheet(wb, sig)
        if sheet is None:
            _fail(ctx, path, "E602", f"no sheet found with the e-way bill column headings ({sig}). This does "
                  f"not look like a merged e-way bill export.")
        rows = [list(r) for r in wb[sheet].iter_rows(min_row=hdr_idx + 1, values_only=True)]
        unread = [n for n in wb.sheetnames if n != sheet]
    finally:
        wb.close()

    header = rows[0]
    hmap = {cc.clean(c): i for i, c in enumerate(header) if c not in (None, "")}
    texts = {i: str(c).strip() for i, c in enumerate(header) if c not in (None, "")}
    cols, mapping, missing_required = {}, [], []
    for f, spec in fields.items():
        col, how, conf = _resolve_field(hmap, spec)
        cols[f] = col
        if col is None and spec.get("required"):
            missing_required.append(f)
        mapping.append(dict(canonical_field=f, raw_header_found=texts.get(col, "") if col is not None else "",
                            raw_column=get_column_letter(col + 1) if col is not None else "", matched_by=how,
                            confidence=conf, status="OK" if col is not None else
                            ("NOT_IN_THIS_FILE" if spec.get("absent_ok") else "MISSING"),
                            rows_filled=0, rows_blank=0, note=""))
    if missing_required:
        names = [fields[f]["label"] for f in missing_required]
        _fail(ctx, path, "E603", f"required column(s) {names} not found on sheet '{sheet}'. Headings found: "
              f"{sorted(texts.values())}. Every check that reads this file is skipped for the run.",
              "Add the new heading to ewb_mapping.json, or supply the file as the merger script produced it.")
    optional_missing = [fields[f]["label"] for f, c in cols.items() if c is None and not fields[f].get("absent_ok")]
    if optional_missing:
        ctx.issue("I601", "INFO", sheet, ", ".join(optional_missing), "all",
                  f"E-Way Bill: column(s) {optional_missing} not found; read as blank/0 where a check uses them.")
    mapped = {c for c in cols.values() if c is not None}
    unmapped = {i: t for i, t in texts.items() if i not in mapped}

    table_rows, fill = [], {f: [0, 0] for f in fields}
    skipped_header_repeats, unparsable_dates = 0, 0
    for ridx, r in enumerate(rows[1:], start=hdr_idx + 2):
        if not any(c not in (None, "") for c in r):
            continue
        # A literal repeat of the header text as a data row (seen on real merges of files with slightly
        # different structure) -- dropped, same as the original reader.
        if str(r[0] or "").strip() in ("EWB No.", ""):   # exact literal, matching the original reader
            skipped_header_repeats += 1
            continue
        rec = dict(src_row=ridx)
        for f, spec in fields.items():
            col = cols[f]
            if col is None:
                rec[f] = None
                continue
            v = r[col] if col < len(r) else None
            fill[f][0 if v not in (None, "") else 1] += 1
            rec[f] = _num(v) if spec["type"] == "num" else v
        if cols.get("ewb_combined") is not None:
            _, d, _t = _split_no_dt(rec.get("ewb_combined"))
            if d is None:
                unparsable_dates += 1
        rec["extra"] = "; ".join(f"{t}={r[i]}" for i, t in unmapped.items() if i < len(r) and r[i] not in (None, ""))
        table_rows.append(rec)
    for m in mapping:
        if m["canonical_field"] in fill:
            m["rows_filled"], m["rows_blank"] = fill[m["canonical_field"]]
    if unparsable_dates:
        ctx.issue("W601", "WARNING", sheet, "ewb_combined", "all",
                  f"{unparsable_dates} row(s) have an 'EWB No. & Dt.' value that could not be read as a date, "
                  f"so they have no month and are excluded from every month-based check.",
                  "those rows in every month-based E-Way Bill check", "Check the row(s) in the source file.")
    if unread:
        ctx.issue("I602", "INFO", "", "", "all", f"Sheets not used: {unread}.")

    # FY / date range -- informational only (used for the canonical filename and META), never stored per
    # row: the canonical rows always keep the raw combined text, exactly as read.
    fys, dmin, dmax = set(), None, None
    for rec in table_rows:
        _, d, _t = _split_no_dt(rec.get("ewb_combined"))
        if d is None:
            continue
        dmin = d if dmin is None or d < dmin else dmin
        dmax = d if dmax is None or d > dmax else dmax
        fy_start = d.year if d.month >= 4 else d.year - 1
        fys.add(f"{fy_start}-{str(fy_start + 1)[2:]}")

    gstin = _CONTEXT["gstin"] or ""
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    meta = [("schema_version", SCHEMA_VERSION), ("recipient_gstin", gstin), ("fy", ", ".join(sorted(fys))),
            ("date_range", f"{dmin} to {dmax}" if dmin else ""), ("source_files", "; ".join(ctx.source_files)),
            ("source_layout", "flat sheet, one row per e-way bill, no period markers"),
            ("adapter_version", ADAPTER_VERSION), ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("row_count", str(len(table_rows))), ("header_repeat_rows_skipped", str(skipped_header_repeats)),
            ("unread_sheets", "; ".join(unread)),
            ("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK"))]
    return dict(meta=meta, rows=table_rows, mapping=mapping, issues=ctx.issues)


# ----------------------------------------------------------------------------------------
# write / read the canonical workbook
# ----------------------------------------------------------------------------------------
_WRITTEN = {}       # canonical path -> source files it was written from (this process)


def canonical_filename(data, source_stem="EWB"):
    m = dict(data["meta"])
    return f"Canonical_EWB_{m.get('recipient_gstin') or 'UNKNOWN'}_{source_stem}.xlsx"


def write_canonical(data, out_dir, source_path=None):
    """Write Canonical_EWB_<GSTIN>_<file-stem>.xlsx -- the file's own stem (inward_eway_bill_merged /
    outward_eway_bill_merged / ...) is part of the name, since (unlike every other source) there is no FY
    in the file itself to distinguish an inward export from an outward one, or one year's pooled file from
    another's, and two different files must never overwrite each other."""
    stem = os.path.splitext(os.path.basename(source_path))[0] if source_path else "EWB"
    name = canonical_filename(data, stem)
    src = dict(data["meta"]).get("source_files")
    base, ext = os.path.splitext(name)
    n = 1
    while True:
        cand = os.path.join(out_dir, name)
        if _WRITTEN.get(os.path.abspath(cand), src) == src:
            break
        n += 1
        name = f"{base}_{n}{ext}"
    sheets = [(ROWS_SHEET, _canon_cols(), data["rows"]), ("MAPPING_REPORT", MAP_COLS, data["mapping"]),
              ("ISSUES", ISSUE_COLS, data["issues"])]
    path = cc.write_workbook(out_dir, name, data["meta"], sheets, max_col_width=32)
    _WRITTEN[os.path.abspath(path)] = src
    return path


def is_canonical_file(path):
    return cc.is_canonical(path, "EWB-", ("META", ROWS_SHEET, "MAPPING_REPORT"))


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return dict(meta=cc.read_meta(wb), rows=cc.read_table(wb, ROWS_SHEET, _canon_cols()),
                    mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS),
                    issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------
# engine-facing access
# ----------------------------------------------------------------------------------------
def _index(data):
    return dict(meta=dict(data["meta"]), rows=data["rows"], _parsed=None)


def _build_fn(path):
    return build_canonical(path)


def _write_fn(data, out_dir):
    # CanonicalCache.get() calls this with (data, out_dir) only -- the source path is threaded through
    # via a closure set immediately before each get() call (see get_data()), so the canonical filename can
    # carry the file's own stem without changing the shared CanonicalCache signature.
    return write_canonical(data, out_dir, source_path=_CURRENT_PATH[0])


_CURRENT_PATH = [None]
_SRC = cc.CanonicalCache("E-Way Bill", is_canonical_file, read_canonical, _build_fn, _write_fn, _index)


def get_data(path, out_dir=None):
    _CURRENT_PATH[0] = path
    return _SRC.get(path, out_dir)


def clear_cache():
    _SRC.clear()
    _WRITTEN.clear()


def parse_annual_ewb(path):
    """Legacy gst_parsers_returns.parse_annual_ewb(path): list of dicts, one per e-way bill, with the
    number/date/time split out of the raw combined text exactly as the original reader did it."""
    ix = get_data(path)["_index"]
    if ix["_parsed"] is not None:
        return ix["_parsed"]
    out = []
    for r in ix["rows"]:
        ewbno, ewbdate, ewbtime = _split_no_dt(r.get("ewb_combined"))
        docno, docdate, _t = _split_no_dt(r.get("doc_combined"))
        docno = docno.strip().upper()
        month = f"{_MON[ewbdate.month - 1]}-{str(ewbdate.year)[2:]}" if ewbdate else None
        out.append(dict(
            ewbno=_s(r.get("ewbno_raw")) or ewbno, ewbdate=ewbdate, ewbtime=ewbtime, month=month,
            docno=docno, docdate=docdate,
            from_gstin=_gstin_of(r.get("from_combined")),
            from_name=str(r.get("from_combined") or "").split("/", 1)[-1].strip(),
            to_gstin=_gstin_of(r.get("to_combined")),
            to_name=str(r.get("to_combined") or "").split("/", 1)[-1].strip(),
            from_place=_s(r.get("from_place")), to_place=_s(r.get("to_place")),
            assess=r.get("assess") or 0.0, taxval=r.get("taxval") or 0.0,
            hsn=_s(r.get("hsn")), hsn_desc=_s(r.get("hsn_desc")), vehicle=_s(r.get("vehicle")),
            rate=r.get("rate") or 0.0,
        ))
    ix["_parsed"] = out
    return out


def filter_by_month(ewb_rows, month_key):
    """month_key e.g. 'Jan-23' -- matches on EWB date's month (not doc date). Pure in-memory filter, kept
    identical to gst_parsers_returns.filter_by_month (not file-reading, so it never needed converting)."""
    return [r for r in ewb_rows if r["month"] == month_key]
