#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GSTR-9 / GSTR-9C CANONICAL LAYER  (spec: docs/GSTR9_CANONICAL_SPEC.md)
========================================================================
Only this module knows what the government's GSTR-9 (annual return) and GSTR-9C (reconciliation statement) Excel exports look like.

    raw GSTR-9 / 9C  --build_canonical()-->  canonical data
    data             --write_canonical()-->  Canonical_GSTR9_<GSTIN>_<FY>.xlsx  /  Canonical_GSTR9C_<GSTIN>_<FY>.xlsx
    .xlsx            --read_canonical()  -->  data
    data             --parse_gstr9() / parse_gstr9c()  -->  exactly what gst_parsers_dept returns

Both exports are one sheet per Item / Part (names truncated to 31 characters). The original located every figure by its ROW LABEL and then took
the numbers from fixed columns of that row. The canonical file keeps, for every figure the engine uses (a "fact"), the matched row - its label, its
source sheet and row number, and the numbers of its first eight columns - plus the Part I label / value rows and the two kinds of list the engine
reads (Item 19 late-fee rows, the auditor's typed reason lines). The reader takes the same columns the original did and keeps its logic (sums,
the "system draft" test, the notes on what was not found).
"""

import datetime as _dt
import os

import openpyxl

import canonical_common as cc
import gst_core as mpu

SCHEMA = {"gstr9": "R9-1.0", "gstr9c": "R9C-1.0"}
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "gstr9_mapping.json")
NCOL = 8
KV_COLS = ["key", "src_row", "label", "value"]
FACT_COLS = ["fact", "sheet_key", "src_sheet", "src_row", "label"] + [f"c{i}" for i in range(NCOL)]
LIST_COLS = ["list_key", "src_row", "text", "n2", "n3"]
MAP_COLS = ["canonical_field", "raw_sheet", "raw_header_found", "raw_column", "matched_by", "confidence", "status", "src_row", "note"]
ISSUE_COLS = cc.ISSUE_COLS
FORMS = ("gstr9", "gstr9c")
# sheets / rows whose absence the original reported in `notes` (WARNING here); everything else is INFO
_NOTED_SHEETS = {"gstr9": {"part1", "item4", "item5", "item6", "item7", "item8", "item9", "part5"},
                 "gstr9c": {"part1", "item5", "item7", "item9", "item12"}}
_NOTED_FACTS = {"gstr9": {"t4_b2b", "t4_cn", "t6a", "t13"}, "gstr9c": set()}


def _forms():
    return cc.load_json(_MAPPING_FILE)["forms"]


def _num(v):
    """Legacy gst_parsers_dept._gxl_num: None stays None; blank / '-' / NA -> 0.0; unparsable text -> None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("₹", "").strip()
    if s in ("", "-", "–", "NA", "N/A"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return None


def _sheet(wb, prefixes):
    """First sheet whose name starts with any of the prefixes (legacy _gxl_sheet)."""
    for name in wb.sheetnames:
        for pfx in prefixes:
            if name.startswith(pfx):
                return name
    return None


def _row_by_label(rows, frags, occurrence=1, col_label=1, max_scan=60):
    """(1-based row number, row) of the n-th row within the first `max_scan` whose label column contains all `frags` (legacy _gxl_row_by_label)."""
    hits = 0
    for i, row in enumerate(rows[:max_scan], start=1):
        cell = row[col_label] if col_label < len(row) else None
        if cell is None:
            continue
        low = str(cell).lower()
        if all(f.lower() in low for f in frags):
            hits += 1
            if hits == occurrence:
                return i, row
    return None, None


def build_canonical(path, form):
    """Convert one GSTR-9 / GSTR-9C workbook into canonical data. A file that cannot be opened raises PeriodParseError ('[E1301] ...')."""
    if isinstance(path, (list, tuple)):
        path = path[0]
    fdef = _forms()[form]
    label = fdef["label"]
    ctx = cc.IssueLog(label, [os.path.basename(path)])
    try:
        wb = openpyxl.load_workbook(path, data_only=True)      # full (not read-only) load: government exports can carry a wrong <dimension>
    except Exception as e:  # noqa: BLE001
        ctx.fail("E1301", f"{label} \"{os.path.basename(path)}\": cannot open the file ({type(e).__name__}: {e}). It may be corrupt, "
                 f"password-protected or not an Excel file.", skipped=f"every check that reads the {label}", action="Re-download the file.")
    sheets, rows_of = {}, {}
    for skey, prefixes in fdef["sheets"].items():
        name = _sheet(wb, prefixes)
        sheets[skey] = name
        if name:
            rows_of[skey] = [list(r) for r in wb[name].iter_rows(values_only=True)]
    kv_rows, fact_rows, list_rows, mapping = [], [], [], []
    kv = fdef["kv"]
    if sheets.get(kv["sheet"]):
        for ridx, row in enumerate(rows_of[kv["sheet"]], start=1):
            if kv["label_col"] == 0:
                if not row or row[0] is None:
                    continue
                text = str(row[0]).strip()
                val = row[1] if len(row) > 1 else None
                low = text.lower()
                for key, rule in kv["rules"]:
                    if text.startswith(rule["startswith"]) and rule.get("contains", "") in low:
                        kv_rows.append(dict(key=key, src_row=ridx, label=text, value=val))
                        break
            else:
                if not row or len(row) < 2:
                    continue
                text = str(row[1] or "").strip().lower()
                val = row[2] if len(row) > 2 else None
                for key, rule in kv["rules"]:
                    if text == rule["equals"]:
                        kv_rows.append(dict(key=key, src_row=ridx, label=text, value=val))
                        break
    noted = _NOTED_SHEETS[form]
    for skey, name in sheets.items():
        if name is None:
            ctx.issue("W1302" if skey in noted else "I1302", "WARNING" if skey in noted else "INFO", "", skey, "all",
                      f"{label}: the sheet starting {fdef['sheets'][skey]} was not found in the workbook; the figures read from it are "
                      f"not available.")
    for fact, spec in fdef["facts"].items():
        name = sheets.get(spec["sheet"])
        rn, row = (None, None)
        if name:
            rn, row = _row_by_label(rows_of[spec["sheet"]], spec["frags"], spec.get("occurrence", 1))
        found = row is not None
        if found:
            rec = dict(fact=fact, sheet_key=spec["sheet"], src_sheet=name, src_row=rn, label=str(row[1]).strip())
            for i in range(NCOL):
                rec[f"c{i}"] = _num(row[i]) if i < len(row) else None
            fact_rows.append(rec)
        elif name and fact in _NOTED_FACTS[form]:
            ctx.issue("W1303", "WARNING", name, fact, "all",
                      f"{label} sheet '{name}': no row whose label contains {spec['frags']} was found in its first 60 rows, so that figure is not available.")
        mapping.append(dict(canonical_field=fact, raw_sheet=name or "", raw_header_found=(rec["label"] if found else " + ".join(spec["frags"])),
                            raw_column="B", matched_by="row label" if found else "", confidence="HIGH" if found else "",
                            status="OK" if found else ("MISSING" if name else "SHEET_MISSING"), src_row=rn or "", note=""))
    for lkey, spec in fdef["lists"].items():
        name = sheets.get(spec["sheet"])
        if not name:
            continue
        for ridx, r in enumerate(rows_of[spec["sheet"]], start=1):
            if spec["mode"] == "letters":
                if r and r[0] in spec["letters"]:
                    list_rows.append(dict(list_key=lkey, src_row=ridx, text=r[0], n2=_num(r[2]) if len(r) > 2 else None,
                                          n3=_num(r[3]) if len(r) > 3 else None))
            elif r and len(r) > 2 and r[2] and str(r[0] or "").strip().isalpha():
                list_rows.append(dict(list_key=lkey, src_row=ridx, text=str(r[2]).strip(), n2=None, n3=None))
    wb.close()
    if not any(sheets.values()):
        ctx.issue("E1304", "ERROR", "", "", "all", f"None of the {label} sheets was found: this does not look like a {label} Excel export.",
                  f"every check that reads the {label}", "Check that this is the government's Excel export.")
    kvd = {r["key"]: r["value"] for r in kv_rows}
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    meta = [("schema_version", SCHEMA[form]), ("form", form), ("recipient_gstin", str(kvd.get("gstin") or "")),
            ("legal_name", str(kvd.get("legal_name") or "")), ("fy", str(kvd.get("fy") or "")),
            ("source_files", "; ".join(ctx.source_files)), ("adapter_version", ADAPTER_VERSION),
            ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))]
    meta += [(f"sheet_{k}", v or "") for k, v in sheets.items()]
    meta.append(("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK")))
    return dict(meta=meta, kv=kv_rows, facts=fact_rows, lists=list_rows, mapping=mapping, issues=ctx.issues)


# ----------------------------------------------------------------------------------------
# write / read
# ----------------------------------------------------------------------------------------
_WRITTEN = {}


def write_canonical(data, out_dir):
    m = dict(data["meta"])
    name = cc.canonical_filename(_forms()[m["form"]]["canon_prefix"], data["meta"])
    src = m.get("source_files")
    base, ext = os.path.splitext(name)
    n = 1
    while _WRITTEN.get(os.path.abspath(os.path.join(out_dir, name)), src) != src:
        n += 1
        name = f"{base}_{n}{ext}"
    sheets = [("R9_KV", KV_COLS, data["kv"]), ("R9_FACTS", FACT_COLS, data["facts"]), ("R9_LISTS", LIST_COLS, data["lists"]),
              ("MAPPING_REPORT", MAP_COLS, data["mapping"]), ("ISSUES", ISSUE_COLS, data["issues"])]
    path = cc.write_workbook(out_dir, name, data["meta"], sheets, max_col_width=40)
    _WRITTEN[os.path.abspath(path)] = src
    return path


def is_canonical_file(path):
    return cc.is_canonical(path, "R9", ("META", "R9_FACTS", "R9_KV", "MAPPING_REPORT"))


def canonical_form(path):
    """'gstr9' / 'gstr9c' for a canonical workbook, else None."""
    if not is_canonical_file(path):
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return dict(cc.read_meta(wb)).get("form")
    finally:
        wb.close()


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return dict(meta=cc.read_meta(wb), kv=cc.read_table(wb, "R9_KV", KV_COLS), facts=cc.read_table(wb, "R9_FACTS", FACT_COLS),
                    lists=cc.read_table(wb, "R9_LISTS", LIST_COLS), mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS),
                    issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


def _index(d):
    meta = dict(d["meta"])
    lists = {}
    for r in d["lists"]:
        lists.setdefault(r["list_key"], []).append(r)
    return dict(meta=meta, facts={r["fact"]: r for r in d["facts"]}, lists=lists)


_SRC = {f: cc.CanonicalCache(_forms()[f]["label"], is_canonical_file, read_canonical, lambda p, f=f: build_canonical(p, f),
                             write_canonical, _index) for f in FORMS}


def get_data(path, form, out_dir=None):
    return _SRC[form].get(path, out_dir)


def clear_cache():
    for c in _SRC.values():
        c.clear()
    _WRITTEN.clear()


# ----------------------------------------------------------------------------------------
# engine-facing readers
# ----------------------------------------------------------------------------------------
def parse_gstr9(path):
    """Legacy gst_parsers_dept.parse_gstr9(path)."""
    out = dict(available=False, reason=None, is_system_draft=None, fy=None, gstin=None, legal_name=None,
               table4_b2b_taxable=None, table4_b2b_igst=None, table4_b2b_cgst=None, table4_b2b_sgst=None,
               table4_cn_taxable=None, table4_cn_igst=None, table4_cn_cgst=None, table4_cn_sgst=None,
               table5_zero_rated=None, table5_sez=None, table5_exempted=None, table5_nil_rated=None,
               table5_nongst=None, table5_all_zero=None,
               table6a_cgst=None, table6a_sgst=None, table6a_igst=None, table6a_cess=None, table6a_total=None,
               table9_liability_igst=None, table9_liability_cgst=None, table9_liability_sgst=None,
               table9_late_fee_payable=None, table9_late_fee_paid=None,
               table9_interest_payable=None, table9_interest_paid=None,
               notes=[],
               arn=None, date_of_filing=None,
               table4_rcm_inward_taxable=None, table4_net_supplies_tax_igst=None,
               table6_total_itc_availed=None, table7_total_reversed=None,
               table8_itc_as_per_2b=None, table8_itc_as_per_6b6h=None, table8_difference=None,
               table13_itc_cgst=None, table13_itc_sgst=None, table13_itc_igst=None, table13_itc_cess=None,
               table12_itc_reversed_cgst=None, table12_itc_reversed_sgst=None,
               table12_itc_reversed_igst=None, table12_itc_reversed_cess=None)
    if not path or not os.path.exists(path):
        out["reason"] = "GSTR-9 not supplied for this taxpayer/FY."
        return out
    try:
        data = get_data(path, "gstr9")
    except mpu.PeriodParseError as ex:
        out["reason"] = f"Could not read GSTR-9 workbook: {ex}"
        return out
    ix, meta = data["_index"], data["_index"]["meta"]
    facts, lists = ix["facts"], ix["lists"]
    has = lambda k: bool(meta.get(f"sheet_{k}"))

    def c(fact, i):
        r = facts.get(fact)
        v = r[f"c{i}"] if r else None
        return None if v is None else float(v)

    def s(fact, cols):
        return sum(c(fact, i) or 0 for i in cols)

    if has("part1"):
        for r in data["kv"]:
            out[r["key"]] = r["value"]
    else:
        out["notes"].append("Part I (Basic Details) sheet not found.")
    out["is_system_draft"] = not bool(out["arn"])
    if out["is_system_draft"]:
        out["notes"].append("No ARN found on Part I -- this may be the pre-filing 'System Drafted "
                            "(For Reference Only)' auto-draft rather than the as-filed return.")
    if has("item4"):
        if "t4_b2b" in facts:
            out["table4_b2b_taxable"], out["table4_b2b_cgst"], out["table4_b2b_sgst"], out["table4_b2b_igst"] = (
                c("t4_b2b", 2), c("t4_b2b", 3), c("t4_b2b", 4), c("t4_b2b", 5))
        else:
            out["notes"].append("Table 4B (B2B outward) row not found.")
        if "t4_cn" in facts:
            out["table4_cn_taxable"], out["table4_cn_cgst"], out["table4_cn_sgst"], out["table4_cn_igst"] = (
                c("t4_cn", 2), c("t4_cn", 3), c("t4_cn", 4), c("t4_cn", 5))
        else:
            out["notes"].append("Table 4I (credit notes) row not found.")
        if "t4_rcm" in facts:
            out["table4_rcm_inward_taxable"] = c("t4_rcm", 2)
        if "t4_net" in facts:
            out["table4_net_supplies_tax_igst"] = c("t4_net", 5)
    else:
        out["notes"].append("Item 4 sheet not found.")
    if has("item5"):
        t5 = {k: (c(f"t5_{k}", 2) if f"t5_{k}" in facts else None) for k in ("zero_rated", "sez", "exempted", "nil_rated", "nongst")}
        out["table5_zero_rated"], out["table5_sez"] = t5["zero_rated"], t5["sez"]
        out["table5_exempted"], out["table5_nil_rated"], out["table5_nongst"] = t5["exempted"], t5["nil_rated"], t5["nongst"]
        vals5 = [v for v in t5.values() if v is not None]
        out["table5_all_zero"] = bool(vals5) and all(abs(v) < 0.01 for v in vals5)
    else:
        out["notes"].append("Item 5 sheet not found.")
    if has("item6"):
        if "t6a" in facts:
            out["table6a_cgst"], out["table6a_sgst"], out["table6a_igst"], out["table6a_cess"] = (
                c("t6a", 3), c("t6a", 4), c("t6a", 5), c("t6a", 6))
            out["table6a_total"] = sum(v or 0 for v in (out["table6a_cgst"], out["table6a_sgst"], out["table6a_igst"], out["table6a_cess"]))
        else:
            out["notes"].append("Table 6A (ITC availed via 3B) row not found.")
        if "t6_total" in facts:
            out["table6_total_itc_availed"] = s("t6_total", (3, 4, 5, 6))
    else:
        out["notes"].append("Item 6 sheet not found.")
    if has("item7"):
        if "t7_total" in facts:
            out["table7_total_reversed"] = s("t7_total", (2, 3, 4, 5))
    else:
        out["notes"].append("Item 7 sheet not found.")
    if has("item8"):
        if "t8_2b" in facts:
            out["table8_itc_as_per_2b"] = s("t8_2b", (2, 3, 4, 5))
        if "t8_6b6h" in facts:
            out["table8_itc_as_per_6b6h"] = s("t8_6b6h", (2, 3, 4, 5))
        if "t8_diff" in facts:
            out["table8_difference"] = s("t8_diff", (2, 3, 4, 5))
    else:
        out["notes"].append("Item 8 sheet not found.")
    if has("item9"):
        for fact, key in (("t9_igst", "igst"), ("t9_cgst", "cgst"), ("t9_sgst", "sgst")):
            if fact in facts:
                out[f"table9_liability_{key}"] = c(fact, 2)
        if "t9_interest" in facts:
            out["table9_interest_payable"], out["table9_interest_paid"] = c("t9_interest", 2), c("t9_interest", 3)
        if "t9_latefee" in facts:
            out["table9_late_fee_payable"], out["table9_late_fee_paid"] = c("t9_latefee", 2), c("t9_latefee", 3)
    else:
        out["notes"].append("Item 9 sheet not found.")
    if has("part5"):
        if "t13" in facts:
            out["table13_itc_cgst"], out["table13_itc_sgst"], out["table13_itc_igst"], out["table13_itc_cess"] = (
                c("t13", 3), c("t13", 4), c("t13", 5), c("t13", 6))
        else:
            out["notes"].append("Table 13 (ITC availed for the previous financial year) row not found.")
        if "t12" in facts:
            out["table12_itc_reversed_cgst"], out["table12_itc_reversed_sgst"], \
                out["table12_itc_reversed_igst"], out["table12_itc_reversed_cess"] = (c("t12", 3), c("t12", 4), c("t12", 5), c("t12", 6))
    else:
        out["notes"].append("Part V (Items 10-14) sheet not found -- Table 13 'ITC availed for "
                            "the previous financial year' not available.")
    rows19 = lists.get("late_fee_19", [])
    if has("item19") and rows19:
        out["table9_late_fee_payable"] = sum(float(r["n2"] or 0) for r in rows19)
        out["table9_late_fee_paid"] = sum(float(r["n3"] or 0) for r in rows19)
    out["available"] = True
    return out


def parse_gstr9c(path):
    """Legacy gst_parsers_dept.parse_gstr9c(path)."""
    out = dict(available=False, reason=None, fy=None, gstin=None, legal_name=None, arn=None, arn_date=None,
               turnover_audited_bs=None, turnover_after_adjustments=None, turnover_declared_gstr9=None,
               turnover_unreconciled=None, exempt_nil_nongst_adjustment=None, taxable_turnover_after_adj=None,
               taxable_turnover_declared=None, itc_per_books=None, itc_declared_gstr9=None, itc_unreconciled=None,
               itc_booked_earlier_claimed_now=None, itc_booked_now_claimed_later=None, tax_payable_total=None,
               tax_paid_declared=None, notes=[],
               turnover_diff_reasons=[], taxable_turnover_diff_reasons=[], itc_diff_reasons=[])
    if not path or not os.path.exists(path):
        out["reason"] = "GSTR-9C not supplied for this taxpayer/FY."
        return out
    try:
        data = get_data(path, "gstr9c")
    except mpu.PeriodParseError as ex:
        out["reason"] = f"Could not read GSTR-9C workbook: {ex}"
        return out
    ix, meta = data["_index"], data["_index"]["meta"]
    facts, lists = ix["facts"], ix["lists"]
    has = lambda k: bool(meta.get(f"sheet_{k}"))

    def c(fact, i):
        r = facts.get(fact)
        v = r[f"c{i}"] if r else None
        return None if v is None else float(v)

    def s(fact, cols):
        return sum(c(fact, i) or 0 for i in cols)

    if has("part1"):
        for r in data["kv"]:
            out[r["key"]] = r["value"]
    else:
        out["notes"].append("Part I (Basic Details) sheet not found.")
    if has("item5"):
        out["turnover_audited_bs"] = c("c5_audited", 3)
        out["turnover_after_adjustments"] = c("c5_after", 3)
        out["turnover_declared_gstr9"] = c("c5_declared", 3)
        out["turnover_unreconciled"] = c("c5_unrec", 3)
    else:
        out["notes"].append("Item 5 sheet not found.")
    if has("item6"):
        out["turnover_diff_reasons"] = [r["text"] for r in lists.get("turnover_diff_reasons", [])]
    if has("item7"):
        out["taxable_turnover_after_adj"] = c("c7_taxable_adj", 2)
        out["exempt_nil_nongst_adjustment"] = c("c7_exempt", 2)
        out["taxable_turnover_declared"] = c("c7_declared", 2)
    else:
        out["notes"].append("Item 7 sheet not found.")
    if has("item8"):
        out["taxable_turnover_diff_reasons"] = [r["text"] for r in lists.get("taxable_turnover_diff_reasons", [])]
    if has("item9"):
        if "c9_payable" in facts:
            out["tax_payable_total"] = s("c9_payable", (3, 4, 5, 6))
        if "c9_paid" in facts:
            out["tax_paid_declared"] = s("c9_paid", (3, 4, 5, 6))
    else:
        out["notes"].append("Item 9 sheet not found.")
    if has("item12"):
        out["itc_per_books"] = c("c12_books", 2)
        out["itc_declared_gstr9"] = c("c12_declared", 2)
        out["itc_unreconciled"] = c("c12_unrec", 2)
        out["itc_booked_earlier_claimed_now"] = c("c12_earlier", 2)
        out["itc_booked_now_claimed_later"] = c("c12_later", 2)
    else:
        out["notes"].append("Item 12 sheet not found.")
    if has("item13"):
        out["itc_diff_reasons"] = [r["text"] for r in lists.get("itc_diff_reasons", [])]
    if not out["exempt_nil_nongst_adjustment"]:
        out["notes"].append("Table 7B (exempt/nil/non-GST adjustment) not found or zero -- Rule R13 "
                            "(turnover-gap check) needs this figure; verify manually if a gap is otherwise expected.")
    out["available"] = True
    return out
