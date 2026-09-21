#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GSTR-3B CANONICAL LAYER  (spec: docs/GSTR3B_CANONICAL_SPEC.md)
================================================================
Only this module knows what a raw GSTR-3B workbook looks like.

    raw / merged GSTR-3B   --build_canonical()-->  canonical 3B data
    canonical 3B data      --write_canonical()-->  Canonical_GSTR3B_<GSTIN>_<FY>.xlsx
    canonical .xlsx        --read_canonical()  -->  canonical 3B data
    canonical 3B data      --parse_gstr3b() / month_fields() / arn_dates_by_month() / months()
                            --> exactly what the four legacy readers return today

3B is a FORM (label on the left, numbers to the right), one sheet per month. Every line the engine
uses is found by its label text (rules in gstr3b_mapping.json), never by row/column position, and
each month's sheet is identified by its own 'Year' + 'Tax Period' rows, never by the sheet name.

ERROR-level problems are raised as gst_core.PeriodParseError with a full plain message: the engine's
per-month loop then reports "*** ERROR processing <month> ... SKIPPED" instead of quietly running on
zeros (the old behaviour whenever a required line was not found).
"""

import datetime as _dt
import json
import os
import re

import openpyxl

import canonical_common as cc
import gst_core as mpu

SCHEMA_VERSION = "3B-1.0"
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "gstr3b_mapping.json")

RETURN_COLS = ["month", "months_covered", "fy", "tax_period", "gstin", "legal_name", "trade_name", "arn",
               "arn_date", "generated_on", "src_sheet", "sheet_no", "used"]
LINE_COLS = ["month", "code", "description", "raw_label", "taxable", "igst", "cgst", "sgst", "cess", "n_values",
             "values_in_row_order", "src_sheet", "src_row"]
MAP_COLS = ["code", "description", "rule", "required", "months_found", "months_missing", "status", "note"]
ISSUE_COLS = cc.ISSUE_COLS

_NUMERIC_TEXT = re.compile(r"^-?[\d,\.]+$")
_DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y")
_CONTEXT = {"gstin": None}


# ----------------------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------------------
def set_context(gstin=None):
    if gstin:
        _CONTEXT["gstin"] = gstin


def _mapping():
    return cc.load_json(_MAPPING_FILE)


_clean, _num = cc.clean, cc.num


def _is_number_cell(c):
    """Same test as gst_parsers_returns.parse_gstr3b's vals_after."""
    return isinstance(c, (int, float)) or (isinstance(c, str) and
                                           _NUMERIC_TEXT.match(str(c).replace(",", "").strip()) is not None)


def _row_text(r):
    return " ".join(str(c) for c in r if c is not None)


def _parse_date(s):
    s = str(s or "").strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _Ctx(source_files):
    """Issues + mapping report for one build (shared implementation: canonical_common)."""
    return cc.IssueLog("GSTR-3B", source_files)


def _fail(ctx, iid, message, skipped="", action=""):
    ctx.fail(iid, message, skipped=skipped, action=action)


# ----------------------------------------------------------------------------------------
# reading one sheet
# ----------------------------------------------------------------------------------------
def _sheet_meta(rows):
    """{'year','tax_period','gstin',...} from the form's label/value rows. Same scan as the legacy
    readers: the first non-empty cell of a row is the label, the next one its value."""
    labels = _mapping()["meta_labels"]
    out = {}
    for row in rows:
        cells = [str(c).strip() for c in row if c not in (None, "")]
        if len(cells) < 2:
            continue
        key = cells[0].lower()
        for field, names in labels.items():
            if key in names and field not in out:
                out[field] = cells[1]
    return out


def _header_columns(rows_above):
    """{column index: canonical column name} from the nearest header row (a row with >= 2 of the
    known column headings such as 'Total Taxable value', 'Integrated Tax')."""
    known = set(_mapping()["header_cells"])
    for r in reversed(rows_above):
        found = {i: _clean(c) for i, c in enumerate(r) if c not in (None, "") and _clean(c) in known}
        if len(found) >= 2:
            names = {"totaltaxablevalue": "taxable", "integratedtax": "igst", "centraltax": "cgst",
                     "stateuttax": "sgst", "cess": "cess"}
            return {i: names[k] for i, k in found.items()}
    return {}


def _rule_matches(rule, text):
    if "all" in rule and not all(s in text for s in rule["all"]):
        return False
    if "any" in rule and not any(s in text for s in rule["any"]):
        return False
    if "none" in rule and any(s in text for s in rule["none"]):
        return False
    if "lower_all" in rule and not all(s.lower() in text.lower() for s in rule["lower_all"]):
        return False
    if "startswith" in rule and not text.strip().startswith(rule["startswith"]):
        return False
    return True


def _extract_lines(rows, sheet_name):
    """{code: line dict} for one sheet. Several matches for a code: the LAST wins (as the legacy
    reader did); the list of all matching rows is returned separately for a warning."""
    lines, hits = {}, {}
    rules = _mapping()["lines"]

    def make(code, r, ridx):
        nums = [_num(c) for c in r if _is_number_cell(c)]
        cols = _header_columns(rows[:ridx])
        named = {"taxable": None, "igst": None, "cgst": None, "sgst": None, "cess": None}
        for i, name in cols.items():
            if i < len(r) and r[i] is not None and _is_number_cell(r[i]):
                named[name] = _num(r[i])
        label = next((str(c).strip() for c in r if c not in (None, "") and not _is_number_cell(c)), "")
        return dict(code=code, description=rules[code]["desc"], raw_label=label, n_values=len(nums),
                    values_in_row_order=";".join(repr(x) for x in nums), src_sheet=sheet_name, src_row=ridx + 1,
                    **named)

    # section-anchored lines (4B): only between 'B. ITC Reversed' and 'C. Net ITC available'
    b = {}
    for code, rule in rules.items():
        if "section" not in rule:
            continue
        s, e = rule["section"]["start"], rule["section"]["end"]
        start = end = None
        for i, r in enumerate(rows):
            j = _row_text(r).strip()
            if j.startswith(s):
                start = i
            elif j.startswith(e) and start is not None:
                end = i
                break
        b[code] = (start, end)
    for code, rule in rules.items():
        if "section" not in rule:
            continue
        start, end = b[code]
        if start is None or end is None:
            continue
        for ridx in range(start + 1, end):
            r = rows[ridx]
            j = _row_text(r).strip()
            if not j or not any(_is_number_cell(c) for c in r):
                continue
            if any(j.startswith(p) for p in rule["row_startswith"]):
                # legacy order: '(2) Others' is tested before '(1)', so a '(2) Others' row is never 4B1
                if code == "4B1" and j.startswith("(2) Others"):
                    continue
                lines[code] = make(code, r, ridx)
                hits.setdefault(code, []).append(ridx + 1)
    # label-matched lines: the whole sheet, last match wins
    for ridx, r in enumerate(rows):
        text = _row_text(r)
        for code, rule in rules.items():
            if "section" in rule:
                continue
            if _rule_matches(rule, text):
                lines[code] = make(code, r, ridx)
                hits.setdefault(code, []).append(ridx + 1)
    return lines, hits


# ----------------------------------------------------------------------------------------
# build
# ----------------------------------------------------------------------------------------
def build_canonical(raw_paths):
    if isinstance(raw_paths, str):
        raw_paths = [raw_paths]
    ctx = _Ctx([os.path.basename(p) for p in raw_paths])
    returns, lines_out, not_3b = [], [], []
    seen_month = {}
    found_by_code = {}
    codes = list(_mapping()["lines"])
    for path in raw_paths:
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            _fail(ctx, "E304", f"GSTR-3B \"{os.path.basename(path)}\": cannot open the file "
                  f"({type(e).__name__}: {e}). It may be corrupt, password-protected or not an Excel file.",
                  "every check that reads GSTR-3B", "Re-download the file.")
        try:
            for sn in wb.sheetnames:
                rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
                meta = _sheet_meta(rows)
                if not (meta.get("year") and meta.get("tax_period")):
                    not_3b.append(f"{os.path.basename(path)}:{sn}")
                    continue
                try:
                    labels = mpu.months_for_tax_period(meta["year"], meta["tax_period"])
                except mpu.PeriodParseError as e:
                    ctx.issue("W305", "WARNING", sn, "tax_period", "",
                              f"Sheet '{sn}': Year '{meta['year']}' / Tax Period '{meta['tax_period']}' not "
                              f"recognised ({e}); the sheet is not used.", "that month", "")
                    not_3b.append(f"{os.path.basename(path)}:{sn}")
                    continue
                month = labels[0]
                seen_month[month] = seen_month.get(month, 0) + 1
                used = "Y" if seen_month[month] == 1 else "N"
                ret = dict(month=month, months_covered=";".join(labels), fy=meta["year"],
                           tax_period=meta["tax_period"], gstin=meta.get("gstin", ""),
                           legal_name=meta.get("legal_name", ""), trade_name=meta.get("trade_name", ""),
                           arn=meta.get("arn", ""), arn_date=None, generated_on=meta.get("generated_on", ""),
                           src_sheet=sn, sheet_no=seen_month[month], used=used)
                ad = meta.get("arn_date", "")
                ret["arn_date"] = _parse_date(ad) or (ad or None)
                if ad and not isinstance(ret["arn_date"], _dt.date):
                    ctx.issue("W306", "WARNING", sn, "arn_date", month,
                              f"Sheet '{sn}': 'Date of ARN' value {ad!r} is not a recognised date; filing-date "
                              f"checks for {month} cannot use it.", "late-fee / interest for that month", "")
                returns.append(ret)
                if used == "N":
                    ctx.issue("W301", "WARNING", sn, "", month,
                              f"Sheet '{sn}' is a second GSTR-3B for {month}; the first sheet "
                              f"('{[r['src_sheet'] for r in returns if r['month'] == month][0]}') is used and this "
                              f"one is ignored.", "", "Keep one GSTR-3B per period.")
                    continue
                lines, hits = _extract_lines(rows, sn)
                for code, h in hits.items():
                    if len(h) > 1:
                        ctx.issue("W302", "WARNING", sn, code, month,
                                  f"{month}: {code} matched {len(h)} rows {h} on sheet '{sn}'; the last is used.",
                                  "", "")
                for code in codes:
                    if code in lines:
                        found_by_code.setdefault(code, []).append(month)
                        lines_out.append(dict(month=month, **{k: v for k, v in lines[code].items()}))
        finally:
            wb.close()

    if not returns:
        _fail(ctx, "E301", f"GSTR-3B \"{', '.join(ctx.source_files)}\": no sheet has readable 'Year' and "
              f"'Tax Period' rows, so no month could be identified. Sheets seen: {not_3b}.",
              "every check that reads GSTR-3B", "Check that this is a GSTR-3B workbook as downloaded from the portal.")
    months = [r["month"] for r in returns if r["used"] == "Y"]
    rules = _mapping()["lines"]
    mapping = []
    for code in codes:
        rule = rules[code]
        found = found_by_code.get(code, [])
        missing = [m for m in months if m not in found]
        status = "OK" if not missing else ("MISSING" if not found else "PARTIAL")
        if missing and rule.get("absent_ok"):
            status = "NOT_IN_THIS_FORM" if not found else "PARTIAL"
        desc_rule = json.dumps({k: v for k, v in rule.items() if k not in ("desc", "required", "zero_default", "absent_ok")},
                               ensure_ascii=False)
        mapping.append(dict(code=code, description=rule["desc"], rule=desc_rule,
                            required="Y" if rule.get("required") else "N", months_found=len(found),
                            months_missing="; ".join(missing), status=status, note=""))
        if missing and rule.get("required"):
            ctx.issue("E302", "ERROR", "", code, "; ".join(missing),
                      f"Required GSTR-3B line {code} ({rule['desc']}) not found in {len(missing)} of "
                      f"{len(months)} month(s): {'; '.join(missing)}. Looked for a row matching "
                      f"{desc_rule}. Those months cannot be checked and are skipped.",
                      "all comparisons for those months",
                      "Give the 3B as downloaded from the portal, or add the new wording to gstr3b_mapping.json.")
        elif missing and rule.get("absent_ok") and not found:
            ctx.issue("I402", "INFO", "", code, "", f"Line {code} ({rule['desc']}) is not in this form version "
                      f"(expected: newer forms do not have it); treated as 0/absent.")
        elif missing:
            ctx.issue("W303", "WARNING", "", code, "; ".join(missing),
                      f"GSTR-3B line {code} ({rule['desc']}) not found in {len(missing)} of {len(months)} "
                      f"month(s): {'; '.join(missing)}. "
                      + ("Zeros are used for it (as before) - checks that read it are unreliable for those months."
                         if rule.get("zero_default") else "It is treated as absent for those months."),
                      "checks that read this line", "Add the new wording to gstr3b_mapping.json if the form changed.")
    if not_3b:
        ctx.issue("I401", "INFO", "", "", "", f"Sheets that are not a 3B return (no Year/Tax Period): {not_3b}.")
    fys = sorted({r["fy"] for r in returns})
    gstin = _CONTEXT["gstin"] or next((r["gstin"] for r in returns if r["gstin"]), "")
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    meta_rows = [("schema_version", SCHEMA_VERSION), ("recipient_gstin", gstin),
                 ("legal_name", next((r["legal_name"] for r in returns if r["legal_name"]), "")),
                 ("fy", ", ".join(fys)), ("source_files", "; ".join(ctx.source_files)),
                 ("source_layout", "one sheet per month"), ("adapter_version", ADAPTER_VERSION),
                 ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                 ("months_present", ";".join(sorted({m for r in returns for m in r["months_covered"].split(";")}))),
                 ("sheets_not_used", "; ".join(not_3b)),
                 ("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK"))]
    return dict(meta=meta_rows, returns=returns, lines=lines_out, mapping=mapping, issues=ctx.issues)


# ----------------------------------------------------------------------------------------
# write / read the canonical workbook
# ----------------------------------------------------------------------------------------
def canonical_filename(data):
    return cc.canonical_filename("GSTR3B", data["meta"])


def write_canonical(data, out_dir):
    return cc.write_workbook(out_dir, canonical_filename(data), data["meta"], [
        ("RETURNS", RETURN_COLS, data["returns"]), ("LINES", LINE_COLS, data["lines"]),
        ("MAPPING_REPORT", MAP_COLS, data["mapping"]), ("ISSUES", ISSUE_COLS, data["issues"])],
        meta_key_width=26, max_col_width=46)


def is_canonical_file(path):
    return cc.is_canonical(path, "3B-", ("META", "LINES", "RETURNS"))


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return dict(meta=cc.read_meta(wb), returns=cc.read_table(wb, "RETURNS", RETURN_COLS),
                    lines=cc.read_table(wb, "LINES", LINE_COLS), mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS),
                    issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------
# engine-facing readers
# ----------------------------------------------------------------------------------------
_SRC = cc.CanonicalCache("GSTR-3B", is_canonical_file, read_canonical, build_canonical, write_canonical,
                         lambda data: _index(data))


def get_data(path, out_dir=None):
    return _SRC.get(path, out_dir)


def clear_cache():
    _SRC.clear()


def _floats(text, n):
    return [float(x) for x in str(text).split(";")] if n else []


def _index(data):
    used = [r for r in data["returns"] if r["used"] == "Y"]
    by_month = {r["month"]: r for r in used}
    lines = {}
    for ln in data["lines"]:
        lines.setdefault(ln["month"], {})[ln["code"]] = ln
    codes = json.loads(json.dumps(_mapping()["lines"]))
    req = {c for c, r in codes.items() if r.get("required")}
    zero = {c for c, r in codes.items() if r.get("zero_default")}
    return dict(by_month=by_month, lines=lines, used=used, req=req, zero=zero,
                months={m for r in used for m in str(r["months_covered"]).split(";") if m})


def months(path):
    """Set of 'Mon-YY' months (legacy gst_core._gstr3b_months)."""
    return set(get_data(path)["_index"]["months"])


def gstin_and_name(path):
    d = get_data(path)
    first = next((r for r in d["_index"]["used"] if r["gstin"]), None)
    return ((first or {}).get("gstin") or None, (first or {}).get("legal_name") or None)


def parse_gstr3b(path, month):
    """Drop-in for gst_parsers_returns.parse_gstr3b(path, month)."""
    d = get_data(path)
    ix = d["_index"]
    ret = ix["by_month"].get(month)
    if ret is None:
        raise mpu.PeriodParseError(
            f"Month {month!r} not found as a GSTR-3B sheet in {path!r}. Months present: {sorted(ix['by_month'])}")
    lines = ix["lines"].get(month, {})
    missing_required = sorted(c for c in ix["req"] if c not in lines)
    if missing_required:
        msgs = [i["message"] for i in d["issues"] if i["id"] == "E302" and month in str(i["periods_affected"])]
        raise mpu.PeriodParseError(
            f"[E302] GSTR-3B {month}: required line(s) {missing_required} not found on sheet "
            f"'{ret['src_sheet']}'. " + (msgs[0] if msgs else ""))
    g = {}
    for code, ln in lines.items():
        if code == "4D175":
            continue
        g[code] = _floats(ln["values_in_row_order"], ln["n_values"])
    for code in ix["zero"]:
        g.setdefault(code, [0.0, 0.0, 0.0, 0.0])
    return g


def month_fields(path, month):
    """Drop-in for gst_checks_hsn_fraud._gstr3b_month_fields(path, month); None if the month is absent."""
    d = get_data(path)
    ix = d["_index"]
    ret = next((r for r in ix["used"] if month in str(r["months_covered"]).split(";")), None)
    if ret is None:
        return None
    lines = ix["lines"].get(ret["month"], {})

    def first(code):
        ln = lines.get(code)
        vals = _floats(ln["values_in_row_order"], ln["n_values"]) if ln else []
        return vals[0] if vals else 0.0

    fd = ret["arn_date"]
    filing = fd.date() if isinstance(fd, _dt.datetime) else (fd if isinstance(fd, _dt.date) else None)
    return dict(filing_date=filing,
                rcm_liability_igst=first("3.1d"),        # legacy quirk kept: the name says IGST, the value is the
                                                          # row's first number = taxable value (unused downstream)
                itc_all_other_igst=first("4A5"),
                itc_reversed_rule42_43=first("4B1"),     # 4(B)(1), whatever the form's wording (rules 42/43, or
                                                          # rules 38/42/43 + s.17(5) since Aug-2022)
                itc_reversed_others=first("4B2"),
                itc_ineligible_175=first("4D175"))


def arn_dates_by_month(path):
    """Drop-in for gst_checks_forensic.gstr3b_arn_dates_by_month(path)."""
    d = get_data(path)
    out = {}
    for r in d["_index"]["used"]:
        fd = r["arn_date"]
        date = fd.date() if isinstance(fd, _dt.datetime) else (fd if isinstance(fd, _dt.date) else None)
        for lbl in str(r["months_covered"]).split(";"):
            if lbl:
                out[lbl] = {"arn": (r["arn"] or None), "date": date}
    return out
