#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GST SOURCE TRACE
================
Final pass over the finished workbook (called from master_build.main just before save).
Adds a "Source & Formula" column to every sheet, and a "Source Map" index sheet, so any
number in the output can be re-checked by opening the named input file / sheet / row and
re-doing the named formula.

Read-only with respect to existing content: it only APPENDS one column per sheet (to the
right of the last used column) and one new sheet. No existing cell is changed. Any failure
on a sheet is swallowed for that sheet (the audit output must never be lost to a cosmetic
traceability column) and reported in the returned status string.

Input files are content-detected by the tool (file names vary per client), so locations are
given as "<document type> file -> sheet -> row/column".
"""
import re
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

COL_TITLE = "Source & Formula"

# ---- where each document type's numbers physically live -----------------------------------
G1B2B = "GSTR-1 file -> sheet 'b2b, sez, de_inv' (this month's block)"
G1 = "GSTR-1 file"
G3B = "GSTR-3B file -> this month's sheet"
EINV = "E-Invoice file -> month block"
G2B = "GSTR-2B file -> sheets b2b / b2b-cdnr (invoice-level; summed by 'ITC Availability')"

# Comparison sheet: key = (Left source label, Right source label) -> (left location, right location, formula)
_CMP_LOC = {
    "GSTR-1 net": f"[{G1B2B} + b2cl + exp + b2cs] taxable/tax columns, MINUS credit notes (sheets cdnr + cdnur)",
    "GSTR-1 gross": "same sheets as 'GSTR-1 net' but BEFORE deducting credit notes",
    "GSTR-1 inv net": f"GSTR-1 invoice sheets (b2b, b2cl, exp, b2cs) less cdnr/cdnur, summed from invoice lines ({G1})",
    "GSTR-1 HSN": f"{G1} -> sheet 'hsn' (Table 12), month block, summed over all HSN rows",
    "GSTR-1 (all rate-lines)": f"{G1B2B}: every rate-line incl. lines with blank invoice no.",
    "GSTR-1 B2B (named)": f"{G1B2B}: only lines WITH an invoice number",
    "GSTR-1 B2B named-inv count": f"{G1B2B}: count of distinct non-blank 'Invoice Number'",
    "GSTR-1 blank-invno lines": f"{G1B2B}: lines where 'Invoice Number' is blank",
    "E-Invoice": f"{EINV}: IRN rows for the month (B2B)",
    "E-Invoice (all rate-lines)": f"{EINV}: all rate-lines for the month",
    "GSTR-3B 3.1(a)": f"{G3B} -> row 'Outward taxable supplies (other than zero rated, nil rated and exempted)' (3.1(a))",
    "GSTR-3B 3.1(d)": f"{G3B} -> row 'Inward supplies liable to reverse charge' (3.1(d))",
    "GSTR-3B 4(A)(3)": f"{G3B} -> Table 4(A)(3) 'Inward supplies liable to reverse charge' ITC row",
    "GSTR-3B 4(A)(5)": f"{G3B} -> Table 4(A)(5) 'All other ITC' row",
    "GSTR-3B 4(B)(2)": f"{G3B} -> Table 4(B)(2) 'Others' (ITC reversed) row",
    "GSTR-2B (gross)": f"{G2B}: ITC-available invoices, credit notes NOT deducted",
    "GSTR-2B (net CN)": f"{G2B}: ITC-available invoices MINUS credit notes (b2b-cdnr)",
    "GSTR-2B CN": f"GSTR-2B file -> sheet b2b-cdnr: credit-note ITC for the month",
    "GSTR-2B RCM": f"{G2B}: only rows with 'Reverse Charge' = Y",
    "This month 4(B)(2)": f"{G3B} -> Table 4(B)(2)",
    "Trailing range (other months)": "Same 3B Table 4(B)(2) cell in the other months' sheets (min-max)",
}

# Comparison sheet: row-label keyword -> which column(s) the figure is taken from
_HEAD_COL = {"igst": "IGST", "cgst": "CGST", "sgst": "SGST", "cess": "CESS", "taxable": "Taxable value"}


def _head(label):
    low = (label or "").lower()
    for k in ("igst", "cgst", "sgst", "cess"):
        if re.search(rf"\b{k}\b", low):
            return _HEAD_COL[k]
    if "taxable" in low:
        return _HEAD_COL["taxable"]
    return None


def _comparison_text(check, lsrc, rsrc):
    left = _CMP_LOC.get(lsrc)
    right = _CMP_LOC.get(rsrc)
    if not left and not right:
        return None
    head = _head(check)
    col = f" [column: {head}]" if head else ""
    parts = []
    if left:
        parts.append(f"LEFT ({lsrc}) = {left}{col}")
    if right:
        parts.append(f"RIGHT ({rsrc}) = {right}{col}")
    parts.append("Difference = Left - Right (tolerance Rs 1.00; MATCH if |Difference| <= tolerance)")
    return " || ".join(parts)


# ---- Analysis (14-check engine) -----------------------------------------------------------
_ANALYSIS = {
    "#0": f"A: {G1B2B}, all lines vs named-invoice lines. Formula: gross = sum(all lines); net = gross - credit notes ('cdnr'). Informational totals.",
    "#1": f"A: {G1} -> sheet 'exemp' (Table 8) nil + exempt + non-GST. B: {G3B} -> 3.1(c) + 3.1(e). Formula: Diff = A - B.",
    "#2": f"A: {G1} -> sheet 'cdnr'/'cdnur' credit-note tax. B: {G3B} -> 3.1(a) tax. Formula: tax after CN = G1 gross - CN; compared with 3B 3.1(a).",
    "#3": f"{G3B} -> Table 4. Formula: Expected 4C = 4A(5) + 4A(3) - 4B(1) - 4B(2); Diff = reported 4C - Expected 4C.",
    "#4": f"A: {G1B2B} (tax / taxable). B: {G3B} -> 3.1(a) (tax / taxable). Formula: effective rate = tax / taxable x 100 on each side; flag if gap > 0.10 pp.",
    "#5": f"{G1B2B} vs {EINV}. Formula: GSTR-1 lines with no invoice no. are re-linked to an e-invoice by (value, rate, tax) composite key.",
    "#6": f"{G1B2B}. Formula: count of 'Invoice Number' values that appear on more than one line (multi-rate invoices are expected; same invoice on repeated identical lines is not).",
    "#7": f"{EINV} error / auto-population field vs GSTR-1 lines. Formula: per matched invoice, value or tax differences.",
    "#8": f"{EINV} IRN date vs GSTR-1 filing date (ARN marker row in the GSTR-1 file). Formula: lag = filing date - IRN date; flag if > 30 days.",
    "#9": f"{G1} -> sheet 'hsn' rate-wise rows vs {G1B2B} rates. Informational split.",
    "#10": f"GSTR-1 and GSTR-3B ARN filing dates (marker rows in each file). Formula: gap = GSTR-3B filing date - GSTR-1 filing date, flag if > 20 days.",
    "#11": f"{G1B2B}: 'Place Of Supply' vs supplier state (first 2 digits of the taxpayer GSTIN). Formula: different state => tax must be IGST; same state => CGST+SGST.",
    "#12": f"{G1B2B} 'Reverse Charge' = Y lines vs {G3B} -> 3.1(d) and 4(A)(3).",
    "#13": f"A: {G1} -> sheet 'hsn' IGST. B: {G1B2B} named-invoice IGST. Formula: Gap = HSN IGST - named-invoice IGST (after credit notes).",
    "#14": f"{G3B}: 4(A)(5) ITC vs 3.1(a) output tax. Formula: ratio = ITC / output tax; flag if > 95%.",
}

# ---- EWB matrix ---------------------------------------------------------------------------
_EWB = {
    "#1": "Outward E-Way-Bill file (doc no.) vs " + G1B2B + " invoice numbers. Formula: EWB-Out doc numbers not found in GSTR-1.",
    "#2": "Outward EWB assessable value vs " + G1B2B + " taxable value, per matched invoice. Formula: flag if |diff| > 1% of value.",
    "#3": "Outward EWB docs (tax-bearing) with no matching GSTR-1 invoice no.",
    "#4": G1B2B + " inter-state invoices > Rs 50,000 with no outward EWB (EWB_THRESHOLD = 50,000).",
    "#5": "Outward EWB vs E-Invoice IRN list (matched on document number).",
    "#6": "Outward EWB value vs E-Invoice value, matched documents.",
    "#10": "Inward EWB file vs GSTR-2B invoice-level b2b: invoice-number match.",
    "#12": "Inward EWB (tax-bearing) with no matching GSTR-2B invoice (see 'EWB Full-Year Reconciliation' for the FY-wide re-match).",
    "#13": "GSTR-2B inter-state invoices > Rs 50,000 vs Inward EWB list.",
    "#15": "Same document number present in both Outward EWB and Inward EWB (self-transaction).",
    "#16": "EWB generation date vs document date in the EWB file. Formula: lag = generation date - document date.",
    "#17": "Three-way match: " + G1B2B + " vs E-Invoice vs Outward EWB on invoice number.",
    "#18": "EWB-Out HSN-wise assessable value vs GSTR-1 'hsn' sheet and E-Invoice.",
    "#22": "EWB validity/expiry fields vs supply dates (needs validity column in the EWB export).",
    "#24": "EWB cancel date vs return filing date (needs cancel column in the EWB export).",
    "#25": "Outward EWB assessable total / " + G1B2B + " taxable total. Formula: ratio = EWB value / GSTR-1 B2B taxable (services HSN carved out).",
    "#26": "Inward EWB assessable total / GSTR-2B b2b taxable total (ratio).",
    "#27": "Outward EWB: same vehicle number repeated between the same GSTIN pair.",
    "#28": "Inward EWB documents matched to GSTR-2B supplier (supplier GSTIN present in 2B).",
}

# ---- Sheet-level descriptions (default per sheet, regex on sheet name) --------------------
# (pattern, source files, formula / logic, code location)
_SHEETS = [
    (r"^Master Dashboard$", "Every other sheet in this workbook (no input file read directly)",
     "Collects every FLAG / MISMATCH / REVIEW row from Comparison, Analysis, EWB, Doc-Series, HSN & Fraud, Forensic, Cancelled E-Inv and Flow sheets. Severity band from the rupee amount; PARENT / VIEW / ROLLUP level removes double counting. Column 'Source Reference' names the originating sheet + check ref.",
     "master_build.write_master_dashboard"),
    (r"Comparison$", "GSTR-1, GSTR-3B, E-Invoice, GSTR-2B files",
     "Left value - Right value = Difference, row by row. The exact cells for each side are in this column.",
     "gst_parsers_returns.build_comparisons / gst_report.write_comparison"),
    (r"Analysis$", "GSTR-1, GSTR-3B, E-Invoice files", "14-check engine; see the row's Ref.", "gst_checks_monthly.run_checks"),
    (r"EWB$", "Outward / Inward E-Way-Bill files, GSTR-1, E-Invoice, GSTR-2B", "28-check E-Way-Bill matrix; see the row's Ref.", "gst_checks_monthly (eway)"),
    (r"EWB Detail$", "Outward / Inward E-Way-Bill files",
     "Row-level detail behind the matching EWB sheet's findings; each row is one EWB document from the EWB file (same month).", "gst_checks_monthly (eway)"),
    (r"^Doc-Series Integrity$", "GSTR-1 file -> sheet 'doc_issue' (Table 13), GSTR-1 b2b invoice numbers, Cancelled E-Invoices, B2CS",
     "Table-13 declared range -> list every serial in the range -> missing = declared serial not present in GSTR-1 invoices; extra = invoice beyond declared range; gaps explained by cancelled e-invoices / B2CS before being flagged.",
     "master_build.write_doc_series"),
    (r"^Rectification Pairs$", "GSTR-1 file -> sheets 'b2ba', 'cdnra' (amendments) and b2b of earlier months",
     "Each amendment row is paired with the earlier month's original invoice. Delta taxable = revised - original; Delta tax = revised tax - original tax.",
     "master_build.build_rectification_pairs"),
    (r"^CN-DN ITC Impact", "GSTR-2B file (b2b-cdnr credit/debit notes), GSTR-3B 4(B)(2)", "ITC effect of credit / debit notes by month and FY.", "gst_checks_flow.build_cn_dn_impact_data"),
    (r"^HSN & Fraud Pattern Checks$", "GSTR-1 'hsn' sheet, GSTR-1 b2b, E-Invoice, EWB, GSTR-2B, HSN master tables",
     "Numbered pattern checks (Ref column); each row's Detail names the invoice / HSN / month. Rate checks use the date-versioned rate table in gst_checks_hsn_fraud.",
     "gst_checks_hsn_fraud.run_all"),
    (r"^Filing Compliance", "GSTR-1 and GSTR-3B ARN filing dates (marker rows), statutory due dates",
     "Late days = filing date - due date. Late fee (s.47) per day, nil-return rate lower, capped per notification; interest (s.50) = net cash tax x 18% x late days / 365.",
     "gst_checks_forensic.month_filing_compliance"),
    (r"^Forensic Checks", "GSTR-9, GSTR-9C, Table 8A, GSTR-1 'exemp', BS/P&L input",
     "R13: 9C Table 7B exempt/nil adjustment vs GSTR-1 Table 8. R14: four-way ITC: GSTR-9 6A vs GSTR-2B FY total vs Table 8A vs GSTR-9C 12A.",
     "gst_checks_forensic"),
    (r"^Cancelled E-Invoices$", "E-Invoice file (cancel column), GSTR-1 b2b, Outward EWB",
     "D2a: cancelled IRN that still has a live GSTR-1 b2b row. D2b: cancelled IRN that still has a live outward EWB.", "gst_checks_forensic.build_cancelled_einvoice_findings"),
    (r"^EWB Pattern - Annual$", "Month EWB sheets (#13 / #15 / #17 / #27 findings)", "Consolidated per subject: months fired out of 12, total value. Column 'Source (month tabs)' names the monthly sheets.", "master_build.write_ewb_pattern_annual"),
    (r"^IRN Late-Gen - Annual$", "E-Invoice file (IRN date), invoice date, Outward EWB date", "Delay days = IRN date - invoice date; EWB-before-IRN timing check.", "master_build.write_irn_late_annual"),
    (r"^ITC-Blocked Inv - Annual$", "GSTR-2B file ('ITC Availability' = No)", "Invoices GSTR-2B itself marks as ITC not available; occurrences per supplier.", "master_build.write_itc_blocked_annual"),
    (r"^Round-Number Inv - Annual$", "GSTR-1 b2b", "Lines where taxable value AND tax are both exact round numbers.", "master_build.write_round_number_annual"),
    (r"^HSN Timeline - Annual$", "GSTR-1 'hsn' sheet by month", "First month an HSN appears (new code mid-year); month-on-month HSN share change >= 10 pp.", "master_build.write_hsn_timeline_annual"),
    (r"^Annual Ledger Walkthrough$", "Cash Ledger, Credit Ledger, Liability Register, Portal Comparison files",
     "Per month: Credit ledger opening + availed - utilised = closing; cash deposits vs debits; liability declared vs paid; portal figure vs filed figure.", "gst_report.build_monthly_rows / write_monthly"),
    (r"^FY Total vs BIFA$", "Ledger files (above) + BO/360 Profile 'Financial Information'", "FY total from the ledgers vs the BO Profile figure (BO Profile values are in LAKHS, x 100,000 to compare).", "gst_report.write_fy_total_vs_bifa"),
    (r"^Related-Party Alerts$", "BO/360 Profile (related-party ITC sheets)", "Copied from the BO Profile; no recomputation.", "gst_report.write_related_party"),
    (r"^Top Counterparties$", "BO/360 Profile (top suppliers / buyers)", "Copied from the BO Profile; no recomputation.", "gst_report.write_top_counterparties"),
    (r"^Purchase vs Sales & Stock$", "GSTR-2B (purchases), GSTR-1 (sales), BS/P&L input", "Implied inventory change = purchases - sales, compared with the BS/P&L inventory movement.", "gst_checks_flow (F1)"),
    (r"^ITC (Yearly|Annual) Summary$", "GSTR-3B Table 4, Credit Ledger, GSTR-2B", "Per-month ITC claimed / reversed / net; 4(B)(1) vs 4(B)(2) split; month-on-month outliers (>2x or <0.5x).", "gst_checks_flow (F2-F7)"),
    (r"^GSTR-2B ITC No & Unconfirmed$", "GSTR-2B file", "Invoices with ITC Availability = No / Unconfirmed; 'Source' column names the origin.", "gst_checks_flow"),
    (r"^Zero-Tax Invoice Scan$", "GSTR-1 b2b / GSTR-2B", "Invoices with taxable value but zero tax.", "gst_checks_flow"),
    (r"^ITC Roll-Forward", "GSTR-3B Table 4 (4A, 4B, 4C) by month", "Opening + claimed - reversed = closing, month by month.", "gst_checks_flow"),
    (r"^Credit Reversal Recon$", "GSTR-3B 4(B), Credit Ledger, Liability Register, DRC-03", "Reversal / reclaim reconciliation across the sources.", "gst_credit_reversal"),
    (r"^3-Way GSTR1 3B EWB$", "GSTR-1 taxable, GSTR-3B 3.1(a), de-duplicated Outward EWB value", "Three values side by side per month; gaps = difference between each pair.", "gst_checks_flow (F4/F5)"),
    (r"^ITC 3B vs 2B$", "GSTR-3B 4(A)(5), GSTR-2B ITC-available", "Diff = 3B claimed - 2B available, FY-wide and by month.", "gst_checks_flow (F7)"),
    (r"^RCM Triangulation$", "GSTR-3B 3.1(d) and 4(A)(3), Cash Ledger RCM debits", "3.1(d) liability vs 4(A)(3) ITC vs cash-ledger RCM debits (s.49(4)).", "gst_checks_flow (F8)"),
    (r"^DRC-03 & Ledger Movements$", "Liability Register, Cash Ledger, BO Profile", "DRC-03 / voluntary payments, refund debits, Rule 86A blocks matched to the BO Profile. 'Source' column names the origin.", "gst_checks_flow (F9)"),
    (r"^Turnover Growth vs Tax$", "GSTR-1 / GSTR-3B turnover, Cash Ledger", "Turnover growth rate vs cash tax share; persistent low cash share is flagged.", "gst_checks_flow (F10)"),
    (r"^Counterparty Transactions$", "GSTR-1 b2b, GSTR-2B b2b", "Same-day repeats; counterparties who are both buyer and seller.", "gst_checks_flow (F11)"),
    (r"^Top Counterparties \(Computed\)$", "GSTR-1 b2b, GSTR-2B b2b; BO Profile for cross-reference", "Tool's own top-10 by value, compared with the BO Profile's top-10.", "gst_checks_flow (F12)"),
    (r"^B2B to B2C Shift$", "GSTR-1 b2b vs b2cs / b2cl", "B2B share this month - B2B share last month; flag if >= 10 pp.", "gst_checks_flow (F6)"),
    (r"^GSTR-2A ", "GSTR-2A file, GSTR-2B file, GSTR-3B, Credit Ledger", "2A vs 2B invoice match, 2A vs 3B / ledger, supplier filing dates, cancelled-supplier ITC.", "gst_checks_flow (G1-G10)"),
    (r"^Potential Blocked Credits$", "GSTR-2B b2b supplier names + blocked-ITC keyword master", "Supplier name matched against s.17(5) keywords; confidence capped at Medium (GSTR-2B has no HSN).", "gst_blocked_credit.build_and_write"),
    (r"^Machinery HSN Scan$", "GSTR-1 'hsn', Outward / Inward EWB, GSTR-2B supplier names, machinery HSN master", "4-digit HSN prefix match; purchase side also uses supplier-name keywords.", "gst_machinery_scan.build_and_write"),
    (r"^QA Summary$", "Dashboard rows", "Counts and rupee exposure of the QA layer.", "master_build.write_qa_summary"),
    (r"^Action Required$", "Dashboard rows", "Only PARENT-level FLAG / MISMATCH rows with severity >= Medium.", "master_build.write_qa_action_required"),
    (r"^EWB Full-Year Reconciliation$", "Inward EWB, GSTR-2B (whole FY)", "Re-matches EWB-In gaps (#12) against the whole-FY GSTR-2B index, separating timing gaps from true absences.", "master_build.write_qa_ewb_reconciliation"),
    (r"^Reviewed Master Dashboard$", "Master Dashboard rows", "Every dashboard row plus a QA status from the QA_REF_RULES lookup.", "master_build.write_qa_reviewed_dashboard"),
]

_MAXLEN = 1500


def _sheet_meta(name):
    for pat, src, how, code in _SHEETS:
        if re.search(pat, name):
            return src, how, code
    return None


def _find_header(ws, scan=15):
    """Header row = first row (within `scan`) with >=3 non-empty string cells."""
    for r in range(1, min(ws.max_row, scan) + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if sum(1 for v in vals if isinstance(v, str) and v.strip()) >= 3:
            return r
    return None


def _clip(s):
    return s if len(s) <= _MAXLEN else s[:_MAXLEN - 3] + "..."


_BIG = 2000   # sheets with more data rows than this get ONE shared note instead of a per-row copy


def _annotate_rows(ws, name, hdr, col):
    """Returns the number of cells written. ws.max_row / ws.max_column / iter_rows are O(cells)
    per call in openpyxl, so the cell dict is read once instead (otherwise the 62k-row
    counterparty sheet takes minutes)."""
    base = _sheet_meta(name)
    default = None
    if base:
        src, how, code = base
        default = f"Source: {src}. Logic: {how} Code: {code}."
    is_cmp = name.endswith("Comparison")
    is_ana = name.endswith("Analysis")
    is_ewb = name.endswith("EWB")
    grid = {}
    for (r, c), cell in ws._cells.items():
        if r > hdr and c < col and cell.value is not None:
            grid.setdefault(r, {})[c] = cell.value
    last = max(grid) if grid else hdr
    big = (last - hdr) > _BIG
    writes = []
    in_hsn = False
    for r in range(hdr + 1, last + 1):
        g = grid.get(r)
        if not g:
            continue
        text = None
        if is_cmp:
            if str(g.get(1) or "").startswith("HSN RATE REVIEW"):
                in_hsn = True
            if in_hsn:
                text = (f"{G1} -> sheet 'hsn' (Table 12), this month's block: one row per HSN x rate; "
                        "rate compared with the date-versioned HSN rate table (gst_checks_hsn_fraud).")
            else:
                text = _comparison_text(str(g.get(2) or ""), g.get(3), g.get(5))
        elif is_ana or is_ewb:
            m = re.match(r"(#\d+)", str(g.get(1) or "").strip())
            if m:
                text = (_ANALYSIS if is_ana else _EWB).get(m.group(1))
        if text is None and default and not big and not (is_cmp and len(g) == 1):
            text = default
        if text:
            writes.append((r, text))
    if big and default:
        writes.append((hdr + 1, "APPLIES TO EVERY ROW BELOW -- " + default))
    for r, text in writes:
        ws.cell(r, col, _clip(text)).alignment = Alignment(wrap_text=True, vertical="top")
    return len(writes)


def add_source_trace(wb):
    """Append the column + the Source Map sheet. Returns a one-line status string."""
    done, failed = 0, []
    for ws in list(wb.worksheets):
        try:
            if ws.max_row < 2 or ws.max_column < 1:
                continue
            hdr = _find_header(ws)
            if hdr is None:
                continue
            col = ws.max_column + 1
            h = ws.cell(hdr, col, COL_TITLE)
            h.font = Font(bold=True, color="FFFFFF")
            h.fill = PatternFill("solid", fgColor="305496")
            h.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = 70
            _annotate_rows(ws, ws.title, hdr, col)
            done += 1
        except Exception as ex:  # never lose the audit output to a traceability column
            failed.append(f"{ws.title}: {ex!r}")

    try:
        _write_source_map(wb)
    except Exception as ex:
        failed.append(f"Source Map: {ex!r}")
    msg = f"Source & Formula column added to {done} sheet(s); 'Source Map' sheet written."
    if failed:
        msg += f" {len(failed)} problem(s): " + "; ".join(failed[:5])
    return msg


def _write_source_map(wb):
    ws = wb.create_sheet("Source Map", 1)
    ws["A1"] = "SOURCE MAP  -  where every sheet's numbers come from"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Each sheet also has a 'Source & Formula' column on its last column with the row-level trace. "
                "Input files are content-detected, so they are named by document type. Within the merged GSTR-1 / E-Invoice / GSTR-2B files each month is a block located by its 'Tax Period' marker row; GSTR-3B has one sheet per month located by its in-sheet 'Tax Period' row.")
    hdr = ["Sheet", "Input file(s)", "Formula / logic", "Code location"]
    for i, t in enumerate(hdr, 1):
        c = ws.cell(4, i, t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    r = 5
    seen = set()
    for sheet in wb.sheetnames:
        if sheet == "Source Map":
            continue
        meta = _sheet_meta(sheet)
        # monthly sheets repeat x12 -> list each kind once
        key = re.sub(r"^[A-Z][a-z]{2}-\d{2} ", "<Month> ", sheet)
        if key in seen:
            continue
        seen.add(key)
        src, how, code = meta if meta else ("(see sheet)", "(no entry in gst_source_trace._SHEETS)", "")
        for i, v in enumerate((key, src, how, code), 1):
            c = ws.cell(r, i, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    for letter, w in zip("ABCD", (34, 55, 90, 42)):
        ws.column_dimensions[letter].width = w
