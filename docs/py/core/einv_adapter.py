#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
E-INVOICE CANONICAL LAYER  (spec: docs/EINV_CANONICAL_SPEC.md)
===============================================================
Only this module knows what a raw (merged) E-Invoice workbook looks like.

    raw / merged E-Invoice   --build_canonical()-->  canonical E-Invoice data
    canonical data           --write_canonical()-->  Canonical_EINV_<GSTIN>_<FY>.xlsx
    canonical .xlsx          --read_canonical()  -->  canonical data
    canonical data           --parse_einv() / read_einv_lines() / read_einv_invoices() / months()
                             --> exactly what the three legacy E-Invoice readers return today

The engine only ever reads the 'b2b, sez, de' tab (the cdnr / cdnur / exp tabs are not used by any check).
That tab has a title, ONE heading row, a period banner per month and that month's invoices under it. Each
canonical row keeps the source row with canonical column names and the month it belongs to; numbers are
converted (blank = 0), identifiers, text and dates are kept exactly as in the source.

A quarterly banner ("Tax Period: Apr-Jun") cannot say which month an invoice belongs to, so each row of such a
banner is placed in the month of its OWN invoice date (IRN date as a fallback); rows with no usable date are
left out of every month and reported (W502). The original reader copied the whole quarter into every month.

Problems are reported in plain words (ISSUES sheet, [warn]/[error] log lines). A file that cannot be used at all
(cannot be opened, no 'b2b, sez, de' tab, required column missing, no period banner) raises PeriodParseError with
the full message, which the engine turns into "E-Invoice file ... could not be read: <message>".
"""

import datetime as _dt
import os

import openpyxl
from openpyxl.utils import get_column_letter

import canonical_common as cc
import gst_core as mpu

SCHEMA_VERSION = "EINV-1.0"
ADAPTER_VERSION = "1.0.0"
_HERE = os.path.dirname(os.path.abspath(__file__))
_MAPPING_FILE = os.path.join(_HERE, "einv_mapping.json")

TABLE = "b2b, sez, de"          # the raw tab (and the COVERAGE table name)
LINES_SHEET = "EINV_B2B"        # the canonical sheet
LEAD_COLS = ["period", "months_covered", "src_row", "_any"]
RETURN_COLS = ["month", "fy", "tax_period", "marker_text"]
COVERAGE_COLS = ["table", "month", "rows", "note"]
MAP_COLS = ["canonical_sheet", "canonical_field", "raw_sheet", "raw_header_found", "raw_column", "matched_by",
            "confidence", "status", "rows_filled", "rows_blank", "note"]
ISSUE_COLS = cc.ISSUE_COLS

_DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y", "%d/%m/%y")
_MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

_CONTEXT = {"gstin": None}


def set_context(gstin=None):
    if gstin:
        _CONTEXT["gstin"] = gstin


def _mapping():
    return cc.load_json(_MAPPING_FILE)


def _fields():
    return _mapping()["fields"]


def _canon_cols():
    return LEAD_COLS + list(_fields()) + ["extra"]


def _label(labels):
    return labels[0] if len(labels) == 1 else f"{labels[0]}..{labels[-1]}"


def _to_date(v):
    """A date/datetime cell, or a date written as text -> date; anything else -> None."""
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _month_label(d):
    return f"{_MON[d.month - 1]}-{d.year % 100:02d}"


# ----------------------------------------------------------------------------------------
# reading the raw workbook
# ----------------------------------------------------------------------------------------
def _resolve_field(hmap, spec):
    """-> (col or None, matched_by, confidence). `hmap` = {clean heading: column} (last duplicate wins)."""
    for pos, name in enumerate(spec.get("exact", [])):
        col = hmap.get(cc.clean(name))
        if col is not None:
            return col, ("exact" if pos == 0 else "alias"), "HIGH"
    for frag in spec.get("contains", []):
        f = cc.clean(frag)
        for h, col in hmap.items():
            if f in h:
                return col, "contains", "MEDIUM"
    return None, "", ""


def _readme(wb):
    """GSTIN / legal name from a 'Read me' sheet, if the file has one (E-Invoice downloads normally do not)."""
    name = next((n for n in wb.sheetnames if n.lower() == "read me"), None)
    out = {}
    if not name:
        return out
    labels = _mapping()["readme_labels"]
    for row in wb[name].iter_rows(values_only=True):
        cells = [str(c).strip() for c in row if c not in (None, "")]
        if len(cells) < 2:
            continue
        for field, names in labels.items():
            if cells[0].lower() in names:
                out[field] = cells[1]
    return out


def _fail(ctx, path, iid, msg, action="Check that this is the merged E-Invoice workbook downloaded from the portal."):
    ctx.fail(iid, f"E-Invoice \"{os.path.basename(path)}\": {msg}", sheet=TABLE,
             skipped="every check that reads the E-Invoice", action=action)


def build_canonical(raw_paths):
    """Convert one raw/merged E-Invoice workbook into canonical data (a dict of tables). A file that cannot be
    used raises PeriodParseError ('[E50x] ...')."""
    if isinstance(raw_paths, str):
        raw_paths = [raw_paths]
    path = raw_paths[0]
    ctx = cc.IssueLog("E-Invoice", [os.path.basename(p) for p in raw_paths])
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        _fail(ctx, path, "E501", f"cannot open the file ({type(e).__name__}: {e}). It may be corrupt, "
              f"password-protected or not an Excel file.", "Re-download the file.")
    raw_names = [n.lower() for n in _mapping()["raw_names"]]
    fields = _fields()
    try:
        names = list(wb.sheetnames)
        sheet = next((n for n in names if n.lower() in raw_names), None)
        if sheet is None:
            _fail(ctx, path, "E502", f"no '{_mapping()['raw_names'][0]}' tab found (tabs: {names}). "
                  f"This does not look like a merged E-Invoice file.")
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        readme = _readme(wb)
        unread = [n for n in names if n != sheet and n.lower() != "read me"]
    finally:
        wb.close()

    first = next((i for i, r in enumerate(rows) if r and mpu.is_marker_row(r)), None)
    if first is None:
        _fail(ctx, path, "E503", f"tab '{sheet}' has no 'Financial Year | Tax Period' banner rows, so no month can "
              f"be read from it.")
    # The heading row is the row above the first banner that matches the MOST known columns (ties: nearest).
    hdr_idx, best = None, 0
    for i in range(first - 1, max(-1, first - 13), -1):
        if not any(c not in (None, "") for c in rows[i]):
            continue
        hm = {cc.clean(c): j for j, c in enumerate(rows[i]) if c not in (None, "")}
        score = sum(1 for spec in fields.values() if _resolve_field(hm, spec)[0] is not None)
        if hdr_idx is None or score > best:
            hdr_idx, best = i, score
    if hdr_idx is None or best == 0:
        _fail(ctx, path, "E503", f"tab '{sheet}': no recognisable column headings found above the first period "
              f"banner (row {first + 1}).", "Add the new headings to einv_mapping.json.")
    header = rows[hdr_idx]
    hmap = {cc.clean(c): i for i, c in enumerate(header) if c not in (None, "")}
    texts = {i: str(c).strip() for i, c in enumerate(header) if c not in (None, "")}
    cols, mapping, missing_required = {}, [], []
    for f, spec in fields.items():
        col, how, conf = _resolve_field(hmap, spec)
        cols[f] = col
        if col is None and spec.get("required"):
            missing_required.append(f)
        mapping.append(dict(canonical_sheet=LINES_SHEET, canonical_field=f, raw_sheet=sheet,
                            raw_header_found=texts.get(col, "") if col is not None else "",
                            raw_column=get_column_letter(col + 1) if col is not None else "", matched_by=how,
                            confidence=conf,
                            status="OK" if col is not None else ("NOT_IN_PORTAL_FILE" if spec.get("absent_ok") else "MISSING"),
                            rows_filled=0, rows_blank=0, note=""))
    if missing_required:
        names = [fields[f].get("label") or f for f in missing_required]
        _fail(ctx, path, "E504", f"required column(s) {names} not found in tab '{sheet}'. Headings found: "
              f"{sorted(texts.values())}. Every check that reads the E-Invoice is skipped for the run.",
              "Supply the E-Invoice as downloaded from the portal, or add the new heading to einv_mapping.json.")
    optional_missing = [fields[f].get("label") or f for f, c in cols.items()
                        if c is None and not fields[f].get("absent_ok")]
    if optional_missing:
        note = ""
        if cols["status"] is None:
            note = (" Because the e-invoice status column is missing, cancelled e-invoices cannot be identified "
                    "and the cancelled-invoice check reports 'cannot test'.")
        ctx.issue("I501", "INFO", sheet, ", ".join(optional_missing), "all",
                  f"E-Invoice tab '{sheet}': column(s) {optional_missing} not found; they read as blank/0 where "
                  f"a check uses them.{note}")
    mapped = {c for c in cols.values() if c is not None}
    unmapped = {i: t for i, t in texts.items() if i not in mapped}

    table_rows, returns, n_by_month, months_order = [], {}, {}, []
    fill = {f: [0, 0] for f in fields}
    current, current_text, quarter_rows, unassigned = None, "", 0, 0
    for ridx in range(hdr_idx + 1, len(rows)):
        r = rows[ridx]
        if r and mpu.is_marker_row(r):
            try:
                fy, tp, labels = mpu.parse_marker_text(str(r[0]))
            except mpu.PeriodParseError as e:
                _fail(ctx, path, "E506", f"the period banner in row {ridx + 1} is not understood ({e}).")
            current, current_text = labels, str(r[0])
            for lbl in labels:
                returns.setdefault(lbl, dict(month=lbl, fy=fy, tax_period=tp, marker_text=current_text))
                if lbl not in n_by_month:
                    n_by_month[lbl] = 0
                    months_order.append(lbl)
            continue
        if not r or not any(c not in (None, "") for c in r):
            continue
        if current is None:
            _fail(ctx, path, "E505", f"data row {ridx + 1} sits before any period banner, so its month cannot be "
                  f"determined: {r!r}")
        rec = dict(months_covered=";".join(current), src_row=ridx + 1, _any=int(bool(any(r))))
        for f, spec in fields.items():
            col = cols[f]
            if col is None:
                rec[f] = None
                continue
            v = r[col] if col < len(r) else None
            fill[f][0 if v not in (None, "") else 1] += 1
            rec[f] = cc.num(v) if spec["type"] == "num" else v
        rec["extra"] = "; ".join(f"{t}={r[i]}" for i, t in unmapped.items() if i < len(r) and r[i] not in (None, ""))
        if len(current) == 1:
            rec["period"] = current[0]
        else:                                   # quarterly banner: the invoice's own date decides the month
            quarter_rows += 1
            rec["period"] = ""
            for cand in (rec.get("invdate"), rec.get("irndate")):
                d = _to_date(cand)
                if d is not None and _month_label(d) in current:
                    rec["period"] = _month_label(d)
                    break
            if not rec["period"]:
                unassigned += 1
        table_rows.append(rec)
        if rec["period"]:
            n_by_month[rec["period"]] += 1
    for m in mapping:
        if m["canonical_field"] in fill:
            m["rows_filled"], m["rows_blank"] = fill[m["canonical_field"]]
    if quarter_rows:
        ctx.issue("W501", "WARNING", sheet, "invdate", "quarterly",
                  f"E-Invoice tab '{sheet}' has quarterly period banner(s): {quarter_rows} invoice row(s) were "
                  f"placed in the month of their own invoice date (IRN date if that is unusable).",
                  "", "Download the E-Invoice month by month if exact month assignment matters.")
    if unassigned:
        ctx.issue("W502", "WARNING", sheet, "invdate", "quarterly",
                  f"{unassigned} invoice row(s) under a quarterly period banner have no usable invoice date or IRN "
                  f"date inside that quarter, so they are not counted in any month.",
                  "those invoices in every month-wise E-Invoice check", "Download the E-Invoice month by month.")
    if unread:
        ctx.issue("I502", "INFO", "", "", "all", f"Tabs not used by any check: {unread}.")

    fys = sorted({r["fy"] for r in returns.values()})
    gstin = _CONTEXT["gstin"] or readme.get("gstin") or ""
    n_err = sum(1 for i in ctx.issues if i["severity"] == "ERROR")
    n_warn = sum(1 for i in ctx.issues if i["severity"] == "WARNING")
    meta = [("schema_version", SCHEMA_VERSION), ("recipient_gstin", gstin), ("legal_name", readme.get("legal_name", "")),
            ("fy", ", ".join(fys)), ("source_files", "; ".join(ctx.source_files)),
            ("source_layout", "merged with period banners"), ("adapter_version", ADAPTER_VERSION),
            ("created_on", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("months_present", ";".join(months_order)), ("unread_tabs", "; ".join(unread)),
            ("status", "OK_WITH_ERRORS" if n_err else ("OK_WITH_WARNINGS" if n_warn else "OK"))]
    coverage = [dict(table=TABLE, month=m, rows=n_by_month[m], note="") for m in months_order]
    return dict(meta=meta, returns=list(returns.values()), coverage=coverage, lines=table_rows, mapping=mapping,
                issues=ctx.issues)


# ----------------------------------------------------------------------------------------
# write / read the canonical workbook
# ----------------------------------------------------------------------------------------
_WRITTEN = {}       # canonical path -> source files it was written from (this process)


def canonical_filename(data):
    return cc.canonical_filename("EINV", data["meta"])


def write_canonical(data, out_dir):
    """Write Canonical_EINV_<GSTIN>_<FY>.xlsx. Two different E-Invoice files of the same GSTIN/FY in one run
    get '_2', '_3'... so the second never overwrites the first."""
    name = canonical_filename(data)
    src = dict(data["meta"]).get("source_files")
    base, ext = os.path.splitext(name)
    n = 1
    while True:
        cand = os.path.join(out_dir, name)
        if _WRITTEN.get(os.path.abspath(cand), src) == src:
            break
        n += 1
        name = f"{base}_{n}{ext}"
    sheets = [("RETURNS", RETURN_COLS, data["returns"]), ("COVERAGE", COVERAGE_COLS, data["coverage"]),
              (LINES_SHEET, _canon_cols(), data["lines"]),
              ("MAPPING_REPORT", MAP_COLS, data["mapping"]), ("ISSUES", ISSUE_COLS, data["issues"])]
    path = cc.write_workbook(out_dir, name, data["meta"], sheets, max_col_width=30)
    _WRITTEN[os.path.abspath(path)] = src
    return path


def is_canonical_file(path):
    return cc.is_canonical(path, "EINV-", ("META", "RETURNS", "COVERAGE", LINES_SHEET))


def read_canonical(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return dict(meta=cc.read_meta(wb), returns=cc.read_table(wb, "RETURNS", RETURN_COLS),
                    coverage=cc.read_table(wb, "COVERAGE", COVERAGE_COLS),
                    lines=cc.read_table(wb, LINES_SHEET, _canon_cols()),
                    mapping=cc.read_table(wb, "MAPPING_REPORT", MAP_COLS),
                    issues=cc.read_table(wb, "ISSUES", ISSUE_COLS))
    finally:
        wb.close()


# ----------------------------------------------------------------------------------------
# engine-facing access
# ----------------------------------------------------------------------------------------
def _index(data):
    by_month = {}
    for r in data["lines"]:
        if r["period"]:
            by_month.setdefault(str(r["period"]), []).append(r)
    ok = {m["canonical_field"]: m["status"] == "OK" for m in data["mapping"]}
    months = []
    for c in data["coverage"]:
        if c["month"] not in months:
            months.append(c["month"])
    return dict(meta=dict(data["meta"]), by_month=by_month, ok=ok, months=months)


_SRC = cc.CanonicalCache("E-Invoice", is_canonical_file, read_canonical, build_canonical, write_canonical, _index)


def get_data(path, out_dir=None):
    return _SRC.get(path, out_dir)


def clear_cache():
    _SRC.clear()
    _WRITTEN.clear()


def months(path):
    """Months that have a period banner in the E-Invoice tab (legacy gst_core._einv_months)."""
    return set(get_data(path)["_index"]["months"])


def _rows(ix, month):
    return [r for r in ix["by_month"].get(month, []) if r["_any"]]


def _f(v):
    return 0.0 if v is None else float(v)


def _s(v):
    return str(v or "").strip()


def _str_or_none(v):
    """Legacy `str(cell).strip()` - a blank cell reads as the text 'None'."""
    return str(v).strip()


# ---- parse_einv -----------------------------------------------------------------------------------
def parse_einv(path, month):
    """Legacy gst_parsers_returns.parse_einv(path, month)."""
    out = {"taxable": 0.0, "IGST": 0.0, "CGST": 0.0, "SGST": 0.0, "CESS": 0.0, "count": 0, "errors": 0,
           "available": True, "lines": {}, "cancel_col_found": False, "cancel_date_col_found": False,
           "cancelled": []}
    if not path or not os.path.exists(path):
        print(f"[info] E-Invoice file not supplied -> EINV checks skipped for {month}")
        out["available"] = False
        return out
    ix = get_data(path)["_index"]
    if month not in ix["months"]:
        print(f"[info] E-Invoice file does not cover {month} -> EINV checks skipped for {month}")
        out["available"] = False
        return out
    ok = ix["ok"]
    cancelled_values = tuple(_mapping()["cancelled_values"])
    out["cancel_col_found"] = ok["status"]
    out["cancel_date_col_found"] = ok["cancel_date"]
    for r in _rows(ix, month):
        invno = _str_or_none(r["invno"]) if ok["invno"] else "None"
        rate = _f(r["rate"]) if ok["rate"] else 0.0
        gstin_key = _str_or_none(r["gstin"]) if ok["gstin"] else ""
        if ok["status"] and _s(r["status"]).upper() in cancelled_values:
            out["cancelled"].append(dict(
                invno=invno, rate=rate, gstin=gstin_key,
                invdate=_s(r["invdate"]) if ok["invdate"] else "",
                taxable=_f(r["taxable"]), igst=_f(r["igst"]), cgst=_f(r["cgst"]), sgst=_f(r["sgst"]),
                cess=_f(r["cess"]), irn=_s(r["irn"]) if ok["irn"] else "",
                cancel_date=_str_or_none(r["cancel_date"]) if ok["cancel_date"] else None, month=month))
            continue
        out["taxable"] += _f(r["taxable"])
        out["IGST"] += _f(r["igst"])
        out["CGST"] += _f(r["cgst"])
        out["SGST"] += _f(r["sgst"])
        out["CESS"] += _f(r["cess"])
        out["count"] += 1
        L = out["lines"].setdefault((invno, gstin_key, rate), [0.0, 0.0])
        L[0] += _f(r["taxable"])
        L[1] += _f(r["igst"])
        if ok["error"] and _s(r["error"]):
            out["errors"] += 1
    return out


# ---- line / invoice readers (gst_checks_monthly) --------------------------------------------------
def read_einv_lines(path, month):
    """Legacy gst_checks_monthly.read_einv_lines(path, month)."""
    from gst_checks_monthly import _parse_date
    ix = get_data(path)["_index"]
    if month not in ix["months"]:
        return []
    return [dict(gstin=_s(r["gstin"]), invno=_s(r["invno"]), invdate=_parse_date(r["invdate"]), pos=_s(r["pos"]),
                 rate=_f(r["rate"]), taxable=_f(r["taxable"]), igst=_f(r["igst"]), cgst=_f(r["cgst"]),
                 sgst=_f(r["sgst"]), rcm=_s(r["rcm"]), irn=_s(r["irn"]), irndate=_parse_date(r["irndate"]),
                 err=_s(r["error"])) for r in _rows(ix, month)]


def read_einv_invoices(path, month):
    """Legacy gst_checks_monthly.read_einv_invoices(path, month)."""
    ix = get_data(path)["_index"]
    out = {}
    if month not in ix["months"]:
        return out
    for r in _rows(ix, month):
        d = out.setdefault(_s(r["invno"]), dict(taxable=0.0, igst=0.0, cgst=0.0, sgst=0.0, invval=0.0))
        d["taxable"] += _f(r["taxable"])
        d["igst"] += _f(r["igst"])
        d["cgst"] += _f(r["cgst"])
        d["sgst"] += _f(r["sgst"])
        d["invval"] += _f(r["invval"])
    return out
